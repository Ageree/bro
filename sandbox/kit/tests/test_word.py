"""brokit docx fill / edit / compare / text (W2 spec §6.11, §9)."""

import json
import os
import tempfile
import unittest

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from tests.cli_run import kit

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def numbered(document, text, num_id="7", level="1"):
    paragraph = document.add_paragraph(text)
    ppr = paragraph._p.get_or_add_pPr()
    num = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), level)
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), num_id)
    num.append(ilvl)
    num.append(num_id_element)
    ppr.append(num)
    return paragraph


def contract(path, payment="10 (десяти) календарных", extra=True, drop=False, account="40702810000000000111"):
    document = docx.Document()
    document.add_paragraph("ДОГОВОР ПОСТАВКИ № 1")
    numbered(document, "Покупатель оплачивает Товар в течение %s дней с даты поставки." % payment)
    numbered(document, "Сторона, нарушившая сроки поставки, уплачивает пени 0,1 % за каждый день.")
    numbered(document, "Уплата неустойки не освобождает Стороны от исполнения обязательств.")
    if not drop:
        numbered(document, "Совокупная ответственность Исполнителя ограничена суммой договора.")
    if extra:
        document.add_paragraph("Споры рассматриваются в Арбитражном суде города Москвы.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "ИНН 7700000086\nр/с %s\nБИК 044500111" % account
    table.cell(0, 1).text = "ИНН 7800000000"
    document.save(path)


class DocxTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.TemporaryDirectory()

    @classmethod
    def tearDownClass(cls):
        cls.dir.cleanup()

    def path(self, name):
        return os.path.join(self.dir.name, name)

    def write_json(self, name, value):
        with open(self.path(name), "w", encoding="utf-8") as handle:
            json.dump(value, handle, ensure_ascii=False)
        return self.path(name)

    def text(self, path, changes="accept"):
        code, result = kit("docx", "text", path, "--changes", changes)
        self.assertEqual(code, 0, result)
        return result["text"]

    def test_fill_replaces_placeholders_split_across_runs(self):
        document = docx.Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("вручается ")
        bold = paragraph.add_run("{учени")
        bold.bold = True
        paragraph.add_run("ку/ученице} 5 «Б» класса")
        document.add_paragraph().add_run("{{ ФИО }}")
        document.add_paragraph("Дата: ________ г.")
        document.sections[0].header.paragraphs[0].text = "Грамота для {ФИО}"
        document.save(self.path("tpl.docx"))
        values = self.write_json("values.json", {"ученику/ученице": "ученице", "ФИО": "Вороне Ангелине", "Дата": "24.12.2026", "Класс": "5Б"})
        code, result = kit("docx", "fill", self.path("tpl.docx"), values, "-o", self.path("filled.docx"))
        self.assertEqual(code, 0, result)
        self.assertEqual(result["unknown"], ["Класс"])
        self.assertEqual(result["left"], [])
        filled = docx.Document(self.path("filled.docx"))
        first = filled.paragraphs[0]
        self.assertEqual(first.text, "вручается ученице 5 «Б» класса")
        self.assertTrue(next(run for run in first.runs if "ученице" in run.text).bold)
        self.assertEqual(filled.paragraphs[1].text, "Вороне Ангелине")
        date_run = next(run for run in filled.paragraphs[2].runs if "24.12.2026" in run.text)
        self.assertTrue(date_run.underline)
        self.assertEqual(filled.sections[0].header.paragraphs[0].text, "Грамота для Вороне Ангелине")

    def test_fill_with_a_list_makes_one_file_per_item(self):
        document = docx.Document()
        document.add_paragraph("{ФИО}")
        document.save(self.path("merge.docx"))
        values = self.write_json("rows.json", [{"_file": "Иванову.docx", "ФИО": "Иванову Петру"}, {"ФИО": "Ким Алине"}])
        code, result = kit("docx", "fill", self.path("merge.docx"), values, "-o", self.path("merged"))
        self.assertEqual(code, 0, result)
        names = sorted(os.listdir(self.path("merged")))
        self.assertEqual(names, ["002.docx", "Иванову.docx"])
        self.assertEqual(docx.Document(self.path("merged/002.docx")).paragraphs[0].text, "Ким Алине")

    def test_edit_writes_tracked_changes_and_numbered_insertion(self):
        contract(self.path("contract.docx"))
        edits = self.write_json(
            "edits.json",
            [
                {"replace": {"in_paragraph": "Покупатель оплачивает", "old": "10 (десяти) календарных", "new": "5 (пяти) банковских"}},
                {"insert_after": {"anchor": "^Сторона, нарушившая", "text": "Стороны сохраняют конфиденциальность условий Договора.", "same_numbering": True}},
                {"replace_all": {"old": "40702810000000000111", "new": "40702810000000000222"}},
                {"replace": {"in_paragraph": "нет такого", "old": "x", "new": "y"}},
            ],
        )
        code, result = kit("docx", "edit", self.path("contract.docx"), edits, "-o", self.path("edited.docx"), "--author", "Бро", "--date", "2026-10-15T10:00:00Z")
        self.assertEqual(code, 0, result)
        self.assertEqual([a["status"] for a in result["applied"]], ["ok", "ok", "ok", "not_found"])
        self.assertEqual(result["not_found"], [4])
        body = docx.Document(self.path("edited.docx")).element.body
        deleted = ["".join(t.text for t in d.iter(W + "delText")) for d in body.iter(W + "del")]
        inserted = ["".join(t.text or "" for t in i.iter(W + "t")) for i in body.iter(W + "ins")]
        self.assertIn("10 (десяти) календарных", deleted)
        self.assertIn("5 (пяти) банковских", inserted)
        self.assertIn("40702810000000000111", deleted)
        self.assertIn("40702810000000000222", inserted)
        for change in list(body.iter(W + "del")) + list(body.iter(W + "ins")):
            self.assertEqual(change.get(W + "author"), "Бро")
            self.assertEqual(change.get(W + "date"), "2026-10-15T10:00:00Z")
        ids = [c.get(W + "id") for c in list(body.iter(W + "del")) + list(body.iter(W + "ins"))]
        self.assertEqual(len(ids), len(set(ids)))
        paragraphs = list(body.iter(W + "p"))
        anchor = next(p for p in paragraphs if "Сторона, нарушившая" in "".join(t.text or "" for t in p.iter(W + "t")))
        new = anchor.getnext()
        # The whole paragraph is inserted, its mark too, and it is numbered like 6.2.
        self.assertEqual([child.tag for child in new if child.tag != W + "pPr"], [W + "ins"])
        self.assertIsNotNone(new.find("%spPr/%srPr/%sins" % (W, W, W)))
        self.assertEqual(new.find("%spPr/%snumPr/%snumId" % (W, W, W)).get(W + "val"), "7")
        self.assertEqual(new.find("%spPr/%snumPr/%silvl" % (W, W, W)).get(W + "val"), "1")
        # ИНН in the same cell is untouched.
        self.assertNotIn("7700000086", "".join(deleted))
        self.assertEqual(self.text(self.path("edited.docx"), "reject"), self.text(self.path("contract.docx")))
        accepted = self.text(self.path("edited.docx"), "accept")
        self.assertIn("в течение 5 (пяти) банковских дней", accepted)
        self.assertIn("за каждый день.\nСтороны сохраняют конфиденциальность условий Договора.\nУплата неустойки", accepted)

    def test_compare_builds_a_redline_that_accepts_to_b_and_rejects_to_a(self):
        contract(self.path("v1.docx"))
        contract(self.path("v2.docx"), payment="30 (тридцати) календарных", drop=True, account="40702810000000000222")
        document = docx.Document(self.path("v2.docx"))
        document.add_paragraph("Новый пункт о форс-мажоре.")
        document.save(self.path("v2.docx"))
        code, result = kit("docx", "compare", self.path("v1.docx"), self.path("v2.docx"), "-o", self.path("redline.docx"))
        self.assertEqual(code, 0, result)
        self.assertEqual(self.text(self.path("redline.docx"), "accept"), self.text(self.path("v2.docx")))
        self.assertEqual(self.text(self.path("redline.docx"), "reject"), self.text(self.path("v1.docx")))
        body = docx.Document(self.path("redline.docx")).element.body
        deleted = ["".join(t.text for t in d.iter(W + "delText")) for d in body.iter(W + "del")]
        inserted = ["".join(t.text or "" for t in i.iter(W + "t")) for i in body.iter(W + "ins")]
        self.assertIn("10 (десяти)", deleted)
        self.assertIn("30 (тридцати)", inserted)
        self.assertIn("40702810000000000111", deleted)
        self.assertTrue(any("Совокупная ответственность" in text for text in deleted))
        self.assertTrue(all(c.get(W + "author") == "Бро" for c in body.iter(W + "ins")))
        shown = self.text(self.path("redline.docx"), "show")
        self.assertIn("{+Новый пункт о форс-мажоре.+}", shown)

    def test_old_doc_is_refused_with_a_hint(self):
        with open(self.path("old.doc"), "wb") as handle:
            handle.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504)
        code, result = kit("docx", "text", self.path("old.doc"))
        self.assertEqual((code, result["code"]), (2, "unsupported"))
        self.assertIn("soffice", result["hint"])


if __name__ == "__main__":
    unittest.main()
