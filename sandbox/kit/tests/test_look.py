"""`brokit look`: pictures within the budget, the adaptive grid, zoomed regions."""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
POPPLER = shutil.which("pdftoppm") is not None and shutil.which("pdfinfo") is not None


def brokit(*args, cache):
    """The kit as the model's wrapper runs it: another process, `-I`, the limits of `look` on."""
    code = "import sys;sys.path.insert(0,%r);from brokit.cli import main;sys.exit(main(sys.argv[1:]))" % KIT
    result = subprocess.run(
        [sys.executable, "-I", "-c", code, *args],
        env={"PATH": os.environ.get("PATH", "/usr/bin:/bin"), "BROKIT_CACHE": cache, "HOME": cache + "-home"},
        capture_output=True,
        text=True,
        timeout=180,
    )
    return result.returncode, json.loads(result.stdout.strip().splitlines()[-1])


@unittest.skipUnless(POPPLER, "poppler-utils are not installed")
class Look(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import numpy
        from PIL import Image
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        cls.root = tempfile.mkdtemp(prefix="brokit-look-")
        cls.cache = os.path.join(cls.root, "cache")
        cls.pdf = os.path.join(cls.root, "six.pdf")
        document = canvas.Canvas(cls.pdf, pagesize=A4)
        for page in range(1, 7):
            document.setFont("Helvetica", 11)
            for line in range(40):
                document.drawString(60, 780 - line * 18, "Page %d, line %d: the quick brown fox jumps over the lazy dog." % (page, line))
            document.showPage()
        document.save()
        # A phone photo's size whose noise survives the downscale: the hardest JPEG to keep small.
        noise = numpy.random.default_rng(7).integers(0, 256, size=(4032, 3024, 3), dtype=numpy.uint8)
        cls.photo = os.path.join(cls.root, "photo.png")
        Image.fromarray(noise).save(cls.photo, compress_level=1)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.root, True)

    def out(self):
        return tempfile.mkdtemp(dir=self.root)

    def test_four_pictures_at_most_and_where_to_go_on(self):
        out = self.out()
        code, result = brokit("look", self.pdf, "--out", out, cache=self.cache)
        self.assertEqual(code, 0, result)
        (item,) = result["items"]
        self.assertEqual((item["kind"], item["unit"], item["pages_total"]), ("pdf", "pt", 6))
        self.assertEqual([image["page"] for image in item["images"]], [1, 2, 3, 4])
        self.assertTrue(result["truncated"])
        self.assertEqual(result["next"], {"path": self.pdf, "pages": "5-6"})
        for image in item["images"]:
            self.assertEqual(max(image["width"], image["height"]), 1280)
            self.assertLessEqual(os.path.getsize(image["file"]), 400_000)
            self.assertGreater(image["text_chars"], 200)

    def test_a_noisy_photo_stays_within_the_budget(self):
        code, result = brokit("look", self.photo, self.photo, self.photo, "--out", self.out(), cache=self.cache)
        self.assertEqual(code, 0, result)
        files = [image["file"] for item in result["items"] for image in item["images"]]
        sizes = [os.path.getsize(path) for path in files]
        self.assertTrue(all(size <= 400_000 for size in sizes), sizes)
        self.assertLessEqual(sum(sizes), 1_200_000)
        from PIL import Image

        for path in files:
            with Image.open(path) as picture:
                self.assertLessEqual(max(picture.size), 1280)

    def test_the_grid_step_follows_the_page_and_the_region(self):
        code, result = brokit("look", self.pdf, "--out", self.out(), "--pages", "2", "--grid", cache=self.cache)
        self.assertEqual(code, 0, result)
        page = result["items"][0]["images"][0]
        self.assertEqual(page["grid_step"], 50)
        code, result = brokit("look", self.photo, "--out", self.out(), "--grid", cache=self.cache)
        self.assertEqual(result["items"][0]["images"][0]["grid_step"], 250)
        self.assertEqual(result["items"][0]["images"][0]["view"], {"width": 3024, "height": 4032})
        code, result = brokit("look", self.pdf, "--out", self.out(), "--pages", "2", "--grid", "--region", "60,40,260,120", cache=self.cache)
        self.assertEqual(code, 0, result)
        zoomed = result["items"][0]["images"][0]
        self.assertEqual(zoomed["region"], [60, 40, 260, 120])
        self.assertLess(zoomed["grid_step"], page["grid_step"])
        self.assertGreater(zoomed["scale"], page["scale"] * 3)

    def test_render_draws_the_chosen_pages(self):
        out = self.out()
        code, result = brokit("render", self.pdf, "-o", out, "--pages", "2-3", "--dpi", "72", cache=self.cache)
        self.assertEqual(code, 0, result)
        self.assertIsNone(result["pdf"])
        self.assertEqual([page["page"] for page in result["files"]], [2, 3])
        for page in result["files"]:
            # 72 dpi: a pixel per point of A4, as poppler rounds it.
            self.assertLessEqual(abs(page["width"] - 595.3), 1)
            self.assertLessEqual(abs(page["height"] - 841.9), 1)
            self.assertTrue(os.path.isfile(page["path"]))

    def test_refusals_name_what_to_do(self):
        code, result = brokit("look", os.path.join(self.root, "sixx.pdf"), "--out", self.out(), cache=self.cache)
        self.assertEqual((code, result["code"]), (2, "not_found"))
        self.assertIn("six.pdf", result["hint"])
        heic = os.path.join(self.root, "photo.heic")
        with open(heic, "wb") as handle:
            handle.write(b"\0\0\0\x18ftypheic\0\0\0\0mif1heic" + b"\0" * 64)
        code, result = brokit("look", heic, "--out", self.out(), cache=self.cache)
        self.assertEqual((code, result["code"]), (2, "unsupported"))
        code, result = brokit("look", self.pdf, self.photo, "--region", "0,0,10,10", "--out", self.out(), cache=self.cache)
        self.assertEqual((code, result["code"]), (2, "bad_args"))

    def test_a_small_picture_is_shown_as_it_is(self):
        from PIL import Image

        icon = os.path.join(self.root, "stamp.png")
        Image.new("RGB", (200, 150), "red").save(icon)
        for path, region in ((icon, None), (self.photo, "100,100,140,130")):
            args = ["look", path, "--out", self.out()] + (["--region", region] if region else [])
            code, result = brokit(*args, cache=self.cache)
            self.assertEqual(code, 0, result)
            self.assertFalse(result["truncated"], result)
            (image,) = result["items"][0]["images"]
            self.assertTrue(os.path.isfile(image["file"]))

    def test_only_plain_image_formats_are_opened(self):
        # Pillow opens EPS too, and draws it through Ghostscript: a person's file is no program to run.
        eps = os.path.join(self.root, "drawing.eps")
        with open(eps, "w", encoding="ascii") as handle:
            handle.write("%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 10 10\nshowpage\n")
        code, result = brokit("look", eps, "--out", self.out(), cache=self.cache)
        self.assertEqual((code, result["code"]), (2, "unsupported"))
        # A phone's MPO is a JPEG whose second frame is a depth map: one picture, not two pages.
        from PIL import Image

        mpo = os.path.join(self.root, "phone.jpg")
        Image.new("RGB", (640, 480), "green").save(mpo, format="MPO", save_all=True, append_images=[Image.new("RGB", (320, 240))])
        code, result = brokit("look", mpo, "--out", self.out(), cache=self.cache)
        self.assertEqual(code, 0, result)
        self.assertEqual((result["items"][0]["pages_total"], len(result["items"][0]["images"])), (1, 1))

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice is not installed")
    def test_an_office_file_is_seen_as_its_pages(self):
        import docx

        path = os.path.join(self.root, "letter.docx")
        document = docx.Document()
        for number in range(3):
            document.add_paragraph("Абзац %d письма." % number)
        document.save(path)
        code, result = brokit("look", path, "--out", self.out(), cache=self.cache)
        self.assertEqual(code, 0, result)
        item = result["items"][0]
        self.assertEqual((item["kind"], item["unit"], item["pages_total"]), ("office", "pt", 1))
        self.assertEqual(len(os.listdir(os.path.join(self.cache, "render"))), 1)

    @unittest.skipUnless(shutil.which("soffice"), "LibreOffice is not installed")
    def test_a_deck_with_a_wall_of_text_is_named_slide_by_slide(self):
        import pptx
        from pptx.util import Inches

        def words(count):
            return " ".join("слово" for _ in range(count))

        def deck(name, dense):
            presentation = pptx.Presentation()
            for number in range(1, 4):
                slide = presentation.slides.add_slide(presentation.slide_layouts[5])
                slide.shapes.title.text = "Итоги квартала, часть %d" % number
                if number == 2 and dense:
                    # 4 + 56 + 60 = 120: a grouped text box and table cells count, not only text frames.
                    group = slide.shapes.add_group_shape()
                    group.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2)).text_frame.text = words(56)
                    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4.5), Inches(8), Inches(2)).table
                    for cell in (table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)):
                        cell.text = words(15)
                else:
                    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2)).text_frame.text = "— выручка растёт\n— расходы падают"
                    # What the speaker says is no text on the slide.
                    slide.notes_slide.notes_text_frame.text = words(200)
            path = os.path.join(folder, name)
            presentation.save(path)
            return path

        # A folder and cache of its own: other tests count the conversions in the shared
        # cache and expect six.pdf among the first files of the root.
        folder = os.path.join(self.root, "slides")
        os.makedirs(folder, exist_ok=True)
        cache = os.path.join(folder, "cache")
        dense = deck("board.pptx", dense=True)
        code, result = brokit("look", dense, "--out", self.out(), "--pages", "1", cache=cache)
        self.assertEqual(code, 0, result)
        (warning,) = [text for text in result["warnings"] if "слайд" in text]
        self.assertIn("слайд 2: 120 слов", warning)
        self.assertNotIn("слайд 1", warning)
        self.assertNotIn("слайд 3", warning)
        self.assertIn("speaker notes", warning)
        light = deck("light.pptx", dense=False)
        code, result = brokit("look", light, "--out", self.out(), "--pages", "1", cache=cache)
        self.assertEqual(code, 0, result)
        self.assertEqual(result["warnings"], [])


if __name__ == "__main__":
    unittest.main()
