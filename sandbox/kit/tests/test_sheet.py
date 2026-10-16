"""brokit xlsx inspect / recalc and csv sniff (W2 spec §6.11, §9)."""

import datetime
import os
import tempfile
import unittest
import zipfile

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from tests.cli_run import has, kit


def workbook(path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Январь"
    sheet["A1"] = "Продажи за январь"
    sheet.merge_cells("A1:C1")
    sheet.append(["Менеджер", "Сумма", "Статус"])
    sheet.append(["Иванов", 1500.5, "оплачено"])
    sheet.append(["Ким", 200, "отменена"])
    sheet["B5"] = "=SUM(B3:B4)"
    sheet["B6"] = "=B5/0"
    rule = DataValidation(type="list", formula1='"оплачено,отменена"')
    rule.add("C3:C4")
    sheet.add_data_validation(rule)
    hidden = book.create_sheet("old")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "старое"
    book.save(path)


def statement(path):
    """A bank statement as the task agent built it: dates and sums typed as text."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Операции"
    sheet.append(["Дата", "Сумма", "ИНН", "Счёт", "Телефон", "Проведено", "Списано", "Описание"])
    for n in range(1, 8):
        sheet.append(
            [
                "%02d.09.2026 08:%02d" % (n, n),
                "\u2212%d %03d,%02d" % (n, n * 7, n) if n % 2 else "%d,50 ₽" % (n * 100),
                "77070838%02d" % n,
                "4081781009991000%04d" % n,
                "+7916123456%d" % n,
                datetime.datetime(2026, 9, n, 8, n),
                -1234.5 * n,
                "Перевод %d,5 тыс." % n,
            ]
        )
    sheet["G9"] = "1 234,56"
    totals = book.create_sheet("Итоги")
    totals.append(["Остаток на начало", "154 215,37"])
    totals.append(["Остаток на конец", "69 335,42"])
    book.save(path)


def timetable(path):
    """A weekly timetable: days across row 2, times down column A, a lesson merged over its slots."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Неделя"
    sheet["A1"] = "Расписание кружков"
    sheet.merge_cells("A1:H1")
    for column, day in enumerate(["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"], 2):
        sheet.cell(2, column, day)
    for index in range(10):
        sheet.cell(3 + index, 1, datetime.time(*divmod(9 * 60 + 30 * index, 60)))
    sheet["B3"] = "Английский\nкаб. 12"
    sheet.merge_cells("B3:B5")
    sheet["E8"] = "Шахматы"
    book.save(path)


def with_stale_value(source, target, cell, value):
    """A copy whose formula cell carries a wrong cached value, as Excel files often do."""
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode("utf-8")
    marker = '<c r="%s"><f>SUM(B3:B4)</f><v></v></c>' % cell
    assert marker in sheet, sheet
    parts["xl/worksheets/sheet1.xml"] = sheet.replace(marker, marker.replace("<v></v>", "<v>%s</v>" % value)).encode()
    with zipfile.ZipFile(target, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


class SheetTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.TemporaryDirectory()
        cls.book = os.path.join(cls.dir.name, "sales.xlsx")
        workbook(cls.book)

    @classmethod
    def tearDownClass(cls):
        cls.dir.cleanup()

    def path(self, name):
        return os.path.join(self.dir.name, name)

    def test_inspect_lists_what_the_workbook_holds(self):
        code, result = kit("xlsx", "inspect", self.book)
        self.assertEqual(code, 0, result)
        first, hidden = result["sheets"]
        self.assertEqual(first["name"], "Январь")
        self.assertEqual(first["merged"], ["A1:C1"])
        self.assertEqual(
            first["merged_values"],
            [{"range": "A1:C1", "value": "Продажи за январь", "rows": [1, 1], "row_labels": [None, None], "column_label": None}],
        )
        # A title above the header is not data a merge hides.
        self.assertNotIn("hint", first)
        self.assertEqual(first["header_row"], 2)
        self.assertEqual(first["formulas"]["count"], 2)
        self.assertEqual(first["validations"][0]["ranges"], "C3:C4")
        self.assertEqual((hidden["name"], hidden["hidden"]), ("old", True))
        self.assertEqual((result["text_dates"], result["text_numbers"]), ([], []))
        self.assertNotIn("hint", result)

    def test_inspect_names_columns_of_dates_and_sums_typed_as_text(self):
        statement(self.path("statement.xlsx"))
        code, result = kit("xlsx", "inspect", self.path("statement.xlsx"))
        self.assertEqual(code, 0, result)
        self.assertEqual(
            result["text_dates"], [{"sheet": "Операции", "column": "A", "cells": 7, "example": "01.09.2026 08:01"}]
        )
        # ИНН, account and phone columns stay text; one text sum among the
        # numbers of «Списано» is not a text column.
        self.assertEqual(
            result["text_numbers"],
            [
                {"sheet": "Операции", "column": "B", "cells": 7, "example": "\u22121 007,01"},
                {"sheet": "Итоги", "column": "B", "cells": 2, "example": "154 215,37"},
            ],
        )
        self.assertIn("'DD.MM.YYYY HH:MM'", result["hint"])

    def test_inspect_places_merged_values_of_a_timetable(self):
        timetable(self.path("timetable.xlsx"))
        code, result = kit("xlsx", "inspect", self.path("timetable.xlsx"))
        self.assertEqual(code, 0, result)
        sheet = result["sheets"][0]
        # The day row is the header although the rows under it hold a time and at most one lesson.
        self.assertEqual(sheet["header_row"], 2)
        lesson = next(item for item in sheet["merged_values"] if item["range"] == "B3:B5")
        self.assertEqual(
            lesson,
            {
                "range": "B3:B5",
                "value": "Английский\nкаб. 12",
                "rows": [3, 5],
                "row_labels": ["09:00:00", "10:00:00"],
                "column_label": "Пн",
            },
        )
        self.assertIn("covers its whole range", sheet["hint"])

    def test_csv_sniff_reads_cp1251_semicolons_and_decimal_commas(self):
        lines = ["Дата;Товар;Кол-во;Цена"]
        for n in range(1, 300):
            lines.append("01.01.2025;Творог 5%% 200 г;%d;%d,90" % (n % 4 + 1, 100 + n))
        lines[120] = '01.01.2025;"Печенье ""Юбилейное; традиционное""";1;79,00'
        lines[200] = '01.01.2025;Сок "Сады; мультифрукт" 1 л;1;99,00'
        with open(self.path("kassa.csv"), "wb") as handle:
            handle.write("\r\n".join(lines).encode("cp1251") + b"\r\n")
        code, result = kit("csv", "sniff", self.path("kassa.csv"))
        self.assertEqual(code, 0, result)
        self.assertEqual((result["encoding"], result["sep"], result["decimal"]), ("cp1251", ";", ","))
        self.assertEqual(result["header"], ["Дата", "Товар", "Кол-во", "Цена"])
        self.assertEqual(result["rows"], 299)
        self.assertEqual(result["bad_lines"], {"count": 1, "lines": [201]})
        self.assertIn("encoding='cp1251'", result["pandas"])

    def test_csv_sniff_names_the_line_of_a_quote_that_never_closes(self):
        lines = ["Дата;Товар;Сумма", '01.01.2025;"Сок 1 л;99,00']
        lines += ["01.01.2025;Хлеб;%d,50" % n for n in range(9000)]
        with open(self.path("open_quote.csv"), "wb") as handle:
            handle.write("\n".join(lines).encode("cp1251"))
        code, result = kit("csv", "sniff", self.path("open_quote.csv"))
        self.assertEqual(code, 0, result)
        self.assertEqual(result["stopped"]["line"], 2)

    @unittest.skipUnless(has("soffice"), "LibreOffice is not installed")
    def test_recalc_recomputes_stale_values_and_reports_errors(self):
        with_stale_value(self.book, self.path("stale.xlsx"), "B5", "999")
        code, result = kit("xlsx", "recalc", self.path("stale.xlsx"), "-o", self.path("recalc.xlsx"))
        self.assertEqual(code, 0, result)
        self.assertEqual(result["errors"], [{"sheet": "Январь", "cell": "B6", "value": "#DIV/0!"}])
        values = openpyxl.load_workbook(self.path("recalc.xlsx"), data_only=True)["Январь"]
        self.assertEqual(values["B5"].value, 1700.5)
        formulas = openpyxl.load_workbook(self.path("recalc.xlsx"))["Январь"]
        self.assertEqual(formulas["B5"].value, "=SUM(B3:B4)")
        code, result = kit("xlsx", "recalc", self.path("stale.xlsx"), "-o", self.path("checked.xlsx"), "--check")
        self.assertEqual((code, result["code"]), (2, "formula_errors"))

    @unittest.skipUnless(has("soffice"), "LibreOffice is not installed")
    def test_recalc_names_dates_and_sums_typed_as_text(self):
        statement(self.path("statement-recalc.xlsx"))
        code, result = kit("xlsx", "recalc", self.path("statement-recalc.xlsx"), "-o", self.path("statement-out.xlsx"))
        self.assertEqual(code, 0, result)
        self.assertEqual([(c["sheet"], c["column"]) for c in result["text_dates"]], [("Операции", "A")])
        self.assertEqual([(c["sheet"], c["column"]) for c in result["text_numbers"]], [("Операции", "B"), ("Итоги", "B")])


if __name__ == "__main__":
    unittest.main()
