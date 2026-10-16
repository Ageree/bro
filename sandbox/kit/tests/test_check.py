"""brokit check: what share_file holds a finished file back for, and what it lets through."""

import datetime
import os
import shutil
import tempfile
import unittest
import zipfile
from unittest import mock

import numpy as np
import openpyxl
from PIL import Image, ImageDraw

from tests import test_img_scan
from tests.cli_run import kit
from tests.pdf_fixtures import kit as kit_apart
from tests.test_sheet import statement


def check(path):
    """`brokit check` in a child process, as the sandbox runs it: its memory limit stays out of this one."""
    code, result = kit_apart("check", path)
    assert code == 0, result
    return result


def codes(result):
    return [problem["code"] for problem in result["problems"]]


def printed_page():
    """A page of print: lines of short words, as letters blur together on a phone photo."""
    rng = np.random.default_rng(5)
    page = Image.new("RGB", test_img_scan.PAGE, "white")
    draw = ImageDraw.Draw(page)
    for y in range(150, 1600, 32):
        x = 120
        while x < 1080:
            width = int(rng.integers(20, 90))
            draw.rectangle([x, y, x + width, y + 12], fill=(25, 25, 25))
            x += width + 16
    return page


def photo(path):
    """A phone photo of the printed page lying askew on a dark table (test_img_scan's warp)."""
    with mock.patch.object(test_img_scan, "page_image", printed_page):
        test_img_scan.photo(path, test_img_scan.CORNERS)


def sky(path):
    """Houses under a sky: as bright as paper and with dark marks in it, but running off the frame's edges."""
    rng = np.random.default_rng(9)
    image = Image.new("RGB", test_img_scan.PHOTO, (60, 62, 70))
    draw = ImageDraw.Draw(image)
    draw.rectangle([0, 0, test_img_scan.PHOTO[0], 1150], fill=(196, 214, 236))
    for _ in range(150):
        x, y = int(rng.integers(100, 1400)), int(rng.integers(100, 1000))
        draw.rectangle([x, y, x + 14, y + 10], fill=(70, 70, 80))
    image.save(path, quality=90)


def captioned_pdf(photo_path, path):
    """The photo «по центру A4» with a caption above it, as the task agent laid out receipts on 16.10."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(path, pagesize=A4)
    pdf.drawString(40, A4[1] - 40, "Receipt 1 of 1")
    pdf.drawImage(photo_path, 60, 120, A4[0] - 120, (A4[0] - 120) * 4 / 3)
    pdf.showPage()
    pdf.save()


def with_cells(source, target, cells):
    """A copy of an openpyxl workbook whose formula cells carry the given XML, as a spreadsheet stores them."""
    with zipfile.ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = parts["xl/worksheets/sheet1.xml"].decode("utf-8")
    for old, new in cells.items():
        assert old in sheet, sheet
        sheet = sheet.replace(old, new)
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    with zipfile.ZipFile(target, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


class CheckTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.TemporaryDirectory()

    @classmethod
    def tearDownClass(cls):
        cls.dir.cleanup()

    def path(self, name):
        return os.path.join(self.dir.name, name)

    # --- .xlsx

    def formulas(self, name):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Итоги"
        sheet.append(["Статья", "Сумма"])
        sheet.append(["Аренда", 1500.5])
        sheet.append(["Связь", 200])
        sheet["B4"] = "=SUM(B2:B3)"
        sheet["B5"] = "=B4/0"
        book.save(self.path(name))
        return self.path(name)

    def test_formulas_openpyxl_never_computed(self):
        result = check(self.formulas("raw.xlsx"))
        self.assertEqual(result["kind"], "xlsx")
        (problem,) = result["problems"]
        self.assertEqual(problem["code"], "not_recalculated")
        self.assertIn("2 formula cells", problem["where"])
        self.assertIn("Итоги!B4", problem["where"])
        self.assertIn("brokit xlsx recalc", problem["fix"])

    def test_xlsxwriter_zeros_are_no_computed_values(self):
        import xlsxwriter

        book = xlsxwriter.Workbook(self.path("writer.xlsx"))
        sheet = book.add_worksheet()
        sheet.write_column("A1", [2, 3])
        sheet.write_formula("A3", "=A1*A2")
        book.close()
        self.assertEqual(codes(check(self.path("writer.xlsx"))), ["not_recalculated"])

    def test_a_computed_workbook_names_its_formula_errors_only(self):
        with_cells(
            self.formulas("plain.xlsx"),
            self.path("computed.xlsx"),
            {
                '<c r="B4"><f>SUM(B2:B3)</f><v></v></c>': '<c r="B4"><f>SUM(B2:B3)</f><v>1700.5</v></c>',
                '<c r="B5"><f>B4/0</f><v></v></c>': '<c r="B5" t="e"><f>B4/0</f><v>#DIV/0!</v></c>',
            },
        )
        (problem,) = check(self.path("computed.xlsx"))["problems"]
        self.assertEqual((problem["code"], problem["where"]), ("formula_errors", "Итоги!B5 #DIV/0!"))

    def test_an_empty_text_result_is_a_computed_value(self):
        book = openpyxl.Workbook()
        book.active["A1"] = '=IF(1=1,"","x")'
        book.save(self.path("blank.xlsx"))
        # LibreOffice stores an empty text result so.
        with_cells(
            self.path("blank.xlsx"), self.path("blank-computed.xlsx"), {'<c r="A1"><f>': '<c r="A1" t="str"><f>'}
        )
        self.assertEqual(check(self.path("blank-computed.xlsx"))["problems"], [])

    def test_dates_and_sums_typed_as_text(self):
        statement(self.path("statement.xlsx"))
        result = check(self.path("statement.xlsx"))
        self.assertEqual(codes(result), ["text_dates", "text_numbers"])
        self.assertIn("Операции!A (7 cells like 01.09.2026 08:01)", result["problems"][0]["where"])
        self.assertIn("'DD.MM.YYYY'", result["problems"][0]["fix"])

    def test_russian_date_codes(self):
        book = openpyxl.Workbook()
        book.active["A1"] = datetime.datetime(2026, 10, 16)
        book.active["A1"].number_format = "ДД.ММ.ГГГГ"
        book.save(self.path("codes.xlsx"))
        (problem,) = check(self.path("codes.xlsx"))["problems"]
        self.assertEqual((problem["code"], problem["where"]), ("local_date_codes", "ДД.ММ.ГГГГ"))

    def test_a_clean_workbook_has_no_problems(self):
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.append(["Дата", "Сумма", "ИНН"])
        for day in range(1, 9):
            sheet.append([datetime.datetime(2026, 9, day), 1234.5 * day, "77070838%02d" % day])
            sheet.cell(sheet.max_row, 1).number_format = "DD.MM.YYYY"
            sheet.cell(sheet.max_row, 2).number_format = "#,##0.00"
        book.save(self.path("clean.xlsx"))
        self.assertEqual(
            check(self.path("clean.xlsx")),
            {"ok": True, "file": self.path("clean.xlsx"), "kind": "xlsx", "problems": []},
        )

    # --- .pdf

    @unittest.skipUnless(shutil.which("pdfimages"), "poppler is not installed")
    def test_a_photo_of_a_page_on_a4_is_not_a_scan(self):
        photo(self.path("IMG_1.jpg"))
        code, result = kit(
            "pdf", "from-images", self.path("IMG_1.jpg"), self.path("IMG_1.jpg"), "-o", self.path("raw.pdf")
        )
        self.assertEqual(code, 0, result)
        (problem,) = check(self.path("raw.pdf"))["problems"]
        self.assertEqual((problem["code"], problem["where"]), ("photo_not_scanned", "pages 1-2"))
        self.assertIn("brokit img scan", problem["fix"])
        captioned_pdf(self.path("IMG_1.jpg"), self.path("captioned.pdf"))
        self.assertEqual([p["where"] for p in check(self.path("captioned.pdf"))["problems"]], ["page 1"])

    @unittest.skipUnless(shutil.which("pdfimages"), "poppler is not installed")
    def test_the_scan_of_that_photo_passes(self):
        photo(self.path("IMG_2.jpg"))
        code, result = kit("img", "scan", self.path("IMG_2.jpg"), "-o", self.path("scan.pdf"))
        self.assertEqual(code, 0, result)
        self.assertEqual(check(self.path("scan.pdf"))["problems"], [])

    @unittest.skipUnless(shutil.which("pdfimages"), "poppler is not installed")
    def test_a_picture_that_runs_off_its_frame_is_no_page(self):
        sky(self.path("houses.jpg"))
        code, result = kit("pdf", "from-images", self.path("houses.jpg"), "-o", self.path("houses.pdf"))
        self.assertEqual(code, 0, result)
        self.assertEqual(check(self.path("houses.pdf"))["problems"], [])

    @unittest.skipUnless(shutil.which("pdfimages"), "poppler is not installed")
    def test_a_text_pdf_and_a_protected_one(self):
        from reportlab.pdfgen import canvas

        pdf = canvas.Canvas(self.path("letter.pdf"))
        pdf.drawString(72, 720, "Справка с места работы")
        pdf.save()
        self.assertEqual(check(self.path("letter.pdf"))["problems"], [])
        code, result = kit(
            "pdf", "encrypt", self.path("letter.pdf"), "--password", "секрет", "-o", self.path("locked.pdf")
        )
        self.assertEqual(code, 0, result)
        result = check(self.path("locked.pdf"))
        self.assertEqual((result["problems"], result["skipped"]), ([], "a password protects it"))

    # --- .pptx

    def deck(self, name, words=20, chart=None):
        import pptx
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches

        deck = pptx.Presentation()
        for number in range(1, 4):
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            slide.shapes.title.text = "Итоги, часть %d" % number
            count = words if number == 2 else 6
            slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2)).text_frame.text = " ".join(
                ["слово"] * count
            )
        if chart == "picture" or chart == "both":
            deck.slides[2].shapes.add_picture(self.matplotlib_png(), Inches(1), Inches(4), Inches(4))
        if chart == "both":
            data = CategoryChartData()
            data.categories = ["I", "II"]
            data.add_series("Выручка", (3.0, 5.0))
            deck.slides[0].shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(4), Inches(3), data
            )
        if chart == "photo":
            photo(self.path("deck-photo.jpg"))
            deck.slides[2].shapes.add_picture(self.path("deck-photo.jpg"), Inches(1), Inches(4), Inches(3))
        deck.save(self.path(name))
        return self.path(name)

    def matplotlib_png(self):
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        figure, axes = plt.subplots()
        axes.bar(["I", "II"], [3, 5])
        figure.savefig(self.path("chart.png"))
        plt.close(figure)
        return self.path("chart.png")

    def test_a_slide_over_80_words(self):
        # The title's three words count too.
        (problem,) = check(self.deck("dense.pptx", words=109))["problems"]
        self.assertEqual((problem["code"], problem["where"]), ("dense_slides", "slide 2: 112 words"))
        self.assertIn("speaker notes", problem["fix"])
        self.assertEqual(check(self.deck("light.pptx", words=77))["problems"], [])

    def test_a_matplotlib_chart_in_a_deck_without_native_ones(self):
        result = check(self.deck("pictures.pptx", chart="picture"))
        self.assertEqual(result["problems"][0]["code"], "chart_pictures")
        self.assertEqual(result["problems"][0]["where"], "slide 3")
        self.assertIn("add_chart", result["problems"][0]["fix"])
        # A deck that has native charts may hold a picture of one on purpose; a photo is no chart.
        self.assertEqual(check(self.deck("both.pptx", chart="both"))["problems"], [])
        self.assertEqual(check(self.deck("photo.pptx", chart="photo"))["problems"], [])

    # --- the rest

    def test_other_files_have_no_problems(self):
        with open(self.path("notes.txt"), "w", encoding="utf-8") as handle:
            handle.write("Итоги квартала\n")
        self.assertEqual(
            check(self.path("notes.txt")), {"ok": True, "file": self.path("notes.txt"), "kind": "other", "problems": []}
        )
        code, result = kit_apart("check", self.path("missing.pdf"))
        self.assertEqual((code, result["code"]), (2, "not_found"))


if __name__ == "__main__":
    unittest.main()
