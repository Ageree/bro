"""brokit check FILE: what a finished file still needs before the person gets it.

The task agent's `share_file` runs it on every .xlsx, .pdf and .pptx and holds the file back
once when it names a problem (agent/subagents/task/tools/share_file.ts); the model may run it
too. Only what is certain counts: a false alarm costs the person a file they are waiting for.

- .xlsx: formula cells whose stored value is an error, formulas no spreadsheet computed
  (openpyxl and xlsxwriter cannot), both in one pass (`formula_values`); date or sum columns
  stored as text (`text_typed`); Russian date codes in number formats (`xlsx recalc` turns them
  into Excel's).
- .pdf: pages whose picture is a phone photo of a document as it was taken, not a scan made by
  `img scan` (`unscanned_photo`).
- .pptx: slides over 80 words (look's count), and charts drawn by matplotlib and put in as
  pictures, in a deck without one native chart.
- anything else: no problems.

Prints {"file", "kind", "problems": [{"code", "where", "fix"}]}; a PDF that a password
protects is not looked into (`skipped`).
"""

import math
import os
import tempfile
import zipfile

from brokit.look import SLIDE_WORDS, compact, crowded_slides
from brokit.render import is_pdf, pdf_info
from brokit.util import KitError, limits, need_file, parser, run

LISTED = 5

FIXES = {
    "formula_errors": "correct the formulas behind them (a division by an empty or zero cell, a reference to a "
    "missing cell or sheet), then run brokit xlsx recalc again",
    "not_recalculated": "brokit xlsx recalc FILE -o OUT computes every formula and names the ones that give an "
    "error; share OUT",
    "text_dates": "write datetime values into these cells with number_format 'DD.MM.YYYY' (or 'DD.MM.YYYY hh:mm'), "
    "not text",
    "text_numbers": "write numbers (float) into these cells with number_format '#,##0.00', not text",
    "local_date_codes": "number formats take English codes ('DD.MM.YYYY'): brokit xlsx recalc FILE -o OUT rewrites "
    "them",
    "photo_not_scanned": "a photo of a page goes into a PDF only through brokit img scan PHOTO… -o OUT.pdf: it finds "
    "the page, flattens it and whitens the background. Run it on the original photos, not on this PDF",
    "dense_slides": "a slide holds a heading and 3–5 short points, at most %d words: move the details to the speaker "
    "notes" % SLIDE_WORDS,
    "chart_pictures": "a chart goes on a slide as a native one: slide.shapes.add_chart(XL_CHART_TYPE…, x, y, cx, cy, "
    "CategoryChartData) with the numbers from the document, not a matplotlib picture",
}


def problem(code, where):
    return {"code": code, "where": where, "fix": FIXES[code]}


def listed(items, more):
    """The first LISTED items joined, and how many more there are."""
    text = ", ".join(items[:LISTED])
    return text + (" and %d more" % (more - LISTED) if more > LISTED else "")


# --- .xlsx ---------------------------------------------------------------------------------


def sheet_problems(path):
    from lxml import etree

    from brokit.office import sheet

    problems = []
    formulas = sheet.formula_values(path)
    if formulas["error_count"]:
        cells = ["%s!%s %s" % (error["sheet"], error["cell"], error["value"]) for error in formulas["errors"]]
        problems.append(problem("formula_errors", listed(cells, formulas["error_count"])))
    count = formulas["uncomputed"]
    if count:
        where = "%d formula %s without a computed value: %s" % (
            count,
            "cell" if count == 1 else "cells",
            listed(formulas["uncomputed_cells"], count),
        )
        problems.append(problem("not_recalculated", where))
    typed = sheet.text_typed(path)
    for code in ("text_dates", "text_numbers"):
        columns = [
            "%s!%s (%d cells like %s)" % (c["sheet"], c["column"], c["cells"], c["example"]) for c in typed[code]
        ]
        if columns:
            problems.append(problem(code, listed(columns, len(columns))))
    with zipfile.ZipFile(path) as archive:
        if "xl/styles.xml" in archive.namelist():
            root = etree.fromstring(
                archive.read("xl/styles.xml"), etree.XMLParser(huge_tree=True, resolve_entities=False)
            )
            codes = sorted({fmt.get("formatCode") or "" for fmt in sheet.local_date_formats(root)})
            if codes:
                problems.append(problem("local_date_codes", listed(codes, len(codes))))
    return problems


# --- .pdf ----------------------------------------------------------------------------------

PDF_PAGES = 60
PHOTO_PAGES = 30
# A picture counts as a page's content when it covers this share of the page: a photo «по
# центру A4» with a caption covers 0.45–0.7 of it.
PHOTO_SHARE = 0.3
SIDE = 600
GRAY = (0.299, 0.587, 0.114)
# Calibrated 16.10 on the bench's phone photos (receipts on a table, a palm, crumpled paper;
# contract pages on wood; a passport; an act), their `img scan` results, the visa pack scans,
# a screenshot of an e-receipt, a booklet's renders of houses under a sky, a selfie and a logo.
# The frame's outer 4 % is white in a scan (its paper whitened) and in a screenshot (0.89–1.0
# of the band), and is table, palm or floor in a photo (0–0.02).
BORDER_BAND = 0.04
WHITE_BORDER = 0.5
# The page found as `img scan` finds it does not reach the frame's edge in a photo of a page
# (0–0.01 of the edge), while a sky, a wall behind a face or a page cut by the frame does
# (0.2–0.62).
MAX_TOUCH = 0.05
# Paper is near neutral even under warm light (median max−min of RGB 3–49); a logo is not (92).
PAPER_CHROMA = 60
# Print: small dark marks inside the page (236–578 at SIDE px); a face, a white box, a
# signature alone have few (0–17).
MIN_MARKS = 60
# A page that fills the frame and stands straight is as good as flat.
MAX_PAGE_SHARE = 0.85
MIN_TILT = 3.0


def unscanned_photo(image):
    """Whether a picture is a phone photo of a page with the table around it, not flattened by `img scan`."""
    import numpy as np
    from PIL import Image, ImageFilter

    from brokit.img.raster import otsu
    from brokit.img.scan import page_quad, quad_area

    rgb = np.asarray(image.convert("RGB"), dtype=np.float32)
    gray = rgb @ np.array(GRAY, dtype=np.float32)
    chroma = rgb.max(axis=2) - rgb.min(axis=2)
    height, width = gray.shape
    if min(height, width) < 64:
        return False
    band = max(2, round(BORDER_BAND * min(height, width)))
    ring = np.ones((height, width), bool)
    ring[band:-band, band:-band] = False
    white = (gray >= 200) & (chroma <= 40)
    if white[ring].mean() >= WHITE_BORDER:
        return False
    blurred = np.asarray(Image.fromarray(gray.astype(np.uint8)).filter(ImageFilter.GaussianBlur(2)), dtype=np.float32)
    found = page_quad(blurred > otsu(blurred), blurred)
    if found is None:
        return False
    points, _, page = found
    if np.concatenate([page[0], page[-1], page[:, 0], page[:, -1]]).mean() > MAX_TOUCH:
        return False
    if float(np.median(chroma[page])) > PAPER_CHROMA or text_marks(gray, page) < MIN_MARKS:
        return False
    return quad_area(points) / (width * height) <= MAX_PAGE_SHARE or tilt(points) >= MIN_TILT


def text_marks(gray, page):
    """Small marks darker than the paper around them inside the page: letters, digits, lines of print."""
    import numpy as np
    from PIL import Image, ImageFilter

    from brokit.img.raster import border_connected, label_components

    solid = ~border_connected(~page, step=2)
    inner = np.asarray(Image.fromarray(solid.astype(np.uint8) * 255).filter(ImageFilter.MinFilter(9))) > 0
    local = np.asarray(Image.fromarray(gray.astype(np.uint8)).filter(ImageFilter.GaussianBlur(6)), dtype=np.float32)
    labels, count = label_components((gray < local - 20) & inner, 20 * MIN_MARKS)
    if labels is None:
        return count
    sizes = np.bincount(labels.ravel())[1:]
    return int(((sizes >= 3) & (sizes <= 0.01 * inner.sum())).sum())


def tilt(points):
    """The largest angle, in degrees, between an edge of the quadrilateral and the nearest axis."""
    angles = []
    for index in range(4):
        (x0, y0), (x1, y1) = points[index], points[(index + 1) % 4]
        angle = math.degrees(math.atan2(float(y1 - y0), float(x1 - x0))) % 90
        angles.append(min(angle, 90 - angle))
    return max(angles)


def page_pictures(path):
    """{page: index of its largest picture among the page's images} for pages that picture covers PHOTO_SHARE of."""
    total = min(pdf_info(path)["pages"], PDF_PAGES)
    sizes = pdf_info(path, 1, total)["sizes"]
    listing = run(["pdfimages", "-list", "-l", str(total), path], timeout=60)
    if listing.returncode != 0:
        return {}
    largest, seen = {}, {}
    for line in listing.stdout.decode("utf-8", "replace").splitlines()[2:]:
        fields = line.split()
        try:
            page, kind, width, height = int(fields[0]), fields[2], int(fields[3]), int(fields[4])
            x_ppi, y_ppi = float(fields[-4]), float(fields[-3])
        except (IndexError, ValueError):
            continue
        # Extraction numbers every image of the page, soft masks too, in this order.
        index = seen.get(page, 0)
        seen[page] = index + 1
        size = sizes.get(page)
        if kind != "image" or x_ppi <= 0 or y_ppi <= 0 or not size or size[0] * size[1] <= 0:
            continue
        share = (width / x_ppi * 72) * (height / y_ppi * 72) / (size[0] * size[1])
        if share >= PHOTO_SHARE and share > largest.get(page, (0, 0))[0]:
            largest[page] = (share, index)
    return {page: index for page, (_, index) in largest.items()}


def page_picture(path, page, index, work):
    """The page's picture at SIDE px on its long side, or None when it cannot be read."""
    from brokit.img.raster import load_photo

    prefix = os.path.join(work, "p%d" % page)
    try:
        if run(["pdfimages", "-f", str(page), "-l", str(page), "-j", path, prefix], timeout=60).returncode != 0:
            return None
        names = [name for name in os.listdir(work) if name.startswith("p%d-%03d." % (page, index))]
        return load_photo(os.path.join(work, names[0]), SIDE)[0] if names else None
    except KitError:
        return None
    finally:
        for name in os.listdir(work):
            os.remove(os.path.join(work, name))


def pdf_problems(path):
    pictures = page_pictures(path)
    photos = []
    with tempfile.TemporaryDirectory(prefix="brokit-check-") as work:
        for page in sorted(pictures)[:PHOTO_PAGES]:
            picture = page_picture(path, page, pictures[page], work)
            if picture is not None and unscanned_photo(picture):
                photos.append(page)
    if not photos:
        return []
    return [
        problem("photo_not_scanned", ("page " if len(photos) == 1 else "pages ") + compact(photos).replace(",", ", "))
    ]


# --- .pptx ---------------------------------------------------------------------------------


def drawn_by_matplotlib(shape):
    """Whether a picture shape holds a PNG matplotlib saved (it names itself in the PNG's Software text)."""
    import io

    from PIL import Image

    try:
        blob = shape.image.blob
        if not blob.startswith(b"\x89PNG"):
            return False
        with Image.open(io.BytesIO(blob)) as image:
            return "matplotlib" in str(image.info.get("Software", "")).lower()
    except Exception:  # noqa: BLE001 — a linked or broken picture is no chart
        return False


def pictures_of(shapes):
    from pptx.shapes.group import GroupShape
    from pptx.shapes.picture import Picture

    for shape in shapes:
        if isinstance(shape, GroupShape):
            yield from pictures_of(shape.shapes)
        elif isinstance(shape, Picture):
            yield shape


def deck_problems(path):
    problems = []
    crowded = crowded_slides(path)
    if crowded:
        problems.append(problem("dense_slides", "; ".join("slide %d: %d words" % pair for pair in crowded)))
    with zipfile.ZipFile(path) as archive:
        native = any(name.startswith("ppt/charts/chart") and name.endswith(".xml") for name in archive.namelist())
    if native or crowded is None:
        return problems
    import pptx

    slides = [
        number
        for number, slide in enumerate(pptx.Presentation(path).slides, 1)
        if any(drawn_by_matplotlib(shape) for shape in pictures_of(slide.shapes))
    ]
    if slides:
        problems.append(
            problem("chart_pictures", ("slide " if len(slides) == 1 else "slides ") + ", ".join(map(str, slides)))
        )
    return problems


# -------------------------------------------------------------------------------------------


def kind_of(path):
    if is_pdf(path):
        return "pdf"
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
        if "xl/workbook.xml" in names:
            return "xlsx"
        if "ppt/presentation.xml" in names:
            return "pptx"
    return "other"


def main(argv):
    p = parser("check")
    p.add_argument("file")
    args = p.parse_args(argv)
    limits()
    need_file(args.file)
    kind = kind_of(args.file)
    result = {"file": args.file, "kind": kind, "problems": []}
    if kind == "xlsx":
        result["problems"] = sheet_problems(args.file)
    elif kind == "pptx":
        result["problems"] = deck_problems(args.file)
    elif kind == "pdf":
        try:
            result["problems"] = pdf_problems(args.file)
        except KitError as error:
            if error.code != "encrypted":
                raise
            result["skipped"] = "a password protects it"
    return result
