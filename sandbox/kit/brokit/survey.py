"""brokit files survey: one look at every file of a folder (an unpacked archive).

For each file: the real type by its first bytes (not the extension), size,
sha256 with byte-identical duplicates grouped (the copy that looks like a copy
— «копия», «(1)», «Copy», a deeper folder — points to the original), junk
(Thumbs.db, desktop.ini, «~$» lock files, .DS_Store, __MACOSX), broken files
(an image that does not decode, a PDF poppler cannot read, a damaged OOXML),
the EXIF date of photos, the first 300 characters of text and whether it is a
scan (an image, or a PDF without a text layer).
"""

import hashlib
import os
import re
import tempfile

from brokit.util import KitError, parser, run

TEXT_CHARS = 300
MAX_FILES = 2000
JUNK_NAMES = {"thumbs.db", "desktop.ini", ".ds_store", "ehthumbs.db"}
COPY_MARK = re.compile(r"(?i)(копия|copy|\(\d+\)|\s-\s?копия|~\d+)(?=\.[^.]*$|$)")
OOXML_KINDS = (("word/", "docx"), ("xl/", "xlsx"), ("ppt/", "pptx"))
ODF_KINDS = {
    "application/vnd.oasis.opendocument.text": "odt",
    "application/vnd.oasis.opendocument.spreadsheet": "ods",
    "application/vnd.oasis.opendocument.presentation": "odp",
}
IMAGE_TYPES = {"jpeg", "png", "gif", "webp", "tiff", "bmp", "heic"}


def main(argv):
    p = parser("files survey")
    p.add_argument("directory")
    args = p.parse_args(argv)
    if not os.path.isdir(args.directory):
        raise KitError("not_found", "no such directory: " + args.directory)
    paths = []
    for folder, dirs, files in os.walk(args.directory):
        dirs.sort()
        for name in sorted(files):
            path = os.path.join(folder, name)
            if not os.path.islink(path):
                paths.append(path)
    if len(paths) > MAX_FILES:
        raise KitError("too_many", "more than %d files" % MAX_FILES, "survey one subfolder at a time")
    records = [describe(path, os.path.relpath(path, args.directory).replace(os.sep, "/")) for path in paths]
    mark_duplicates(records)
    add_office_text(records, [os.path.join(args.directory, r["path"]) for r in records])
    groups = {}
    for record in records:
        if record["dup_of"]:
            groups.setdefault(record["dup_of"], [record["dup_of"]]).append(record["path"])
    return {
        "files": records,
        "count": len(records),
        "duplicates": list(groups.values()),
        "junk": [r["path"] for r in records if r["junk"]],
        "broken": [r["path"] for r in records if r["broken"]],
    }


def sniff(path, head):
    """The file's type from its first bytes."""
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if head[:6] in (b"GIF87a", b"GIF89a"):
        return "gif"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "webp"
    if head[4:8] == b"ftyp" and head[8:12] in (b"heic", b"heix", b"mif1", b"msf1", b"heim", b"hevc"):
        return "heic"
    if head[4:8] == b"ftyp":
        return "mp4"
    if head[:4] in (b"II*\x00", b"MM\x00*"):
        return "tiff"
    if head[:2] == b"BM":
        return "bmp"
    if head[:4] == b"PK\x03\x04":
        return zip_kind(path)
    if head.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return ole_kind(path)
    if head.startswith(b"Rar!\x1a\x07"):
        return "rar"
    if head.startswith(b"7z\xbc\xaf\x27\x1c"):
        return "7z"
    if head.startswith(b"{\\rtf"):
        return "rtf"
    if head.startswith((b"ID3", b"\xff\xfb", b"\xff\xf3")):
        return "mp3"
    if head.startswith(b"OggS"):
        return "ogg"
    return text_kind(path, head)


def zip_kind(path):
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            if "mimetype" in names:
                return ODF_KINDS.get(archive.read("mimetype").decode("ascii", "replace").strip(), "zip")
            if "[Content_Types].xml" in names:
                for prefix, kind in OOXML_KINDS:
                    if any(name.startswith(prefix) for name in names):
                        return kind
    except (zipfile.BadZipFile, OSError, KeyError):
        return "zip"
    return "zip"


def ole_kind(path):
    """doc / xls / ppt / msg by the stream names in an OLE container's directory."""
    with open(path, "rb") as handle:
        data = handle.read(512 * 1024)
    for stream, kind in (
        ("WordDocument", "doc"),
        ("Workbook", "xls"),
        ("Book", "xls"),
        ("PowerPoint Document", "ppt"),
        ("__substg1.0_", "msg"),
    ):
        if stream.encode("utf-16-le") in data:
            return kind
    return "ole"


def decode_text(data):
    """Text of a small file: UTF-8/UTF-16 by BOM, UTF-8, else cp1251."""
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return data.decode("utf-16", "replace"), "utf-16"
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", "replace"), "utf-8"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError as error:
        if error.start > len(data) - 4:
            return data[: error.start].decode("utf-8"), "utf-8"
    return data.decode("cp1251", "replace"), "cp1251"


def text_kind(path, head):
    if b"\x00" in head and not (head.startswith(b"\xff\xfe") or head.startswith(b"\xfe\xff")):
        return "binary"
    text, _ = decode_text(head)
    stripped = text.lstrip().lower()
    name = path.lower()
    if stripped.startswith("begin:vcard"):
        return "vcf"
    if stripped.startswith("begin:vcalendar"):
        return "ics"
    if stripped.startswith("<?xml") or stripped.startswith("<"):
        return "html" if "<html" in stripped[:500] else "xml"
    if name.endswith(".csv"):
        return "csv"
    if name.endswith(".ini") or stripped.startswith("["):
        return "ini"
    return "txt"


def is_junk(relative):
    name = relative.rsplit("/", 1)[-1]
    lower = name.lower()
    return (
        lower in JUNK_NAMES
        or name.startswith("~$")
        or name.startswith("._")
        or lower.endswith(".tmp")
        or "__MACOSX/" in relative + "/"
    )


def describe(path, relative):
    size = os.path.getsize(path)
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        head = handle.read(64 * 1024)
        digest.update(head)
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    kind = sniff(path, head)
    record = {
        "path": relative,
        "type": kind,
        "bytes": size,
        "sha256": digest.hexdigest(),
        "dup_of": None,
        "junk": is_junk(relative),
        "broken": None,
        "exif_date": None,
        "text": None,
        "scan": False,
        "pages": None,
    }
    if size == 0:
        record["broken"] = "empty file"
        return record
    if record["junk"]:
        return record
    if kind in IMAGE_TYPES:
        inspect_image(path, kind, record)
    elif kind == "pdf":
        inspect_pdf(path, record)
    elif kind in ("docx", "xlsx", "pptx"):
        inspect_ooxml(path, kind, record)
    elif kind in ("txt", "csv", "xml", "html", "vcf", "ics", "ini", "rtf"):
        with open(path, "rb") as handle:
            text, _ = decode_text(handle.read(16 * 1024))
        record["text"] = short(text)
    return record


def short(text):
    text = re.sub(r"[ \t ]+", " ", text or "")
    text = re.sub(r"\s*\n\s*", "\n", text).strip()
    return text[:TEXT_CHARS] or None


def inspect_image(path, kind, record):
    from PIL import Image, UnidentifiedImageError

    record["scan"] = True
    if kind == "heic":
        record["broken"] = None
        return
    try:
        with Image.open(path) as image:
            exif = image.getexif()
            date = exif.get_ifd(0x8769).get(36867) or exif.get(306)
            if date:
                record["exif_date"] = str(date).strip().replace(":", "-", 2)
            image.load()
    except (UnidentifiedImageError, OSError, SyntaxError, ValueError) as error:
        record["broken"] = "image does not decode: " + str(error)[:80]
        return
    if kind == "jpeg":
        with open(path, "rb") as handle:
            handle.seek(max(0, record["bytes"] - 1024))
            if b"\xff\xd9" not in handle.read():
                record["broken"] = "JPEG is cut short"


def inspect_pdf(path, record):
    info = run(["pdfinfo", path], timeout=60)
    if info.returncode != 0:
        message = info.stderr.decode("utf-8", "replace").strip().splitlines()
        record["broken"] = "PDF does not open: " + (message[-1][:80] if message else "pdfinfo failed")
        return
    match = re.search(rb"^Pages:\s+(\d+)", info.stdout, re.M)
    record["pages"] = int(match.group(1)) if match else None
    if re.search(rb"^Encrypted:\s+yes", info.stdout, re.M):
        record["encrypted"] = True
    text = run(["pdftotext", "-l", "3", "-enc", "UTF-8", path, "-"], timeout=60)
    content = text.stdout.decode("utf-8", "replace") if text.returncode == 0 else ""
    record["text"] = short(content)
    pages = min(record["pages"] or 1, 3)
    record["scan"] = len(re.sub(r"\s+", "", content)) < 20 * pages


def inspect_ooxml(path, kind, record):
    import zipfile

    try:
        with zipfile.ZipFile(path) as archive:
            damaged = archive.testzip()
    except (zipfile.BadZipFile, OSError) as error:
        record["broken"] = "damaged archive: " + str(error)[:80]
        return
    if damaged:
        record["broken"] = "damaged part: " + damaged
        return
    try:
        if kind == "docx":
            import docx

            document = docx.Document(path)
            record["text"] = short("\n".join(p.text for p in document.paragraphs if p.text.strip()))
        elif kind == "xlsx":
            record["text"] = short(xlsx_preview(path))
        else:
            import pptx

            deck = pptx.Presentation(path)
            record["pages"] = len(deck.slides)
            lines = []
            for slide in list(deck.slides)[:3]:
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame.text.strip():
                        lines.append(shape.text_frame.text)
            record["text"] = short("\n".join(lines))
    except Exception as error:  # noqa: BLE001 — a library refusing the file is a finding, not a crash
        record["broken"] = "%s does not open: %s" % (kind, str(error)[:80])


def xlsx_preview(path):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        lines = []
        for sheet in workbook.worksheets[:2]:
            lines.append("[" + sheet.title + "]")
            for row in sheet.iter_rows(max_row=8, values_only=True):
                cells = [str(value) for value in row if value not in (None, "")]
                if cells:
                    lines.append(" | ".join(cells))
        return "\n".join(lines)
    finally:
        workbook.close()


def original_rank(record):
    name = record["path"].rsplit("/", 1)[-1]
    return (bool(COPY_MARK.search(name)), record["path"].count("/"), len(name), record["path"])


def mark_duplicates(records):
    by_hash = {}
    for record in records:
        if record["bytes"] and not record["junk"]:
            by_hash.setdefault(record["sha256"], []).append(record)
    for group in by_hash.values():
        if len(group) < 2:
            continue
        group.sort(key=original_rank)
        for copy in group[1:]:
            copy["dup_of"] = group[0]["path"]


def add_office_text(records, paths):
    """Text of old .doc/.xls through LibreOffice, one batch per format."""
    from brokit.office import soffice

    jobs = {"doc": "txt:Text (encoded):UTF8", "xls": "csv:Text - txt - csv (StarCalc):44,34,76,1"}
    for kind, target in jobs.items():
        chosen = [(r, p) for r, p in zip(records, paths) if r["type"] == kind and not r["junk"] and not r["dup_of"]]
        if not chosen:
            continue
        with tempfile.TemporaryDirectory(prefix="brokit-survey-") as work:
            sources = []
            for n, (_, path) in enumerate(chosen):
                source = os.path.join(work, "f%d.%s" % (n, kind))
                os.symlink(os.path.abspath(path), source)
                sources.append(source)
            try:
                outputs = soffice.convert(sources, target, os.path.join(work, "out"))
            except KitError as error:
                if error.code == "unreadable":
                    for record, _ in chosen:
                        record["broken"] = record["broken"] or "LibreOffice: " + error.error
                continue
            for (record, _), output in zip(chosen, outputs):
                if output is None:
                    record["broken"] = "LibreOffice cannot open it"
                    continue
                with open(output, "rb") as handle:
                    record["text"] = short(decode_text(handle.read(16 * 1024))[0])
