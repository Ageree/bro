"""brokit look: files as pictures for the task agent's `look` tool (W2 spec §6.1).

`look FILE… --out DIR [--pages S] [--grid] [--region x0,y0,x1,y1]` writes JPEG
pictures of PDF pages, images (frames of a TIFF or GIF are pages) and office
files (converted to PDF once per content, `brokit.render.office_pdf`) into
DIR, and prints what each shows. The tool reads them and moves the latest two
looks into the model's prompt (agent/lib/model/tool-images.ts).

- A page is drawn from its CropBox after /Rotate: its units are view points,
  from the top-left of the page as displayed, like every kit command; an
  image's units are pixels after its EXIF rotation.
- Budget: a picture's long side is at most 1280 px and its JPEG at most
  400 KB (quality 60 down to 35, then smaller), at most four pictures and
  1.2 MB in one call; what did not fit is `truncated`, and `next` names the
  pages to ask for.
- `--grid` draws coordinates: magenta lines (35 %) every step, the step
  nearest to a sixteenth of the long side in {1, 2, 5, 10, …, 1000}, ticks at
  a fifth of it, labels along the four edges only.
- `--region` zooms into one page of one file: up to 600 dpi for a PDF, up to
  four times for an image, and the grid's step comes from the region.
"""

import io
import math
import os
import re
import shutil
import tempfile

from brokit.render import extension, is_office, is_pdf, office_pdf, pdf_info, render_page, view_size
from brokit.util import KitError, limits, need_file, parser, run

LONG_SIDE = 1280
MAX_PICTURES = 4
PICTURE_BYTES = 400_000
CALL_BYTES = 1_200_000
# Less than this left of the call's budget is not worth another picture.
MIN_PICTURE_BYTES = 40_000
# A picture is not made smaller than this to fit its budget.
MIN_SIDE = 320
GRID_STEPS = (1, 2, 5, 10, 20, 25, 50, 100, 200, 250, 500, 1000)
REGION_DPI = 600
REGION_ZOOM = 4
HEIC_BRANDS = (b"heic", b"heix", b"hevc", b"hevx", b"heim", b"heis", b"mif1", b"msf1")
# The formats `look` opens. Pillow knows many more, some through outside
# programs (EPS runs Ghostscript): a hostile file is opened as one of these
# or refused. JPEG includes MPO, the JPEG of many phone cameras.
IMAGE_FORMATS = ("AVIF", "BMP", "GIF", "JPEG", "PNG", "TIFF", "WEBP")
FONT_PATHS = ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "DejaVuSans.ttf")
# A picture of a slide does not say how much text it carries: a deck of
# 95–164 words a slide at 12 pt looked fine to the task agent.
SLIDE_WORDS = 80


def grid_step(long_side):
    """The step of the grid: the value nearest to a sixteenth of the long side."""
    return min(GRID_STEPS, key=lambda step: (abs(step - long_side / 16), step))


def parse_region(text):
    try:
        values = [float(part) for part in text.split(",")]
    except ValueError:
        values = []
    if len(values) != 4 or values[2] <= values[0] or values[3] <= values[1]:
        raise KitError("bad_region", "bad region: " + text, "x0,y0,x1,y1 with x1 > x0 and y1 > y0, in grid units")
    return values


def wanted_pages(spec, total, warnings, path):
    """Pages of "1-3,5" or "4-" that the file has; a page it lacks is a warning, not a refusal."""
    if spec is None or total <= 1:
        return list(range(1, total + 1))
    pages = []
    for part in spec.split(","):
        match = re.fullmatch(r"\s*(\d+)\s*(?:(-)\s*(\d*))?\s*", part)
        if match is None:
            raise KitError("bad_pages", "bad page spec: " + spec, 'like "1-3,5" or "4-"')
        start = int(match.group(1))
        end = start if match.group(2) is None else int(match.group(3) or total)
        if start < 1 or end < start:
            raise KitError("bad_pages", "bad page spec: " + spec, 'like "1-3,5" or "4-"')
        if start > total:
            warnings.append("%s has %d pages: no page %d" % (path, total, start))
            continue
        pages.extend(page for page in range(start, min(end, total) + 1) if page not in pages)
    return pages


def compact(pages):
    """[5, 6, 7, 9] → "5-7,9"."""
    runs = []
    for page in pages:
        if runs and page == runs[-1][1] + 1:
            runs[-1][1] = page
        else:
            runs.append([page, page])
    return ",".join(str(a) if a == b else "%d-%d" % (a, b) for a, b in runs)


def open_image(path):
    from PIL import Image

    return Image.open(path, formats=IMAGE_FORMATS)


def classify(path):
    """pdf, office or image, from the bytes first; refuses anything else."""
    need_file(path)
    if is_pdf(path):
        return "pdf"
    if is_office(path):
        return "office"
    with open(path, "rb") as handle:
        head = handle.read(32)
    try:
        with open_image(path):
            return "image"
    except Exception:  # noqa: BLE001 — Pillow raises many kinds for a file it cannot read
        if head[4:8] == b"ftyp" and head[8:12] in HEIC_BRANDS:
            raise KitError("unsupported", path + " is a HEIC photo, which cannot be opened here", "ask for it as JPEG (пришли JPEG)")
        raise KitError("unsupported", path + " is not a PDF, an image or an office file (." + extension(path) + ")")


def flattened(image):
    """RGB on white, whatever the mode."""
    from PIL import Image

    if image.mode in ("RGBA", "LA", "PA") or (image.mode == "P" and "transparency" in image.info):
        rgba = image.convert("RGBA")
        canvas = Image.new("RGB", rgba.size, "white")
        canvas.paste(rgba, mask=rgba.getchannel("A"))
        return canvas
    return image.convert("RGB")


class PdfPages:
    """A PDF, or an office file's PDF, page by page; units are view points."""

    unit = "pt"

    def __init__(self, path, kind, warnings, work):
        self.kind = kind
        self.pdf = office_pdf(path) if kind == "office" else path
        info = pdf_info(self.pdf)
        self.total = info["pages"]
        self.work = work
        if info["encrypted"]:
            warnings.append(path + " is encrypted without a user password: opened as is")

    def view(self, page):
        size = pdf_info(self.pdf, page, page)["sizes"].get(page)
        if size is None:
            raise KitError("unreadable", "page %d has no size" % page)
        return view_size(size)

    def picture(self, page, view, region):
        from PIL import Image

        prefix = os.path.join(self.work, "page-%d" % page)
        if region is None:
            path = render_page(self.pdf, page, prefix, scale_to=LONG_SIDE)
        else:
            width, height = region[2] - region[0], region[3] - region[1]
            dpi = min(REGION_DPI, LONG_SIDE * 72.0 / max(width, height))
            pixels = dpi / 72.0
            crop = (
                int(round(region[0] * pixels)),
                int(round(region[1] * pixels)),
                max(1, int(round(width * pixels))),
                max(1, int(round(height * pixels))),
            )
            path = render_page(self.pdf, page, prefix, dpi=dpi, crop=crop)
        with Image.open(path) as image:
            picture = image.convert("RGB")
        os.remove(path)
        return picture

    def text_chars(self, page):
        result = run(["pdftotext", "-q", "-cropbox", "-f", str(page), "-l", str(page), self.pdf, "-"], timeout=60)
        if result.returncode != 0:
            return 0
        return len(re.sub(r"\s+", "", result.stdout.decode("utf-8", "replace")))


class ImageFrames:
    """An image, frame by frame; units are pixels after the EXIF rotation."""

    unit = "px"
    kind = "image"

    def __init__(self, path):
        self.path = path
        try:
            with open_image(path) as image:
                # An MPO's later frames are the camera's depth map or preview, not pages.
                self.total = 1 if image.format == "MPO" else getattr(image, "n_frames", 1)
        except Exception:  # noqa: BLE001
            raise KitError("unreadable", path + " is not a readable image")

    def _frame(self, page, draft):
        from PIL import Image, ImageOps

        try:
            image = open_image(self.path)
            image.seek(page - 1)
            if draft and image.format == "JPEG":
                image.draft("RGB", (LONG_SIDE, LONG_SIDE))
            frame = ImageOps.exif_transpose(image)
            frame.load()
            return image, frame
        except (OSError, ValueError, EOFError, Image.DecompressionBombError):
            raise KitError("unreadable", "%s: frame %d cannot be read" % (self.path, page))

    def view(self, page):
        try:
            with open_image(self.path) as image:
                image.seek(page - 1)
                orientation = image.getexif().get(0x0112, 1)
                width, height = image.size
        except (OSError, ValueError, EOFError):
            raise KitError("unreadable", "%s: frame %d cannot be read" % (self.path, page))
        return (height, width) if orientation in (5, 6, 7, 8) else (width, height)

    def picture(self, page, view, region):
        from PIL import Image

        source, frame = self._frame(page, draft=region is None)
        try:
            if region is None:
                picture = flattened(frame)
                picture.thumbnail((LONG_SIDE, LONG_SIDE), Image.LANCZOS)
                return picture
            # The frame may have been drafted smaller than the view: map the region onto it.
            ratio = frame.width / float(view[0])
            box = tuple(int(round(value * ratio)) for value in region)
            crop = flattened(frame.crop(box))
            long_side = max(crop.size)
            target = min(LONG_SIDE, REGION_ZOOM * long_side) if long_side < LONG_SIDE else LONG_SIDE
            factor = target / float(long_side)
            if abs(factor - 1) > 0.01:
                crop = crop.resize((max(1, round(crop.width * factor)), max(1, round(crop.height * factor))), Image.LANCZOS)
            return crop
        finally:
            source.close()

    def text_chars(self, page):
        return None


def word_count(text):
    """Words of `text`: runs between spaces with a letter or a digit, so bullets and dashes are not words."""
    return sum(1 for token in text.split() if any(char.isalnum() for char in token))


def shape_words(shapes):
    from pptx.shapes.group import GroupShape

    total = 0
    for shape in shapes:
        if isinstance(shape, GroupShape):
            total += shape_words(shape.shapes)
        elif shape.has_text_frame:
            total += word_count(shape.text_frame.text)
        elif getattr(shape, "has_table", False):
            # A merged cell's spanned parts would count its text again.
            total += sum(word_count(cell.text_frame.text) for row in shape.table.rows for cell in row.cells if not cell.is_spanned)
    return total


def crowded_slides(path):
    """[(slide number, words)] of a .pptx's slides over SLIDE_WORDS words (charts and notes are not counted), or None
    for a deck python-pptx cannot read."""
    try:
        import pptx

        counts = [shape_words(slide.shapes) for slide in pptx.Presentation(path).slides]
    except Exception:  # noqa: BLE001 — a deck python-pptx refuses is still drawn by LibreOffice
        return None
    return [(number, words) for number, words in enumerate(counts, 1) if words > SLIDE_WORDS]


def dense_slides(path, warnings):
    """One warning for a .pptx whose slides carry more than SLIDE_WORDS words."""
    from brokit.ru.amount import plural

    dense = ["слайд %d: %d %s" % (number, words, plural(words, ("слово", "слова", "слов"))) for number, words in crowded_slides(path) or []]
    if dense:
        warnings.append(
            "%s has too much text on a slide: %s. A slide holds a heading and 3–5 short points (≤ %d words, body text ≥ 16 pt); "
            "move the rest to the speaker notes." % (path, "; ".join(dense), SLIDE_WORDS)
        )


def font(size):
    from PIL import ImageFont

    for path in FONT_PATHS:
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default(size)


def label_text(value):
    return str(int(value)) if float(value).is_integer() else ("%.2f" % value).rstrip("0").rstrip(".")


def draw_grid(picture, box, step):
    """Lines every `step` units over `box` (x0, y0, x1, y1 the picture shows), labels on the edges."""
    from PIL import Image, ImageDraw

    x0, y0, x1, y1 = box
    width, height = picture.size
    sx = width / float(x1 - x0)
    sy = height / float(y1 - y0)
    overlay = Image.new("RGBA", picture.size, (0, 0, 0, 0))
    draw = ImageDraw.Draw(overlay)
    colour = (255, 0, 255, 89)
    minor = step / 5.0
    tick = 6
    xs = [k * minor for k in range(int(math.ceil(x0 / minor - 1e-9)), int(math.floor(x1 / minor + 1e-9)) + 1)]
    ys = [k * minor for k in range(int(math.ceil(y0 / minor - 1e-9)), int(math.floor(y1 / minor + 1e-9)) + 1)]
    majors_x = []
    majors_y = []
    for value in xs:
        px = int(round((value - x0) * sx))
        if abs(value / step - round(value / step)) < 1e-6:
            draw.line([(px, 0), (px, height)], fill=colour, width=1)
            majors_x.append((value, px))
        else:
            draw.line([(px, 0), (px, tick)], fill=colour, width=1)
            draw.line([(px, height - tick), (px, height)], fill=colour, width=1)
    for value in ys:
        py = int(round((value - y0) * sy))
        if abs(value / step - round(value / step)) < 1e-6:
            draw.line([(0, py), (width, py)], fill=colour, width=1)
            majors_y.append((value, py))
        else:
            draw.line([(0, py), (tick, py)], fill=colour, width=1)
            draw.line([(width - tick, py), (width, py)], fill=colour, width=1)
    picture = Image.alpha_composite(picture.convert("RGBA"), overlay).convert("RGB")

    size = max(14, int(round(max(width, height) / 80.0)))
    face = font(size)
    draw = ImageDraw.Draw(picture)
    ink = (150, 0, 150)
    halo = {"fill": ink, "font": face, "stroke_width": 2, "stroke_fill": (255, 255, 255)}
    band = size + 6
    right_edge = -1
    for value, px in majors_x:
        text = label_text(value)
        left = px + 3
        text_width = draw.textlength(text, font=face)
        if left < right_edge + 4 or left + text_width > width - 2:
            continue
        draw.text((left, 2), text, **halo)
        draw.text((left, height - size - 4), text, **halo)
        right_edge = left + text_width
    bottom_edge = -1
    for value, py in majors_y:
        top = py + 2
        if top < band or top + size > height - band or top < bottom_edge + 2:
            continue
        text = label_text(value)
        text_width = draw.textlength(text, font=face)
        draw.text((3, top), text, **halo)
        draw.text((width - text_width - 3, top), text, **halo)
        bottom_edge = top + size
    return picture


def encoded(base, box, step, cap):
    """The JPEG of `base` (with its grid) within `cap` bytes, or None: quality first, then size.

    The picture as it is always gets a try, however small (an icon, a small
    region); it is made smaller only while it stays MIN_SIDE px long.
    """
    from PIL import Image

    factor = 1.0
    while True:
        if factor == 1.0:
            picture = base
        else:
            size = (max(1, int(round(base.width * factor))), max(1, int(round(base.height * factor))))
            picture = base.resize(size, Image.LANCZOS)
        if step is not None:
            picture = draw_grid(picture, box, step)
        quality = 60
        while True:
            buffer = io.BytesIO()
            picture.save(buffer, "JPEG", quality=quality, optimize=True)
            if buffer.tell() <= cap:
                return buffer.getvalue(), picture.size
            if quality <= 35:
                break
            quality = max(35, int(quality * 0.85))
        factor *= 0.85
        if max(base.size) * factor < MIN_SIDE:
            return None


def clamp_region(region, view, warnings):
    width, height = view
    clamped = [max(0.0, min(region[0], width)), max(0.0, min(region[1], height)), max(0.0, min(region[2], width)), max(0.0, min(region[3], height))]
    if clamped[2] - clamped[0] < 1 or clamped[3] - clamped[1] < 1:
        raise KitError("bad_region", "the region lies outside the page (%g x %g)" % (width, height))
    if clamped != list(region):
        warnings.append("the region was cut to the page (%g x %g)" % (width, height))
    return clamped


def main(argv):
    p = parser("look")
    p.add_argument("files", nargs="+")
    p.add_argument("--out", required=True)
    p.add_argument("--pages")
    p.add_argument("--grid", action="store_true")
    p.add_argument("--region")
    args = p.parse_args(argv)
    limits()
    region = parse_region(args.region) if args.region else None
    if region is not None and len(args.files) != 1:
        raise KitError("bad_args", "--region takes one file and one page")
    kinds = [classify(path) for path in args.files]
    os.makedirs(args.out, exist_ok=True)
    work = tempfile.mkdtemp(prefix="brokit-look-")
    warnings = []
    items = []
    following = None
    pictures_left = MAX_PICTURES
    used = 0
    try:
        for index, (path, kind) in enumerate(zip(args.files, kinds)):
            unit = "px" if kind == "image" else "pt"
            item = {"path": path, "kind": kind, "unit": unit, "pages_total": None, "images": []}
            items.append(item)
            if following is not None:
                warnings.append("not shown, over this call's pictures: " + path)
                continue
            source = ImageFrames(path) if kind == "image" else PdfPages(path, kind, warnings, work)
            if extension(path) == "pptx":
                dense_slides(path, warnings)
            item["pages_total"] = source.total
            pages = wanted_pages(args.pages, source.total, warnings, path)
            if region is not None and args.pages is None:
                pages = pages[:1]
            if region is not None and len(pages) != 1:
                raise KitError("bad_args", "--region takes one file and one page", 'add --pages N')
            for position, page in enumerate(pages):
                cap = min(PICTURE_BYTES, CALL_BYTES - used)
                if pictures_left == 0 or cap < MIN_PICTURE_BYTES:
                    following = {"path": path, "pages": compact(pages[position:])}
                    break
                view = source.view(page)
                shown = clamp_region(region, view, warnings) if region is not None else None
                box = shown or [0.0, 0.0, float(view[0]), float(view[1])]
                step = grid_step(max(box[2] - box[0], box[3] - box[1])) if args.grid else None
                result = encoded(source.picture(page, view, shown), box, step, cap)
                if result is None:
                    following = {"path": path, "pages": compact(pages[position:])}
                    break
                data, (width, height) = result
                file = os.path.join(args.out, "%d-p%d.jpg" % (index + 1, page))
                with open(file, "wb") as handle:
                    handle.write(data)
                used += len(data)
                pictures_left -= 1
                item["images"].append(
                    {
                        "page": page,
                        "file": file,
                        "width": width,
                        "height": height,
                        "bytes": len(data),
                        "view": {"width": round(view[0], 2), "height": round(view[1], 2)},
                        "scale": round(width / float(box[2] - box[0]), 4),
                        "region": [round(value, 2) for value in shown] if shown else None,
                        "grid_step": step,
                        "text_chars": source.text_chars(page),
                    }
                )
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return {"items": items, "truncated": following is not None, "next": following, "warnings": warnings}
