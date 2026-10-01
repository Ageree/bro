"""Smoke test of the sandbox rootfs, run as the sandbox user inside it.

build_rootfs.sh runs it in the chroot before packing (it also warms the matplotlib font cache and the
LibreOffice profile under $HOME, which then ship in the image), and scripts/cloudru-code-host/e2e.py runs it
in a live sandbox. Exits non-zero on the first failed check; prints one JSON line of what it saw.

  python3 verify.py [WORKDIR]
"""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import time

TEXT = "Привет, Бро: отчёт за сентябрь"
checks = {}


def timed(name):
    def wrap(fn):
        started = time.monotonic()
        result = fn()
        checks[name] = {"ms": round((time.monotonic() - started) * 1000), **(result or {})}
        return result
    return wrap


def main():
    work = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="verify-")
    os.makedirs(work, exist_ok=True)
    os.chdir(work)

    @timed("identity")
    def _():
        assert os.getuid() == 1000 and os.getgid() == 1000, (os.getuid(), os.getgid())
        assert os.path.realpath(sys.prefix) == "/opt/py", sys.prefix
        assert shutil.which("python3") == "/opt/py/bin/python3", shutil.which("python3")
        assert os.access(os.path.expanduser("~"), os.W_OK), "HOME is not writable"
        return {"python": sys.version.split()[0], "tz": time.strftime("%Z")}

    @timed("chart")
    def _():
        import matplotlib
        import matplotlib.pyplot as plt
        assert matplotlib.get_backend().lower() == "agg", matplotlib.get_backend()
        fig, ax = plt.subplots(figsize=(4, 3))
        ax.bar(["янв", "фев", "мар"], [3, 5, 2])
        ax.set_title(TEXT)
        fig.savefig("chart.png", dpi=80)
        return {"bytes": os.path.getsize("chart.png")}

    @timed("pptx")
    def _():
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = TEXT
        slide.placeholders[1].text = "Ёжик в тумане — 42 %"
        picture = deck.slides.add_slide(deck.slide_layouts[5])
        picture.shapes.title.text = "График"
        picture.shapes.add_picture("chart.png", Inches(1), Inches(1.5), width=Inches(6))
        deck.save("deck.pptx")
        return {"bytes": os.path.getsize("deck.pptx")}

    @timed("xlsx")
    def _():
        import openpyxl
        import pandas as pd
        import xlsxwriter
        book = openpyxl.Workbook()
        book.active.append(["Месяц", "Сумма"])
        book.active.append(["Сентябрь", 1500])
        book.save("openpyxl.xlsx")
        writer = xlsxwriter.Workbook("xlsxwriter.xlsx")
        sheet = writer.add_worksheet("Итоги")
        sheet.write_row(0, 0, ["Месяц", "Сумма"])
        sheet.write_row(1, 0, ["Октябрь", 2500])
        writer.close()
        frame = pd.concat([pd.read_excel("openpyxl.xlsx"), pd.read_excel("xlsxwriter.xlsx")])
        assert frame["Сумма"].sum() == 4000, frame
        return {"rows": len(frame)}

    @timed("docx")
    def _():
        import docx
        document = docx.Document()
        document.add_heading(TEXT, level=1)
        document.add_paragraph("Съешь же ещё этих мягких французских булок.")
        document.save("note.docx")
        assert docx.Document("note.docx").paragraphs[0].text == TEXT
        return {"bytes": os.path.getsize("note.docx")}

    @timed("pdf_tools")
    def _():
        import bs4
        import numpy
        import pypdf
        import reportlab
        import tabulate
        from PIL import Image
        from reportlab.pdfgen import canvas
        Image.open("chart.png").verify()
        pdf = canvas.Canvas("reportlab.pdf")
        pdf.drawString(72, 720, "reportlab")
        pdf.save()
        assert len(pypdf.PdfReader("reportlab.pdf").pages) == 1
        assert bs4.BeautifulSoup("<p>да</p>", "lxml").p.text == "да"
        assert "|" in tabulate.tabulate([[1, 2]], tablefmt="github")
        return {"numpy": numpy.__version__, "reportlab": reportlab.Version}

    @timed("soffice")
    def _():
        result = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", ".", "deck.pptx"],
                                capture_output=True, text=True, timeout=300)
        assert result.returncode == 0 and os.path.exists("deck.pdf"), (result.returncode, result.stdout,
                                                                       result.stderr)
        import pypdf
        pages = len(pypdf.PdfReader("deck.pdf").pages)
        assert pages >= 1, pages
        text = subprocess.run(["pdftotext", "deck.pdf", "-"], capture_output=True, text=True, check=True).stdout
        assert "Привет" in text, text[:200]
        return {"pages": pages, "bytes": os.path.getsize("deck.pdf")}

    print(json.dumps({"ok": True, "checks": checks}, ensure_ascii=False))


if __name__ == "__main__":
    main()
