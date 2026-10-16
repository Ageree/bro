"""Fixture for f-mailmerge-gramoty-dative.

gramota_shablon.docx  A4 landscape template. The decorative frame is a full-page PNG (29.7×21 cm, 150 dpi) anchored
                      in the header behind the text (wp:anchor behindDoc=1, relative to the page), so it repeats on
                      every page of a merged document; «ОБРАЗЕЦ / SPECIMEN» is printed in the frame's bottom-right
                      corner. Body: «ГРАМОТА», «вручается {ученику/ученице} 5 «Б» класса», «{ФИО}», «за {номинация}»,
                      «Классный руководитель ____ И. И. Образцова», «24 декабря 2026 г.». Each placeholder is a run of
                      its own, formatted like the rest of its line.
klass_5B.xlsx         title row (merged) + header «Фамилия Имя | Номинация» + 25 pupils from datagen.GRAMOTY in the
                      nominative, sorted alphabetically as in a class register; gender is NOT in the table.
                      Nominations are written to follow «за …» (accusative), without the «за».

truth/gramoty_5B.json  rows in sheet order with gender, dative form and the expected certificate lines.
"""

from __future__ import annotations

import io
import math
import pathlib
import sys

import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Emu, Pt, RGBColor
from PIL import Image, ImageDraw

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))
import common as C  # noqa: E402
import datagen as DG  # noqa: E402

CASE = "f-mailmerge-gramoty-dative"
DOCX_NAME = "gramota_shablon.docx"
XLSX_NAME = "klass_5B.xlsx"
WHEN = (2026, 12, 15, 18, 0, 0)
PAGE_CM = (29.7, 21.0)
FRAME_DPI = 150

NOMINATIONS = {
    "Иванов Пётр": "активное участие в олимпиаде по математике",
    "Смирнова Анна": "отличную учёбу",
    "Черных Илья": "победу в школьном турнире по шахматам",
    "Черных Дарья": "успехи в изучении английского языка",
    "Ким Алина": "первое место в конкурсе чтецов",
    "Ким Тимур": "активное участие в спортивной жизни класса",
    "Грицук Мария": "творческий подход к оформлению классного уголка",
    "Шевчук Никита": "победу в районной олимпиаде по информатике",
    "Кузьмина Любовь": "успехи в учёбе и примерное поведение",
    "Лебедь Павел": "помощь одноклассникам в учёбе",
    "Гурамишвили Нино": "призовое место в конкурсе рисунков „Зимняя сказка“",
    "Соловей Лев": "лучший проект на школьной научной конференции",
    "Белых Ксения": "активное участие в жизни класса",
    "Ткаченко Артём": "высокие результаты в школьной спартакиаде",
    "Ткаченко Илья": "волю к победе в соревнованиях по плаванию",
    "Рябой Михаил": "отличное знание истории родного края",
    "Толстая Варвара": "победу в конкурсе юных музыкантов",
    "Кох Ева": "успехи в изучении русского языка и литературы",
    "Никитин Савва": "изобретательность на занятиях по робототехнике",
    "Пушкарь Глеб": "ответственное отношение к обязанностям старосты класса",
    "Ли Даниил": "победу в математическом квесте",
    "Ворона Ангелина": "яркое выступление на новогоднем концерте",
    "Мороз Ярослава": "призовое место в конкурсе „Знатоки природы“",
    "Мороз Ярослав": "активное участие в экологической акции „Чистый двор“",
    "Голубь Николь": "отличную учёбу и активную позицию",
}

DARK_RED = RGBColor(0x8B, 0x1A, 0x1A)
NAVY = RGBColor(0x1B, 0x2A, 0x5C)


def _ru_key(s: str):
    order = "абвгдеёжзийклмнопрстуфхцчшщъыьэюя"
    return [order.index(ch) if ch in order else 100 + ord(ch) for ch in s.lower()]


# ------------------------------------------------------------------ frame image

def _frame_png() -> bytes:
    W = round(PAGE_CM[0] / 2.54 * FRAME_DPI)        # 1754
    H = round(PAGE_CM[1] / 2.54 * FRAME_DPI)        # 1240
    img = Image.new("RGB", (W, H), (255, 253, 246))
    d = ImageDraw.Draw(img)
    gold, navy, red = (184, 141, 46), (27, 42, 92), (190, 32, 32)
    # soft vignette towards the border
    yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
    edge = np.minimum.reduce([xx, yy, W - 1 - xx, H - 1 - yy])
    tint = np.clip(1.0 - edge / 140.0, 0, 1)[..., None]
    base = np.asarray(img, np.float32)
    warm = np.array([246, 236, 206], np.float32)
    img = Image.fromarray((base * (1 - 0.55 * tint) + warm * 0.55 * tint).astype(np.uint8), "RGB")
    d = ImageDraw.Draw(img)
    for inset, width, col in ((34, 14, navy), (58, 4, gold), (70, 2, gold), (84, 2, navy)):
        d.rectangle((inset, inset, W - 1 - inset, H - 1 - inset), outline=col, width=width)
    # corner rosettes
    for cx, cy in ((84, 84), (W - 85, 84), (84, H - 85), (W - 85, H - 85)):
        for k in range(12):
            a = k * math.pi / 6
            ex, ey = cx + 30 * math.cos(a), cy + 30 * math.sin(a)
            d.ellipse((ex - 12, ey - 12, ex + 12, ey + 12), outline=gold, width=3)
        d.ellipse((cx - 22, cy - 22, cx + 22, cy + 22), fill=navy)
        d.ellipse((cx - 9, cy - 9, cx + 9, cy + 9), fill=gold)
    # garland of small diamonds along the inner border
    for x in range(160, W - 160, 46):
        for y in (70, H - 71):
            d.polygon([(x, y - 7), (x + 7, y), (x, y + 7), (x - 7, y)], fill=gold)
    for y in range(160, H - 160, 46):
        for x in (70, W - 71):
            d.polygon([(x, y - 7), (x + 7, y), (x, y + 7), (x - 7, y)], fill=gold)
    # laurel branches at the bottom centre
    cx, cy = W // 2, H - 150
    for side in (-1, 1):
        for k in range(9):
            t = k / 8
            px = cx + side * (40 + 260 * t)
            py = cy - 70 * math.sin(math.pi * t * 0.9)
            ang = math.atan2(-70 * math.pi * 0.9 * math.cos(math.pi * t * 0.9) / 8, side * 260 / 8)
            lx, ly = 26 * math.cos(ang + side * 0.9), 26 * math.sin(ang + side * 0.9)
            d.ellipse((px - 9 + lx * 0.5, py - 9 + ly * 0.5, px + 9 + lx * 0.5, py + 9 + ly * 0.5), fill=(86, 128, 70))
            d.ellipse((px - 7 - lx * 0.4, py - 7 - ly * 0.4, px + 7 - lx * 0.4, py + 7 - ly * 0.4), fill=(110, 150, 86))
    d.ellipse((cx - 16, cy - 16, cx + 16, cy + 16), fill=gold)
    # specimen mark in the corner
    f = C.ff_pil_font("DejaVuSans-Bold", 26)
    d.text((W - 120, H - 118), "ОБРАЗЕЦ / SPECIMEN", font=f, fill=red, anchor="rs")
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=False, compress_level=9)
    return buf.getvalue()


# ------------------------------------------------------------------ template

def _line(doc, parts, *, size, bold=False, italic=False, color=None, space_before=0, space_after=0,
          font="Times New Roman"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Cm(0)
    for text in parts:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.name = font
        C.od_docx_run_fonts(r._r.get_or_add_rPr(), font, "ru-RU")
        r.bold = bold
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
    return p


def _build_template(path: pathlib.Path):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(PAGE_CM[0]), Cm(PAGE_CM[1])
    sec.left_margin = sec.right_margin = Cm(3.2)
    sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(2.6)
    sec.header_distance, sec.footer_distance = Cm(0.5), Cm(0.5)
    C.od_docx_base_style(doc, font="Times New Roman", size_pt=14, space_after_pt=0, justify=False)

    hp = sec.header.paragraphs[0]
    pic = hp.add_run().add_picture(io.BytesIO(_frame_png()), width=Cm(PAGE_CM[0]), height=Cm(PAGE_CM[1]))
    C.od_docx_inline_to_anchor(pic, x_emu=0, y_emu=0, behind=True, shape_id=1, name="Рамка грамоты")
    _ = Emu

    _line(doc, ["ГРАМОТА"], size=60, bold=True, color=DARK_RED, space_before=12, space_after=6,
          font="Times New Roman")
    _line(doc, ["вручается ", "{ученику/ученице}", " 5 «Б» класса"], size=22, italic=True, color=NAVY,
          space_after=14)
    _line(doc, ["{ФИО}"], size=34, bold=True, color=NAVY, space_after=14)
    _line(doc, ["за ", "{номинация}"], size=22, color=NAVY, space_after=54)
    p = _line(doc, ["Классный руководитель", "\t", "____________________", "\t", "И. И. Образцова"], size=15)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(2.5)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0))
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5))
    p.paragraph_format.space_after = Pt(18)
    p = _line(doc, ["24 декабря 2026 г."], size=15)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(2.5)
    C.docx_save_deterministic(doc, path, when=WHEN, author="И. И. Образцова (ОБРАЗЕЦ)")
    C.ooxml_fix_core_dates(path, when=WHEN)


# ------------------------------------------------------------------ class list

def _build_xlsx(path: pathlib.Path, rows):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "5Б"
    ws["A1"] = "5 «Б» класс — номинации для грамот, декабрь 2026 (ОБРАЗЕЦ / SPECIMEN)"
    ws.merge_cells("A1:B1")
    ws["A1"].font = Font(bold=True, size=13)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws["A2"], ws["B2"] = "Фамилия Имя", "Номинация"
    thin = Side(style="thin", color="999999")
    for cell in (ws["A2"], ws["B2"]):
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        cell.alignment = Alignment(horizontal="center")
    for i, (name, nom) in enumerate(rows, start=3):
        ws.cell(i, 1, name).border = Border(top=thin, bottom=thin, left=thin, right=thin)
        c = ws.cell(i, 2, nom)
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        c.alignment = Alignment(wrap_text=False)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 62
    ws.freeze_panes = "A3"
    ws.oddHeader.center.text = "ОБРАЗЕЦ / SPECIMEN"
    wb.properties.creator = "И. И. Образцова (ОБРАЗЕЦ)"
    wb.properties.lastModifiedBy = "И. И. Образцова (ОБРАЗЕЦ)"
    import datetime as _dt
    wb.properties.created = _dt.datetime(*WHEN)
    wb.properties.modified = _dt.datetime(*WHEN)
    wb.save(path)
    C.ooxml_fix_core_dates(path, when=WHEN)


def build(out_dir: pathlib.Path, truth_dir: pathlib.Path) -> list[str]:
    out_dir = pathlib.Path(out_dir)
    truth_dir = pathlib.Path(truth_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / DOCX_NAME
    xlsx_path = out_dir / XLSX_NAME
    _build_template(docx_path)

    assert set(NOMINATIONS) == {n for n, _g, _d in DG.GRAMOTY} and len(DG.GRAMOTY) == 25
    by_name = {n: (g, dat) for n, g, dat in DG.GRAMOTY}
    ordered = sorted(by_name, key=_ru_key)
    rows = [(n, NOMINATIONS[n]) for n in ordered]
    _build_xlsx(xlsx_path, rows)

    truth = {
        "case": CASE,
        "files": {"template": DOCX_NAME, "list": XLSX_NAME},
        "template": {
            "page": "A4 landscape 29.7×21 cm",
            "frame": "word/media/image1.png in header1.xml, wp:anchor behindDoc=1 at page (0,0), 29.7×21 cm",
            "placeholders": ["{ученику/ученице}", "{ФИО}", "{номинация}"],
            "lines": ["ГРАМОТА", "вручается {ученику/ученице} 5 «Б» класса", "{ФИО}", "за {номинация}",
                      "Классный руководитель ____ И. И. Образцова", "24 декабря 2026 г."],
        },
        "sheet": {"name": "5Б", "title_row": 1, "header_row": 2, "first_data_row": 3, "last_data_row": 27,
                  "columns": ["Фамилия Имя", "Номинация"]},
        "rows": [{"sheet_row": i, "nominative": n, "gender": by_name[n][0], "dative": by_name[n][1],
                  "gender_word": "ученику" if by_name[n][0] == "м" else "ученице", "nomination": nom,
                  "expected_lines": [f"вручается {'ученику' if by_name[n][0] == 'м' else 'ученице'} 5 «Б» класса",
                                     by_name[n][1], f"за {nom}"]}
                 for i, (n, nom) in enumerate(rows, start=3)],
        "indeclinable": [d for n, g, d in DG.GRAMOTY if n == d],
    }
    tp = truth_dir / "gramoty_5B.json"
    C.ff_write_json(tp, truth)
    return [str(docx_path), str(xlsx_path), str(tp)]


if __name__ == "__main__":
    root = HERE.parent
    for p in build(root / "out" / CASE, root / "truth"):
        print(p)
