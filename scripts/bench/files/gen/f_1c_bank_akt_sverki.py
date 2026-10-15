"""Fixtures for f-1c-bank-akt-sverki (ИП Образцов О.О. ↔ ООО «Тест-Поставка», III кв. 2026).

kl_to_1c.txt — bank export in 1CClientBankExchange 1.03, cp1251, CRLF. Header + СекцияРасчСчет with
  opening/turnover/closing balances that add up, then one СекцияДокумент=Платежное поручение per
  operation in date order: 5 payments from ООО «Тест-Поставка» ИНН 7700000017 (datagen.PAY_FROM_THEM),
  2 from the decoy ООО «Тест-Поставка Плюс» ИНН 7700000024 (datagen.PAY_FROM_DECOY), and
  datagen.OTHER_BANK_OPS. Payment п/п 361 of 29.09 (their act) is deliberately absent (arrived in Q4).
upd_test_postavka.xlsx — sheet «УПД»: header row 3 «№ | Дата | Покупатель | ИНН | Сумма, руб.»,
  datagen.UPD (6 rows to ООО «Тест-Поставка») interleaved with 3 УПД to other buyers (2 to
  ООО «Другой Клиент», 1 to the decoy «Тест-Поставка Плюс» for 15 000 — matches its payment), and a
  note row «Сальдо на 01.07.2026: долг ООО «Тест-Поставка» 50 000,00».
akt_ot_kontragenta.pdf — the counterparty's reconciliation act (their books, mirrored: our sales are
  their credit, their payments their debit) built from datagen.THEIR_ACT; their columns filled, ours
  blank; their closing 23 700,00 in our favour; signature, round seal «ОБРАЗЕЦ» (vector, no text layer),
  red «ОБРАЗЕЦ / SPECIMEN» mark.

truth/akt_sverki.json: expected numbers (datagen.akt_expected) + the parsed operations of each file.
"""

from __future__ import annotations

import datetime as dt
import pathlib
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))
import common as C  # noqa: E402
import datagen as G  # noqa: E402

CASE = "f-1c-bank-akt-sverki"
F_TXT = "kl_to_1c.txt"
F_XLSX = "upd_test_postavka.xlsx"
F_PDF = "akt_ot_kontragenta.pdf"

PERIOD = (dt.date(2026, 7, 1), dt.date(2026, 9, 30))
CREATED = (dt.date(2026, 10, 1), "09:14:37")
ACC_OPENING = 187_432.15  # our account balance on 01.07.2026 (any value; turnovers must add up)

US_BANK = dict(bank1='АО "Банк Образец"', bank2="г. Москва", bik=G.US["bik"], corr=G.US["corr"])
OTHER_BANK = dict(bank1='ПАО "Тест-Банк"', bank2="г. Москва", bik="044500001", corr="30101810000000000001")

# extra details for counterparties that datagen gives only name/INN for
PARTY = {
    G.THEM["inn"]: dict(name=G.THEM["name"], inn=G.THEM["inn"], kpp=G.THEM["kpp"], acc=G.THEM["acc"], **OTHER_BANK),
    G.DECOY["inn"]: dict(name=G.DECOY["name"], inn=G.DECOY["inn"], kpp=G.DECOY["kpp"], acc=G.DECOY["acc"], **OTHER_BANK),
    "7700000031": dict(name="ООО «Хостинг-Образец»", inn="7700000031", kpp="770001001", acc="40702810000000000031", **OTHER_BANK),
    # treasury-shaped but fictional requisites (the real УФК по г. Москве account/BIK/ЕКС must not appear)
    "7700000048": dict(name="УФК по г. Москве (ИФНС)", inn="7700000048", kpp="770001001", acc="03100643000000000048",
                       bank1="ГУ Банка Образец//УФК по г. Москве (ОБРАЗЕЦ)", bank2="г. Москва", bik="004500048",
                       corr="40102810000000000048"),
    "7700000055": dict(name="ООО «Другой Клиент»", inn="7700000055", kpp="770001001", acc="40702810000000000055", **OTHER_BANK),
    "770000000029": dict(name="ИП Тестова А.А.", inn="770000000029", kpp="", acc="40802810000000000029", **OTHER_BANK),
    "7700000062": dict(name="АО «Банк Образец»", inn="7700000062", kpp="770001001", acc="70601810000000000062",
                       bank1=US_BANK["bank1"], bank2=US_BANK["bank2"], bik=US_BANK["bik"], corr=US_BANK["corr"]),
}

# what each incoming payment from ООО «Тест-Поставка» pays (first one = opening debt from Q2)
THEM_PURPOSE = {
    "214": "Оплата по УПД № 97 от 26.06.2026. НДС не облагается",
    "251": "Оплата по УПД № 101 от 03.07.2026. НДС не облагается",
    "287": "Оплата по УПД № 105 от 17.07.2026. НДС не облагается",
    "305": "Оплата по УПД № 109 от 05.08.2026. НДС не облагается",
    "342": "Оплата по УПД № 112 от 19.08.2026. НДС не облагается",
}
DECOY_PURPOSE = {
    "77": "Оплата по УПД № 110 от 07.08.2026. НДС не облагается",
    "91": "Предоплата по счёту № 131 от 09.09.2026. НДС не облагается",
}
# our own outgoing payment-order numbers, and the payer's numbers of the other incoming payments
OUT_NUM = {"7700000031": "58", "7700000048": "61", "770000000029": "64", "7700000062": "67"}
IN_NUM = {("7700000055", dt.date(2026, 8, 3)): "1187", ("7700000055", dt.date(2026, 9, 28)): "1342"}

# УПД to other buyers (interleaved with datagen.UPD in the register)
OTHER_UPD = [
    ("102", dt.date(2026, 7, 9), "ООО «Другой Клиент»", "7700000055", 64_000.00),
    ("110", dt.date(2026, 8, 7), G.DECOY["name"], G.DECOY["inn"], 15_000.00),
    ("116", dt.date(2026, 9, 14), "ООО «Другой Клиент»", "7700000055", 28_500.00),
]


def _q(name: str) -> str:
    """1C/bank exports use straight quotes: ООО «X» -> ООО "X"."""
    return name.replace("«", '"').replace("»", '"')


def _d(d: dt.date) -> str:
    return d.strftime("%d.%m.%Y")


def _amt(x: float) -> str:
    return f"{x:.2f}"


def bank_operations() -> list[dict]:
    """All account operations of the 1C file, in date order (stable within a date)."""
    ops = []
    for d, num, amt in G.PAY_FROM_THEM:
        ops.append(dict(date=d, num=num, amount=amt, direction="in", party=G.THEM["inn"], purpose=THEM_PURPOSE[num]))
    for d, num, amt in G.PAY_FROM_DECOY:
        ops.append(dict(date=d, num=num, amount=amt, direction="in", party=G.DECOY["inn"], purpose=DECOY_PURPOSE[num]))
    for d, direction, _name, inn, amt, purpose in G.OTHER_BANK_OPS:
        num = OUT_NUM[inn] if direction == "out" else IN_NUM[(inn, d)]
        ops.append(dict(date=d, num=num, amount=amt, direction=direction, party=inn, purpose=purpose))
    ops.sort(key=lambda o: (o["date"], o["direction"] != "in", o["num"]))
    return ops


def _doc_lines(op: dict) -> list[str]:
    us = dict(name=G.US["name"], inn=G.US["inn"], kpp="", acc=G.US["acc"], **US_BANK)
    other = PARTY[op["party"]]
    payer, payee = (other, us) if op["direction"] == "in" else (us, other)
    tax = op["party"] == "7700000048"
    L = [
        "СекцияДокумент=Платежное поручение",
        f"Номер={op['num']}",
        f"Дата={_d(op['date'])}",
        f"Сумма={_amt(op['amount'])}",
        f"ПлательщикСчет={payer['acc']}",
    ]
    L.append(f"ДатаСписано={_d(op['date'])}" if op["direction"] == "out" else "ДатаСписано=")
    L += [
        f"Плательщик=ИНН {payer['inn']} {_q(payer['name'])}",
        f"ПлательщикИНН={payer['inn']}",
        f"Плательщик1={_q(payer['name'])}",
        f"ПлательщикРасчСчет={payer['acc']}",
        f"ПлательщикБанк1={payer['bank1']}",
        f"ПлательщикБанк2={payer['bank2']}",
        f"ПлательщикБИК={payer['bik']}",
        f"ПлательщикКорсчет={payer['corr']}",
        f"ПолучательСчет={payee['acc']}",
    ]
    L.append(f"ДатаПоступило={_d(op['date'])}" if op["direction"] == "in" else "ДатаПоступило=")
    L += [
        f"Получатель=ИНН {payee['inn']} {_q(payee['name'])}",
        f"ПолучательИНН={payee['inn']}",
        f"Получатель1={_q(payee['name'])}",
        f"ПолучательРасчСчет={payee['acc']}",
        f"ПолучательБанк1={payee['bank1']}",
        f"ПолучательБанк2={payee['bank2']}",
        f"ПолучательБИК={payee['bik']}",
        f"ПолучательКорсчет={payee['corr']}",
        "ВидПлатежа=",
        "ВидОплаты=01",
        "Код=0",
        f"СтатусСоставителя={'01' if tax else ''}",
        f"ПлательщикКПП={payer['kpp'] if not tax else '0'}",
        f"ПолучательКПП={payee['kpp']}",
        f"ПоказательКБК={'18201061201010000510' if tax else ''}",
        f"ОКАТО={'0' if tax else ''}",
        f"ПоказательОснования={'0' if tax else ''}",
        f"ПоказательПериода={'0' if tax else ''}",
        f"ПоказательНомера={'0' if tax else ''}",
        f"ПоказательДаты={'0' if tax else ''}",
        "ПоказательТипа=",
        "Очередность=5",
        f"НазначениеПлатежа={op['purpose']}",
        "КонецДокумента",
    ]
    return L


def build_1c_text() -> tuple[str, dict]:
    ops = bank_operations()
    tin = round(sum(o["amount"] for o in ops if o["direction"] == "in"), 2)
    tout = round(sum(o["amount"] for o in ops if o["direction"] == "out"), 2)
    closing = round(ACC_OPENING + tin - tout, 2)
    L = [
        "1CClientBankExchange",
        "ВерсияФормата=1.03",
        "Кодировка=Windows",
        "Отправитель=Банк Образец — интернет-банк (ОБРАЗЕЦ / SPECIMEN)",
        "Получатель=Бухгалтерия предприятия, редакция 3.0",
        f"ДатаСоздания={_d(CREATED[0])}",
        f"ВремяСоздания={CREATED[1]}",
        f"ДатаНачала={_d(PERIOD[0])}",
        f"ДатаКонца={_d(PERIOD[1])}",
        f"РасчСчет={G.US['acc']}",
        "Документ=Платежное поручение",
        "СекцияРасчСчет",
        f"ДатаНачала={_d(PERIOD[0])}",
        f"ДатаКонца={_d(PERIOD[1])}",
        f"РасчСчет={G.US['acc']}",
        f"НачальныйОстаток={_amt(ACC_OPENING)}",
        f"ВсегоПоступило={_amt(tin)}",
        f"ВсегоСписано={_amt(tout)}",
        f"КонечныйОстаток={_amt(closing)}",
        "КонецРасчСчет",
    ]
    for op in ops:
        L += _doc_lines(op)
    L.append("КонецФайла")
    text = "\r\n".join(L) + "\r\n"
    summary = dict(account=G.US["acc"], opening=ACC_OPENING, total_in=tin, total_out=tout, closing=closing,
                   operations=[dict(date=o["date"].isoformat(), number=o["num"], amount=o["amount"],
                                    direction=o["direction"], counterparty_inn=o["party"],
                                    counterparty=PARTY[o["party"]]["name"], purpose=o["purpose"]) for o in ops])
    return text, summary


def build_upd_xlsx(path: pathlib.Path) -> list[dict]:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    rows = [(n, d, G.THEM["name"], G.THEM["inn"], a) for n, d, a in G.UPD] + list(OTHER_UPD)
    rows.sort(key=lambda r: (r[1], r[0]))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "УПД"
    ws["A1"] = "Реестр выставленных УПД — ИП Образцов О. О., III квартал 2026 г.   ОБРАЗЕЦ / SPECIMEN"
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"] = "Все данные вымышленные (ОБРАЗЕЦ / SPECIMEN). Услуги НДС не облагаются (УСН)."
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="C00000")
    head = ["№", "Дата", "Покупатель", "ИНН", "Сумма, руб."]
    thin = Side(style="thin", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, h in enumerate(head, start=1):
        cell = ws.cell(row=3, column=j, value=h)
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    for i, (n, d, buyer, inn, amt) in enumerate(rows, start=4):
        vals = [n, d, buyer, inn, amt]
        for j, v in enumerate(vals, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = Font(name="Arial", size=10)
            cell.border = border
        ws.cell(row=i, column=1).number_format = "@"
        ws.cell(row=i, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=i, column=2).number_format = "DD.MM.YYYY"
        ws.cell(row=i, column=4).number_format = "@"
        ws.cell(row=i, column=5).number_format = "#,##0.00"
    note_row = 4 + len(rows) + 1
    ws.cell(row=note_row, column=1,
            value="Примечание: сальдо на 01.07.2026 — долг ООО «Тест-Поставка» 50 000,00 "
                  "(по акту сверки за II квартал 2026 г., УПД № 97 от 26.06.2026).").font = Font(name="Arial", size=10, italic=True)
    for col, w in zip("ABCDE", (8, 12, 30, 15, 15)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    wb.properties.creator = "ОБРАЗЕЦ / SPECIMEN"
    wb.properties.title = "Реестр УПД (ОБРАЗЕЦ)"
    wb.save(path)
    C.ooxml_fix_core_dates(path)
    return [dict(number=n, date=d.isoformat(), buyer=b, inn=inn, amount=a) for n, d, b, inn, a in rows]


def _their_rows() -> list[tuple]:
    """(date, text, debit, credit) in the counterparty's books (they are the buyer)."""
    rows = []
    for num, d, amt in G.THEIR_ACT["purchases"]:
        if num.startswith("118-К"):
            rows.append((d, "Корректировка (УКД № 118-К от 28.09.2026)", None, amt))
        else:
            rows.append((d, f"Поступление (УПД № {num} от {_d(d)})", None, amt))
    for d, num, amt in G.THEIR_ACT["payments"]:
        rows.append((d, f"Оплата (п/п № {num} от {_d(d)})", amt, None))
    rows.sort(key=lambda r: (r[0], r[2] is not None))
    return rows


def build_their_act_pdf(path: pathlib.Path) -> dict:
    from reportlab.lib.colors import Color

    C.ff_register_rl_fonts("LiberationSans", "LiberationSans-Bold", "DejaVuSans-Bold")
    W, H = C.FF_A4
    c = C.ff_rl_canvas(str(path), title="Акт сверки взаимных расчетов (ОБРАЗЕЦ)",
                       subject="Акт сверки ООО «Тест-Поставка» — ИП Образцов О.О., III кв. 2026 (ОБРАЗЕЦ)")
    them, us = "ООО «Тест-Поставка»", "ИП Образцов Олег Олегович"
    fmt = C.bk_fmt_money

    # specimen marks: text-layer label on top, outline watermark across the page
    c.setFillColorRGB(0.78, 0.1, 0.1)
    c.setFont("LiberationSans-Bold", 9)
    c.drawRightString(W - 28, H - 22, "ОБРАЗЕЦ / SPECIMEN — вымышленные данные")
    C.bk_rl_watermark(c, cx=W / 2, cy=H / 2 - 40, size=58, angle=33, alpha=0.10)

    c.setFillColorRGB(0, 0, 0)
    y = H - 58
    c.setFont("LiberationSans-Bold", 13)
    c.drawCentredString(W / 2, y, "Акт сверки")
    y -= 16
    c.setFont("LiberationSans", 9.5)
    c.drawCentredString(W / 2, y, "взаимных расчетов за период: 01.07.2026 - 30.09.2026")
    y -= 13
    c.drawCentredString(W / 2, y, f"между {them}")
    y -= 13
    c.drawCentredString(W / 2, y, f"и {us}")
    y -= 13
    c.drawCentredString(W / 2, y, "по договору поставки услуг № 12/26 от 15.01.2026")
    y -= 22
    intro = [
        f"Мы, нижеподписавшиеся, Генеральный директор {them} Тестов Тимур Тимурович, с одной стороны,",
        f"и {us}, с другой стороны, составили настоящий акт сверки в том, что состояние",
        "взаимных расчетов по данным учета следующее:",
    ]
    c.setFont("LiberationSans", 8.5)
    for line in intro:
        c.drawString(32, y, line)
        y -= 11.5
    y -= 6

    # table geometry: left half = their data, right half = ours (left blank by the sender)
    x0, x1 = 28, W - 28
    # their: date | document | debit | credit  ||  ours: date | document | debit | credit
    cols = [x0, x0 + 50, x0 + 222, x0 + 272, x0 + 323, x0 + 369, x0 + 449, x0 + 494, x1]
    row_h = 13.0
    c.setLineWidth(0.6)
    c.setStrokeColorRGB(0, 0, 0)

    def hline(yy, a=x0, b=x1):
        c.line(a, yy, b, yy)

    def text_in(col_i, yy, s, *, align="left", font="LiberationSans", size=7.6):
        c.setFont(font, size)
        a, b = cols[col_i], cols[col_i + 1]
        if align == "left":
            c.drawString(a + 2.5, yy, s)
        elif align == "right":
            c.drawRightString(b - 2.5, yy, s)
        else:
            c.drawCentredString((a + b) / 2, yy, s)

    top = y
    # header row 1 (spanning)
    hline(top)
    c.setFont("LiberationSans-Bold", 7.8)
    c.drawCentredString((cols[0] + cols[4]) / 2, top - 10, f"По данным {them}, руб.")
    c.drawCentredString((cols[4] + cols[8]) / 2, top - 10, "По данным ИП Образцов О. О.")
    y = top - row_h - 1
    hline(y)
    for i, h in enumerate(["Дата", "Документ", "Дебет", "Кредит", "Дата", "Документ", "Дебет", "Кредит"]):
        text_in(i, y - 9.5, h, align="center", font="LiberationSans-Bold")
    y -= row_h
    hline(y)
    body_top = y

    rows = _their_rows()
    opening = G.THEIR_ACT["opening"]
    deb_total = round(sum(r[2] for r in rows if r[2] is not None), 2)
    cred_total = round(sum(r[3] for r in rows if r[3] is not None), 2)
    closing = round(opening + cred_total - deb_total, 2)

    def span_row(label_left, deb, cred, *, bold=True, label_right=None):
        nonlocal y
        f = "LiberationSans-Bold" if bold else "LiberationSans"
        c.setFont(f, 7.6)
        c.drawString(cols[0] + 2.5, y - 9.5, label_left)
        if deb is not None:
            text_in(2, y - 9.5, fmt(deb), align="right", font=f)
        if cred is not None:
            text_in(3, y - 9.5, fmt(cred), align="right", font=f)
        if label_right:
            c.setFont(f, 7.6)
            c.drawString(cols[4] + 2.5, y - 9.5, label_right)
        y -= row_h
        hline(y)

    span_row("Сальдо начальное", None, opening, label_right="Сальдо начальное")
    for d, doc, deb, cred in rows:
        text_in(0, y - 9.5, _d(d))
        text_in(1, y - 9.5, doc, size=6.9 if len(doc) > 30 else 7.6)
        if deb is not None:
            text_in(2, y - 9.5, fmt(deb), align="right")
        if cred is not None:
            text_in(3, y - 9.5, fmt(cred), align="right")
        y -= row_h
        hline(y)
    span_row("Обороты за период", deb_total, cred_total, label_right="Обороты за период")
    span_row("Сальдо конечное", None, closing, label_right="Сальдо конечное")
    bottom = y
    # vertical lines: outer + the half split everywhere; inner columns except on span rows
    c.line(cols[0], top, cols[0], bottom)
    c.line(cols[8], top, cols[8], bottom)
    c.setLineWidth(1.1)
    c.line(cols[4], top, cols[4], bottom)
    c.setLineWidth(0.6)
    for i in (1, 2, 3, 5, 6, 7):
        # header row 2 + transaction rows; span rows (opening / turnover / closing) keep the doc column merged
        c.line(cols[i], top - row_h - 1, cols[i], body_top)
        c.line(cols[i], body_top - row_h, cols[i], bottom + 2 * row_h)
        if i in (2, 3, 6, 7):
            c.line(cols[i], body_top, cols[i], body_top - row_h)
            c.line(cols[i], bottom + 2 * row_h, cols[i], bottom)

    y = bottom - 20
    c.setFont("LiberationSans", 8.5)
    c.drawString(32, y, f"По данным {them}")
    y -= 12
    c.setFont("LiberationSans-Bold", 8.5)
    c.drawString(32, y, f"на 30.09.2026 задолженность в пользу {us} {fmt(closing)} руб.")
    y -= 12
    c.setFont("LiberationSans", 8.5)
    c.drawString(32, y, f"({C.bk_rub_words(closing, capitalize=True)})")

    # signatures
    y -= 34
    c.setFont("LiberationSans", 8.5)
    c.drawString(32, y, f"От {them}")
    c.drawString(W / 2 + 10, y, f"От {us}")
    y -= 14
    c.drawString(32, y, "Генеральный директор")
    c.drawString(W / 2 + 10, y, "Индивидуальный предприниматель")
    y -= 24
    c.line(32, y, 150, y)
    c.drawString(156, y + 1, "(Тестов Т. Т.)")
    c.line(W / 2 + 10, y, W / 2 + 128, y)
    c.drawString(W / 2 + 134, y + 1, "(Образцов О. О.)")
    C.bk_rl_signature(c, 44, y - 4, 100, 26, seed=361)
    y -= 14
    c.setFont("LiberationSans", 7.5)
    c.drawString(60, y, "М.П.")
    c.drawString(W / 2 + 40, y, "М.П.")
    C.bk_rl_round_stamp(c, 150, y + 6, 36, "ООО ТЕСТ-ПОСТАВКА • ОБРАЗЕЦ • ИНН 7700000017",
                        ["ОБРАЗЕЦ", "SPECIMEN"])
    y -= 58
    c.setFont("LiberationSans", 7)
    c.setFillColorRGB(0.35, 0.35, 0.35)
    c.drawString(32, y, "Документ сформирован в учебных целях. Все организации, лица и реквизиты вымышлены. ОБРАЗЕЦ / SPECIMEN.")
    c.showPage()
    c.save()
    return dict(opening_credit=opening, debit_turnover=deb_total, credit_turnover=cred_total, closing_credit=closing,
                rows=[dict(date=d.isoformat(), document=doc, debit=deb, credit=cred) for d, doc, deb, cred in rows])


def build(out_dir: pathlib.Path, truth_dir: pathlib.Path) -> list[str]:
    out_dir = pathlib.Path(out_dir)
    truth_dir = pathlib.Path(truth_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    text, bank = build_1c_text()
    p_txt = out_dir / F_TXT
    p_txt.write_bytes(text.encode("cp1251"))

    p_xlsx = out_dir / F_XLSX
    upd = build_upd_xlsx(p_xlsx)

    p_pdf = out_dir / F_PDF
    their = build_their_act_pdf(p_pdf)

    exp = G.akt_expected()
    truth = dict(
        case=CASE,
        note="Ожидаемые числа — datagen.akt_expected(); ниже разбор фикстур, на котором они держатся.",
        expected=exp,
        us=G.US, counterparty=G.THEM, decoy=G.DECOY,
        period=[PERIOD[0].isoformat(), PERIOD[1].isoformat()],
        kl_to_1c=dict(encoding="cp1251", newline="CRLF", format="1CClientBankExchange 1.03", **bank),
        upd_register=dict(sheet="УПД", header_row=3, rows=upd, opening_note_debt=G.AKT_OPENING),
        their_act=their,
        derivation=dict(
            our_debit=[dict(upd=n, date=d.isoformat(), amount=a) for n, d, a in G.UPD],
            our_credit=[dict(pp=n, date=d.isoformat(), amount=a) for d, n, a in G.PAY_FROM_THEM],
            excluded_decoy=[dict(pp=n, date=d.isoformat(), amount=a, inn=G.DECOY["inn"]) for d, n, a in G.PAY_FROM_DECOY],
        ),
    )
    tpath = C.ff_write_json(truth_dir / "akt_sverki.json", truth)
    return [str(p_txt), str(p_xlsx), str(p_pdf), tpath]


if __name__ == "__main__":
    base = HERE.parent
    for p in build(base / "out" / CASE, base / "truth"):
        print(p)
