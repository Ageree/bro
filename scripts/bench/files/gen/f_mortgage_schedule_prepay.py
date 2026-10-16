"""f-mortgage-schedule-prepay has no input files (everything is in the chat message).

This module writes only ground truth:

truth/ipoteka_grafik.json   — datagen.mortgage_expected() plus both full schedules (with and without the
                              two prepayments) from datagen.mortgage(), so xlsx_last_payment / judge can
                              compare any row.
truth/ipoteka_reference.xlsx — a reference answer for checker self-tests: an input block (сумма, ставка,
                              срок, даты, досрочки — coloured), the schedule as live formulas (interest =
                              ROUND(prev balance × $rate$/12; 2), payment level carried per row, prepayment
                              «срок» keeps the payment, «платёж» re-annuitises over ROUNDUP(NPER(...)) payments
                              left — the count left in the schedule BEFORE the prepayment, so the term is kept), a second sheet without prepayments, and totals on top (переплата,
                              экономия, число платежей, дата последнего). No cached values are stored;
                              LibreOffice recalculates on load. Rows after payoff are blank (date and № are "").

The reference is verified by recalculating it with LibreOffice (see __main__ with --verify).
"""

from __future__ import annotations

import datetime as dt
import pathlib
import sys

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))
import common as C  # noqa: E402
import datagen as G  # noqa: E402

CASE = "f-mortgage-schedule-prepay"
ROW0 = 20          # first schedule row
NROWS = 360        # rows reserved (original term)
INPUTS = dict(price=9_800_000, down=2_300_000, rate=0.06, term=360, issue=dt.date(2026, 10, 20),
              first=dt.date(2026, 11, 20), p1=(dt.date(2027, 6, 20), 500_000, "срок"),
              p2=(dt.date(2027, 12, 20), 300_000, "платёж"))


def _schedule_sheet(ws, *, with_prepay: bool, inp: str):
    """Schedule rows ROW0..ROW0+NROWS-1. `inp` is the sheet-name prefix for input cells ('' or "'Ипотека'!")."""
    from openpyxl.styles import Alignment, Font, PatternFill

    head = ["№", "Дата платежа", "Остаток на начало", "Проценты", "Платёж", "Основной долг", "Досрочка",
            "Остаток на конец", "Платёж (уровень)", "Платёж далее"]
    if not with_prepay:
        head = head[:6] + ["", "Остаток на конец", "Платёж (уровень)", ""]
    for j, h in enumerate(head, start=1):
        cell = ws.cell(row=ROW0 - 1, column=j, value=h or None)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    r_ = f"{inp}$B$6/12"
    for i in range(NROWS):
        r = ROW0 + i
        p = r - 1
        alive = f"C{r}>0.005"
        ws[f"A{r}"] = f'=IF({alive},ROW()-{ROW0 - 1},"")'
        ws[f"B{r}"] = f'=IF({alive},EDATE({inp}$B$9,ROW()-{ROW0}),"")'
        ws[f"C{r}"] = f"={inp}$B$5" if i == 0 else f"=H{p}"
        ws[f"D{r}"] = f"=IF({alive},ROUND(C{r}*{r_},2),0)"
        ws[f"I{r}"] = "=$E$3" if (i == 0 and with_prepay) else (f"={inp}$E$3" if i == 0 else (f"=J{p}" if with_prepay else f"=I{p}"))
        ws[f"E{r}"] = f"=IF({alive},IF(C{r}+D{r}>I{r}+0.005,I{r},C{r}+D{r}),0)"
        ws[f"F{r}"] = f"=ROUND(E{r}-D{r},2)"
        if with_prepay:
            ws[f"G{r}"] = (f"=IF(AND(C{r}-F{r}>0.005,OR(B{r}=$B$10,B{r}=$B$13)),"
                           f"MIN(IF(B{r}=$B$10,$B$11,$B$14),C{r}-F{r}),0)")
            ws[f"H{r}"] = f"=ROUND(C{r}-F{r}-G{r},2)"
            ws[f"J{r}"] = (f'=IF(AND(G{r}>0,H{r}>0.005,OR(AND(B{r}=$B$10,$B$12="платёж"),AND(B{r}=$B$13,$B$15="платёж"))),'
                           f"ROUNDUP(H{r}*{r_}/(1-(1+{r_})^(-ROUNDUP(NPER({r_},-I{r},C{r}-F{r}),0))),2),I{r})")
        else:
            ws[f"H{r}"] = f"=ROUND(C{r}-F{r},2)"
        ws[f"B{r}"].number_format = "DD.MM.YYYY"
        for col in "CDEFGHIJ":
            ws[f"{col}{r}"].number_format = '#,##0.00;-#,##0.00;""'
    for col, w in zip("ABCDEFGHIJ", (6, 13, 16, 12, 12, 14, 12, 16, 14, 14)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = f"A{ROW0}"


def build_reference_xlsx(path: pathlib.Path) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    last = ROW0 + NROWS - 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ипотека"
    ws["A1"] = "График платежей по ипотеке (эталон для самопроверки проверок) — ОБРАЗЕЦ / SPECIMEN"
    ws["A1"].font = Font(bold=True, size=12)
    inp_fill = PatternFill("solid", fgColor="FFF2CC")
    rows = [
        ("Стоимость квартиры", INPUTS["price"], "#,##0.00"),
        ("Первоначальный взнос", INPUTS["down"], "#,##0.00"),
        ("Сумма кредита", "=B3-B4", "#,##0.00"),
        ("Ставка, % годовых", INPUTS["rate"], "0.00%"),
        ("Срок, мес.", INPUTS["term"], "0"),
        ("Дата выдачи", INPUTS["issue"], "DD.MM.YYYY"),
        ("Дата первого платежа", INPUTS["first"], "DD.MM.YYYY"),
        ("Досрочка 1 — дата", INPUTS["p1"][0], "DD.MM.YYYY"),
        ("Досрочка 1 — сумма", INPUTS["p1"][1], "#,##0.00"),
        ("Досрочка 1 — уменьшить (срок/платёж)", INPUTS["p1"][2], "@"),
        ("Досрочка 2 — дата", INPUTS["p2"][0], "DD.MM.YYYY"),
        ("Досрочка 2 — сумма", INPUTS["p2"][1], "#,##0.00"),
        ("Досрочка 2 — уменьшить (срок/платёж)", INPUTS["p2"][2], "@"),
    ]
    for i, (label, val, fmt) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label)
        c = ws.cell(row=i, column=2, value=val)
        c.number_format = fmt
        if not (isinstance(val, str) and val.startswith("=")):
            c.fill = inp_fill
    last_b = f"'Без досрочек'!"
    totals = [
        ("Аннуитетный платёж", "=ROUNDUP(B5*(B6/12)/(1-(1+B6/12)^(-B7)),2)", "#,##0.00"),
        ("Платёж после досрочки 2", f"=IFERROR(INDEX(E{ROW0}:E{last},MATCH(B13,B{ROW0}:B{last},0)+1),E3)", "#,##0.00"),
        ("Число платежей", f'=COUNTIF(E{ROW0}:E{last},">0")', "0"),
        ("Дата последнего платежа", f"=MAX(B{ROW0}:B{last})", "DD.MM.YYYY"),
        ("Переплата (проценты) с досрочками", f"=SUM(D{ROW0}:D{last})", "#,##0.00"),
        ("Переплата без досрочек", f"=SUM({last_b}D{ROW0}:D{last})", "#,##0.00"),
        ("Экономия на досрочках", "=E8-E7", "#,##0.00"),
        ("Всего выплачено (с досрочками)", f"=SUM(E{ROW0}:E{last})+SUM(G{ROW0}:G{last})", "#,##0.00"),
    ]
    for i, (label, f, fmt) in enumerate(totals, start=3):
        ws.cell(row=i, column=4, value=label).font = Font(bold=True)
        c = ws.cell(row=i, column=5, value=f)
        c.number_format = fmt
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["D"].width = 34
    _schedule_sheet(ws, with_prepay=True, inp="")
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["D"].width = 34
    ws2 = wb.create_sheet("Без досрочек")
    ws2["A1"] = "Тот же кредит без досрочных погашений — ОБРАЗЕЦ / SPECIMEN"
    ws2["E3"] = "='Ипотека'!E3"
    _schedule_sheet(ws2, with_prepay=False, inp="'Ипотека'!")
    for r in range(ROW0, ROW0 + NROWS):  # level column refers to own E3 copy
        if r == ROW0:
            ws2[f"I{r}"] = "=$E$3"
    wb.properties.creator = "ОБРАЗЕЦ / SPECIMEN"
    wb.save(path)
    C.ooxml_fix_core_dates(path)


def build(out_dir: pathlib.Path, truth_dir: pathlib.Path) -> list[str]:
    out_dir = pathlib.Path(out_dir)
    truth_dir = pathlib.Path(truth_dir)
    out_dir.mkdir(parents=True, exist_ok=True)  # the case has no input files
    exp = G.mortgage_expected()
    truth = dict(
        case=CASE,
        fixtures=[],
        inputs=dict(price=INPUTS["price"], down=INPUTS["down"], principal=INPUTS["price"] - INPUTS["down"],
                    rate=INPUTS["rate"], term_months=INPUTS["term"], issue=INPUTS["issue"].isoformat(),
                    first_payment=INPUTS["first"].isoformat(),
                    prepayments=[dict(date=INPUTS["p1"][0].isoformat(), amount=INPUTS["p1"][1], mode="term", after_payment=8),
                                 dict(date=INPUTS["p2"][0].isoformat(), amount=INPUTS["p2"][1], mode="payment", after_payment=14)]),
        conventions=("interest = round(balance × rate / 12, 2); payment rounded up to the kopeck; the last payment is "
                     "the short one; a prepayment is applied right after the regular payment of that date; «срок» "
                     "keeps the payment, «платёж» re-annuitises over the payments left in the current schedule"),
        expected=exp,
        schedule_with=G.mortgage(),
        schedule_without=G.mortgage(prepay=()),
    )
    t1 = C.ff_write_json(truth_dir / "ipoteka_grafik.json", truth)
    ref = truth_dir / "ipoteka_reference.xlsx"
    build_reference_xlsx(ref)
    return [t1, str(ref)]


def verify_reference(path: pathlib.Path) -> dict:
    """Recalculate the reference with LibreOffice, also with the rate set to 7 %, and compare with datagen."""
    import shutil
    import tempfile

    import openpyxl

    exp = G.mortgage_expected()
    v = C.xlsx_recalc_values(path)["Ипотека"]
    got = dict(annuity=v[(3, 5)], payment_after_dec2027=v[(4, 5)], payments_with=v[(5, 5)], last_date_with=v[(6, 5)],
               interest_with=v[(7, 5)], interest_without=v[(8, 5)], saving=v[(9, 5)],
               formula_cells=None)
    wb = openpyxl.load_workbook(path)
    got["formula_cells"] = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row
                               if isinstance(c.value, str) and c.value.startswith("="))
    with tempfile.TemporaryDirectory() as td:
        cp = pathlib.Path(td) / "probe.xlsx"
        shutil.copy(path, cp)
        wb = openpyxl.load_workbook(cp)
        wb["Ипотека"]["B6"] = 0.07
        wb.save(cp)
        got["annuity_if_rate_7"] = C.xlsx_recalc_values(cp)["Ипотека"][(3, 5)]
    return dict(expected={k: exp.get(k) for k in got if k in exp}, got=got)


if __name__ == "__main__":
    base = HERE.parent
    for p in build(base / "out" / CASE, base / "truth"):
        print(p)
    if "--verify" in sys.argv:
        import json

        print(json.dumps(verify_reference(base / "truth" / "ipoteka_reference.xlsx"), ensure_ascii=False, default=str, indent=1))
