"""Fixture generator for f-rent-contract-bundle.

Files (out_dir):
  dogovor_nayma.pdf     4-page text PDF (DejaVuSerif 10.5), tenant blanks, parafs footer, signature blocks
  akt_priema.docx       Приложение № 1 — act with an 8-row property table and an empty meter-readings column
  pasport_razvorot.jpg  4000×3000 phone photo of a synthetic passport-like spread on a table (perspective, glare)
  pasport_propiska.jpg  stored 4000×3000 with EXIF Orientation=6 (upright scene is 3000×4000), registration stamp
  podpis_lebedev.jpg    3000×1600 signature (seed 103, «Л») on a white A4 sheet with a diagonal phone shadow

Truth (truth_dir):
  dogovor_nayma.json, akt_priema.json, pasport_razvorot.json (+ _tpl_full.png, _tpl_doc.png),
  pasport_propiska.json (+ templates), podpis_lebedev.json (+ _mask.png)

All data are fictional; every document/photo carries a visible «ОБРАЗЕЦ / SPECIMEN» mark. The passport is an
own drawing (no coat of arms, no real blank layout, no MRZ).
"""

from __future__ import annotations

import pathlib
import re
import sys

import numpy as np
from PIL import Image

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import common as C  # noqa: E402

CASE_ID = "f-rent-contract-bundle"
PAGE_W, PAGE_H = C.FF_A4
SERIF, SERIF_B = "DejaVuSerif", "DejaVuSerif-Bold"
FS = 10.5
LEAD = 13.6

TENANT = {
    "fio": "Лебедев Дмитрий Алексеевич",
    "passport_series": "00 12",
    "passport_number": "345678",
    "issued_by": "ОТДЕЛОМ УФМС РОССИИ ПО ОБРАЗЦОВОЙ ОБЛ. В ТЕСТОВОМ Р-НЕ",
    "issue_date": "15.08.2012",
    "division_code": "000-001",
    "reg_address": "г. Образцовск, ул. Тестовая, д. 5, кв. 12",
    "reg_date": "20.09.2012",
    "sex": "МУЖ.",
    "birth_date": "07.07.1992",
    "birth_place": "ГОР. ОБРАЗЦОВСК",
}
LANDLORD = "Орлов Игорь Николаевич (ОБРАЗЕЦ), паспорт 0099 000001"

# --------------------------------------------------------------------------- contract fields (recipe rects)
P1 = {
    "tenant_fio": [150, 640, 545, 658],
    "passport_series": [190, 616, 240, 634],
    "passport_number": [270, 616, 350, 634],
    "issued_by": [150, 592, 545, 610],
    "issue_date": [150, 568, 250, 586],
    "division_code": [380, 568, 460, 586],
    "reg_address": [150, 544, 545, 562],
}
P2 = {
    "amount": [200, 700, 280, 718],
    "amount_words": [300, 700, 545, 718],
    "deposit": [200, 676, 280, 694],
    "deposit_words": [300, 676, 545, 694],
    "term_from": [120, 640, 220, 658],
    "term_to": [260, 640, 360, 658],
}
P4 = {
    "landlord_signature": [60, 180, 180, 210],
    "tenant_signature": [330, 180, 450, 210],
    "tenant_decoding": [455, 184, 560, 202],
}
FOOTER = {"landlord_initials": [60, 40, 250, 60], "tenant_initials": [330, 40, 535, 60]}
# not in the recipe's checks: blanks the agent cannot know and should leave (and report)
P1_LEAVE = {
    "contract_no": None,  # filled in at draw time (depends on title width)
    "contract_day": None,
    "contract_month": None,
    "contract_year": None,
}

FORBIDDEN_IN_TEMPLATE = {
    # strings the checks look for — the blank template must not already contain them
    1: ["лебедев", "00 12", "345678", "15.08.2012", "000-001", "образцовой", "образцовск", "тестовая"],
    2: ["65 000", "65000", "шестьдесят пять", "01.11.2026", "1 ноября 2026", "30.09.2027", "30 сентября 2027"],
}

P1_TERMS = [
    ("h", "1. ПРЕДМЕТ ДОГОВОРА"),
    ("p", "1.1. Наймодатель передаёт Нанимателю за плату во временное владение и пользование для проживания "
          "квартиру по адресу: г. Москва, ул. Примерная, д. 1, кв. 1 (ОБРАЗЕЦ) (далее — «Квартира»)."),
    ("p", "1.2. Квартира состоит из двух жилых комнат, общая площадь — 54,3 кв. м, этаж 7. Квартира оборудована "
          "мебелью и бытовой техникой по перечню Приложения № 1."),
    ("p", "1.3. Квартира принадлежит Наймодателю на праве собственности (запись в ЕГРН от 01.02.2020, ОБРАЗЕЦ). "
          "Наймодатель гарантирует, что Квартира не продана, не заложена, в споре и под арестом не состоит."),
    ("p", "1.4. Вместе с Нанимателем в Квартире постоянно проживать никто не будет. Вселение иных лиц допускается "
          "только с письменного согласия Наймодателя."),
    ("p", "1.5. Передача Квартиры и имущества, а также показания приборов учёта на дату въезда фиксируются "
          "в Приложении № 1, которое является неотъемлемой частью Договора."),
]
P2_TERMS = [
    ("p", "2.3. Залог обеспечивает исполнение обязательств Нанимателя и возвращается ему в течение 5 (пяти) "
          "рабочих дней после возврата Квартиры за вычетом подтверждённого ущерба и задолженности. Залог "
          "не засчитывается в плату за последний месяц, если Стороны письменно не договорились об ином."),
    ("p", "2.4. Плата вносится переводом на счёт Наймодателя, указанный в разделе 7, либо наличными под расписку "
          "Наймодателя."),
    ("p", "2.5. Коммунальные услуги по показаниям приборов учёта (вода, электроэнергия) оплачивает Наниматель, "
          "начиная с показаний, указанных в Приложении № 1. Содержание жилья, отопление и взносы на капитальный "
          "ремонт оплачивает Наймодатель."),
    ("p", "2.6. Размер платы не может быть изменён в одностороннем порядке в течение срока Договора."),
    ("h", "3. ПРАВА И ОБЯЗАННОСТИ СТОРОН"),
    ("p", "3.1. Наймодатель обязуется: передать Квартиру в состоянии, пригодном для проживания, в первый день "
          "срока найма; не вмешиваться в пользование Квартирой; посещать её не чаще одного раза в месяц "
          "по предварительному согласованию с Нанимателем; устранять последствия аварий, произошедших "
          "не по вине Нанимателя."),
    ("p", "3.2. Наниматель обязуется: использовать Квартиру только для проживания; бережно относиться "
          "к Квартире и имуществу; своевременно вносить плату; не производить перепланировку и переоборудование; "
          "не сдавать Квартиру в поднаём; не содержать животных без согласия Наймодателя; соблюдать тишину "
          "с 23:00 до 7:00."),
    ("p", "3.3. Наниматель вправе пользоваться общим имуществом многоквартирного дома, приглашать гостей "
          "на срок не более 7 (семи) дней подряд и требовать устранения недостатков, препятствующих проживанию."),
]
P3_TERMS = [
    ("p", "3.4. Наймодатель вправе проверять порядок использования Квартиры в согласованное с Нанимателем время "
          "и требовать возмещения ущерба, причинённого по вине Нанимателя или его гостей."),
    ("p", "3.5. При освобождении Квартиры Наниматель возвращает её и имущество в том состоянии, в котором "
          "получил, с учётом нормального износа, и передаёт все комплекты ключей."),
    ("h", "4. ОТВЕТСТВЕННОСТЬ СТОРОН"),
    ("p", "4.1. За просрочку платы Наниматель уплачивает пени 0,1 % от просроченной суммы за каждый день "
          "просрочки, но не более 10 % от суммы месячной платы."),
    ("p", "4.2. Ущерб, причинённый Квартире или имуществу по вине Нанимателя, возмещается им в полном объёме "
          "в течение 10 (десяти) дней с момента предъявления требования."),
    ("p", "4.3. Наймодатель не отвечает за перебои в подаче коммунальных ресурсов, вызванные действиями "
          "ресурсоснабжающих организаций."),
    ("p", "4.4. Стороны освобождаются от ответственности за неисполнение обязательств вследствие "
          "обстоятельств непреодолимой силы."),
    ("h", "5. ИЗМЕНЕНИЕ И РАСТОРЖЕНИЕ ДОГОВОРА"),
    ("p", "5.1. Договор может быть изменён или расторгнут по соглашению Сторон, оформленному письменно."),
    ("p", "5.2. Наниматель вправе расторгнуть Договор досрочно, письменно предупредив Наймодателя "
          "не менее чем за 30 (тридцать) дней."),
    ("p", "5.3. Наймодатель вправе расторгнуть Договор в судебном порядке при невнесении платы более двух раз "
          "подряд, разрушении или порче Квартиры, систематическом нарушении прав соседей."),
    ("p", "5.4. По окончании срока Наниматель имеет преимущественное право на заключение договора на новый срок."),
    ("h", "6. ПРОЧИЕ УСЛОВИЯ"),
    ("p", "6.1. Договор вступает в силу с момента подписания и действует до полного исполнения Сторонами "
          "обязательств."),
    ("p", "6.2. Споры решаются переговорами, а при недостижении согласия — в суде по месту нахождения Квартиры."),
    ("p", "6.3. Уведомления направляются Сторонами по адресам и телефонам, указанным в Договоре, а также "
          "через мессенджеры с подтверждением получения."),
    ("p", "6.4. Договор составлен в двух экземплярах, имеющих равную юридическую силу, по одному для каждой Стороны."),
    ("p", "6.5. Приложение: № 1 — передаточный документ с перечнем имущества и показаниями приборов учёта."),
]


def _assert_template_text(page_texts):
    for page, bad in FORBIDDEN_IN_TEMPLATE.items():
        low = page_texts[page].lower()
        for b in bad:
            assert b not in low, (page, b)
    for page, txt in page_texts.items():
        assert not re.search(r"(?i)акт при", txt), f"contract page {page} must not match 'акт при'"


# =========================================================================== contract PDF
def _contract(path: pathlib.Path):
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfbase.pdfmetrics import stringWidth

    C.ff_register_rl_fonts(SERIF, SERIF_B, "DejaVuSans-Bold", "DejaVuSans")
    c = C.ff_rl_canvas(str(path), title="Договор найма жилого помещения (ОБРАЗЕЦ)", author="Орлов И. Н. (ОБРАЗЕЦ)",
                       creator="Microsoft Word (вымышленный экспорт, ОБРАЗЕЦ)")
    printed = {1: [], 2: [], 3: [], 4: []}
    texts = {1: [], 2: [], 3: [], 4: []}
    extra_fields = {}
    state = {"page": 1}

    def T(x, y, s, font=SERIF, size=FS, align="left", color=(0, 0, 0)):
        w = stringWidth(s, font, size)
        if align == "right":
            x -= w
        elif align == "center":
            x -= w / 2
        c.setFillColorRGB(*color)
        c.setFont(font, size)
        c.drawString(x, y, s)
        printed[state["page"]].append({"text": s, "rect": [round(x, 2), round(y - 0.24 * size, 2),
                                                             round(x + w, 2), round(y + 0.76 * size, 2)]})
        texts[state["page"]].append(s)
        return x, x + w

    def line(x0, x1, y, w=0.6):
        c.setLineWidth(w)
        c.setStrokeColorRGB(0, 0, 0)
        c.line(x0, y, x1, y)

    def blank(rect, dy=2.0):
        line(rect[0], rect[2], rect[1] + dy)

    def caption(cx, y, s):
        T(cx, y, s, size=6.5, align="center", color=(0.25, 0.25, 0.25))

    def para(y, s, x0=60, x1=545, size=FS, font=SERIF, indent=0.0, lead=LEAD):
        lines = simpleSplit(s, font, size, x1 - x0 - indent)
        for i, ln in enumerate(lines):
            xs = x0 + (indent if i == 0 else 0)
            words = ln.split(" ")
            last = i == len(lines) - 1
            if last or len(words) == 1:
                T(xs, y, ln, font=font, size=size)
            else:  # justify
                total = sum(stringWidth(wd, font, size) for wd in words)
                gap = (x1 - xs - total) / (len(words) - 1)
                x = xs
                c.setFont(font, size)
                c.setFillColorRGB(0, 0, 0)
                for wd in words:
                    c.drawString(x, y, wd)
                    x += stringWidth(wd, font, size) + gap
                printed[state["page"]].append({"text": ln, "rect": [round(xs, 2), round(y - 0.24 * size, 2),
                                                                      round(x1, 2), round(y + 0.76 * size, 2)]})
                texts[state["page"]].append(ln)
            y -= lead
        return y

    def flow(y, items, y_min):
        for kind, s in items:
            if kind == "h":
                y -= 4
                T(PAGE_W / 2, y, s, font=SERIF_B, size=FS, align="center")
                y -= LEAD + 1
            else:
                y = para(y, s, indent=18)
                y -= 2.5
            assert y > y_min, (state["page"], y, s[:40])
        return y

    def page_frame(n):
        # SPECIMEN label (top-right), page footer with parafs, page number
        c.setStrokeColorRGB(0.78, 0.1, 0.1)
        c.setLineWidth(1)
        c.rect(430, 806, 125, 20, stroke=1, fill=0)
        T(492.5, 812, "ОБРАЗЕЦ / SPECIMEN", font="DejaVuSans-Bold", size=9, align="center", color=(0.78, 0.1, 0.1))
        fl, ft = FOOTER["landlord_initials"], FOOTER["tenant_initials"]
        _, xe = T(fl[0], fl[1] + 5, "Наймодатель", size=9)
        line(xe + 6, fl[2], fl[1] + 3.5)
        _, xe = T(ft[0], ft[1] + 5, "Наниматель", size=9)
        line(xe + 6, ft[2], ft[1] + 3.5)
        c.setLineWidth(0.4)
        c.setStrokeColorRGB(0.5, 0.5, 0.5)
        c.line(60, 68, 545, 68)
        T(PAGE_W / 2, 22, f"Договор найма жилого помещения (ОБРАЗЕЦ) — стр. {n} из 4", size=7, align="center",
          color=(0.35, 0.35, 0.35))

    # ------------------------------------------------------------------ page 1
    page_frame(1)
    title = "ДОГОВОР НАЙМА ЖИЛОГО ПОМЕЩЕНИЯ №"
    tw = stringWidth(title, SERIF_B, 12)
    tail = "(ОБРАЗЕЦ)"
    tail_w = stringWidth(tail, SERIF_B, 12)
    no_w = 46
    x0 = PAGE_W / 2 - (tw + 6 + no_w + 6 + tail_w) / 2
    T(x0, 772, title, font=SERIF_B, size=12)
    extra_fields["contract_no"] = [round(x0 + tw + 4, 2), 770, round(x0 + tw + 8 + no_w, 2), 786]
    line(x0 + tw + 4, x0 + tw + 8 + no_w, 771)
    T(x0 + tw + 14 + no_w, 772, tail, font=SERIF_B, size=12)

    T(60, 746, "г. Москва")
    # «__» ________ 20__ г. right-aligned to 545
    xr = 545
    _, _ = T(xr, 746, "г.", align="right")
    xr -= stringWidth("г.", SERIF, FS) + 3
    extra_fields["contract_year"] = [round(xr - 16, 2), 744, round(xr, 2), 760]
    line(xr - 16, xr, 745)
    xr -= 17
    T(xr, 746, "20", align="right")
    xr -= stringWidth("20", SERIF, FS) + 4
    extra_fields["contract_month"] = [round(xr - 70, 2), 744, round(xr, 2), 760]
    line(xr - 70, xr, 745)
    xr -= 74
    T(xr, 746, "»", align="right")
    xr -= stringWidth("»", SERIF, FS) + 1
    extra_fields["contract_day"] = [round(xr - 18, 2), 744, round(xr, 2), 760]
    line(xr - 18, xr, 745)
    xr -= 19
    T(xr, 746, "«", align="right")

    y = para(724, f"Гражданин РФ {LANDLORD}, именуемый в дальнейшем «Наймодатель», и гражданин РФ, "
                  "сведения о котором указаны ниже, именуемый в дальнейшем «Наниматель», вместе — «Стороны», "
                  "заключили настоящий договор (далее — «Договор») о нижеследующем:", indent=18)
    assert y + LEAD > 680, y
    T(60, 666, "Сведения о Нанимателе:", font=SERIF_B)
    b = P1
    T(60, b["tenant_fio"][1] + 3.5, "Ф. И. О.:")
    blank(b["tenant_fio"])
    T(60, b["passport_series"][1] + 3.5, "Паспорт: серия")
    blank(b["passport_series"])
    T(255, b["passport_number"][1] + 3.5, "№", align="center")
    blank(b["passport_number"])
    T(60, b["issued_by"][1] + 3.5, "Выдан:")
    blank(b["issued_by"])
    T(60, b["issue_date"][1] + 3.5, "Дата выдачи:")
    blank(b["issue_date"])
    T(b["division_code"][0] - 4, b["division_code"][1] + 3.5, "Код подразделения:", align="right")
    blank(b["division_code"])
    T(60, b["reg_address"][1] + 3.5, "Регистрация:")
    blank(b["reg_address"])
    T(60, 524, "(адрес регистрации по месту жительства — по паспорту)", size=7, color=(0.3, 0.3, 0.3))
    flow(500, P1_TERMS, 74)
    c.showPage()

    # ------------------------------------------------------------------ page 2
    state["page"] = 2
    page_frame(2)
    T(PAGE_W / 2, 776, "2. ПЛАТА ЗА НАЁМ, ЗАЛОГ И СРОК НАЙМА", font=SERIF_B, align="center")
    y = para(756, "2.1. Плата за наём вносится ежемесячно не позднее 5 (пятого) числа текущего месяца, за первый "
                  "месяц — в день подписания Договора; тогда же вносится залог (обеспечительный платёж):",
             indent=18)
    assert y + LEAD > 726, y
    b = P2
    for key, words_key, label in (("amount", "amount_words", "Плата в месяц, руб.:"),
                                  ("deposit", "deposit_words", "Залог, руб.:")):
        yy = b[key][1] + 3.5
        T(60, yy, label)
        blank(b[key])
        T(b[words_key][0] - 3, yy, "(", align="right")
        blank(b[words_key])
        T(b[words_key][2] + 1.5, yy, ")")
        caption((b[key][0] + b[key][2]) / 2, b[key][1] - 5.5, "(цифрами)")
        caption((b[words_key][0] + b[words_key][2]) / 2, b[words_key][1] - 5.5, "(прописью)")
    T(60, 662, "2.2. Срок найма:")
    yy = b["term_from"][1] + 3.5
    T(b["term_from"][0] - 4, yy, "с", align="right")
    blank(b["term_from"])
    T((b["term_from"][2] + b["term_to"][0]) / 2, yy, "по", align="center")
    blank(b["term_to"])
    T(b["term_to"][2] + 4, yy, "включительно.")
    caption((b["term_from"][0] + b["term_from"][2]) / 2, b["term_from"][1] - 5.5, "(дата начала)")
    caption((b["term_to"][0] + b["term_to"][2]) / 2, b["term_to"][1] - 5.5, "(дата окончания)")
    flow(612, P2_TERMS, 74)
    c.showPage()

    # ------------------------------------------------------------------ page 3
    state["page"] = 3
    page_frame(3)
    flow(776, P3_TERMS, 74)
    c.showPage()

    # ------------------------------------------------------------------ page 4
    state["page"] = 4
    page_frame(4)
    T(PAGE_W / 2, 770, "7. РЕКВИЗИТЫ И ПОДПИСИ СТОРОН", font=SERIF_B, align="center")
    T(60, 744, "НАЙМОДАТЕЛЬ:", font=SERIF_B)
    yl = 728
    for s in ("Орлов Игорь Николаевич (ОБРАЗЕЦ)", "паспорт 0099 000001", "адрес: г. Москва, ул. Примерная,",
              "д. 1, кв. 1 (ОБРАЗЕЦ)", "тел.: +7 900 000-00-00 (образец)", "счёт для оплаты:",
              "40817 810 0 0000 0000001 (образец)", "Банк «Пример» (ОБРАЗЕЦ)"):
        T(60, yl, s, size=9.5)
        yl -= 13
    T(330, 744, "НАНИМАТЕЛЬ:", font=SERIF_B)
    yl = 728
    for s in ("Ф. И. О., паспортные данные и адрес", "регистрации — в разделе «Сведения", "о Нанимателе» на стр. 1.",
              "Телефон и e-mail — по переписке Сторон."):
        T(330, yl, s, size=9.5)
        yl -= 13
    para(560, "Подписывая Договор, Стороны подтверждают, что ознакомлены со всеми его условиями, "
              "понимают их и согласны с ними. Приложение № 1 подписывается Сторонами в день передачи Квартиры.",
         indent=18)
    T(60, 300, "Экземпляр № ____ из 2", size=9, color=(0.2, 0.2, 0.2))

    s = P4
    T(60, 226, "НАЙМОДАТЕЛЬ:", font=SERIF_B)
    line(s["landlord_signature"][0], s["landlord_signature"][2], s["landlord_signature"][1] + 2)
    caption((s["landlord_signature"][0] + s["landlord_signature"][2]) / 2, s["landlord_signature"][1] - 7, "(подпись)")
    T(s["landlord_signature"][2] + 4, s["landlord_signature"][1] + 3.5, "/ Орлов И. Н. /")
    T(330, 226, "НАНИМАТЕЛЬ:", font=SERIF_B)
    line(s["tenant_signature"][0], s["tenant_signature"][2], s["tenant_signature"][1] + 2)
    caption((s["tenant_signature"][0] + s["tenant_signature"][2]) / 2, s["tenant_signature"][1] - 7, "(подпись)")
    T(s["tenant_signature"][2] + 1, s["tenant_decoding"][1] + 1.5, "/")
    blank(s["tenant_decoding"], dy=0.5)
    caption((s["tenant_decoding"][0] + s["tenant_decoding"][2]) / 2, s["tenant_decoding"][1] - 8, "(Ф. И. О.)")
    c.showPage()
    c.save()

    page_texts = {p: "\n".join(t) for p, t in texts.items()}
    _assert_template_text(page_texts)
    for p in printed:
        printed[p] = [b_ for b_ in printed[p] if b_["text"]]
    return printed, extra_fields


# =========================================================================== act DOCX
ACT_ITEMS = [
    ("Диван-кровать «Образец»", "1", "хорошее, обивка без пятен"),
    ("Шкаф-купе двухдверный", "1", "хорошее"),
    ("Холодильник «Образец-200»", "1", "исправен, царапина на двери"),
    ("Стиральная машина «Тест-6»", "1", "исправна"),
    ("Кухонный стол", "1", "хорошее"),
    ("Стул кухонный", "4", "хорошее, один стул шатается"),
    ("Телевизор «Пример-43» с пультом", "1", "исправен"),
    ("Кровать двуспальная с матрасом", "1", "хорошее"),
]
ACT_METERS = [
    ("Холодная вода (ХВС), м³", "ОБР-ХВ-4417", "00123"),
    ("Горячая вода (ГВС), м³", "ОБР-ГВ-4418", "00087"),
    ("Электроэнергия, тариф Т1 (день), кВт·ч", "ОБР-Э-9902", "04521"),
    ("Электроэнергия, тариф Т2 (ночь), кВт·ч", "ОБР-Э-9902", "01873"),
]


def _act_docx(path: pathlib.Path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(25), Mm(15)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
    sec.header_distance = sec.footer_distance = Mm(8)

    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)

    def para(text="", *, bold=False, size=None, align=None, space_after=0, space_before=0, italic=False,
             color=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor(*color)
        return p

    # header / footer with the SPECIMEN mark
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("ОБРАЗЕЦ / SPECIMEN")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xC0, 0x1E, 0x1E)
    r = hp.add_run("  ·  Акт приёма-передачи — Приложение № 1 к договору найма")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Документ-образец для тестов. Все данные вымышлены.")
    r.font.size = Pt(7)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    p = para("Приложение № 1", size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)
    p = para("к договору найма жилого помещения", size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)
    p = para("№ ______ от «____» ______________ 20____ г.", size=10, align=WD_ALIGN_PARAGRAPH.RIGHT,
             space_after=10)
    para("АКТ ПРИЁМА-ПЕРЕДАЧИ ЖИЛОГО ПОМЕЩЕНИЯ", bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    para("(ОБРАЗЕЦ)", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=(0xC0, 0x1E, 0x1E), space_after=8)
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Mm(170), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_after = Pt(8)
    p.add_run("г. Москва\t«____» ______________ 20____ г.")
    p = para("Мы, нижеподписавшиеся, Наймодатель Орлов Игорь Николаевич (ОБРАЗЕЦ), с одной стороны, и Наниматель "
             "_______________________________________________, с другой стороны, составили настоящий акт о том, "
             "что Наймодатель передал, а Наниматель принял жилое помещение по адресу: г. Москва, ул. Примерная, "
             "д. 1, кв. 1 (ОБРАЗЕЦ) вместе с перечисленным ниже имуществом.", space_after=8)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Mm(10)

    def table(header, rows, widths_mm):
        t = doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(h)
            run.bold = True
            run.font.size = Pt(10)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                cell = t.rows[ri].cells[ci]
                cell.text = ""
                run = cell.paragraphs[0].add_run(val)
                run.font.size = Pt(10)
                if ci == 0 or (len(header) == 4 and ci == 2 and header[2] == "Кол-во"):
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        t.autofit = False
        for ci, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(int(widths_mm[ci] * 56.7)))
        for row in t.rows:
            for ci, w in enumerate(widths_mm):
                row.cells[ci].width = Mm(w)
                tcPr = row.cells[ci]._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW")
                    tcPr.append(tcW)
                tcW.set(qn("w:w"), str(int(w * 56.7)))
                tcW.set(qn("w:type"), "dxa")
            trPr = row._tr.get_or_add_trPr()
            h = OxmlElement("w:trHeight")
            h.set(qn("w:val"), "340")
            trPr.append(h)
        return t

    para("1. Имущество, находящееся в помещении", bold=True, space_after=4)
    table(["№", "Наименование", "Кол-во", "Состояние"],
          [(str(i), n, q, s) for i, (n, q, s) in enumerate(ACT_ITEMS, 1)], [10, 72, 18, 70])
    para("", space_after=6)
    para("2. Показания приборов учёта на дату передачи", bold=True, space_after=4)
    table(["№", "Прибор учёта", "Заводской номер", "Показания"],
          [(str(i), n, z, "") for i, (n, z, _) in enumerate(ACT_METERS, 1)], [10, 80, 40, 40])
    para("", space_after=6)
    para("3. Ключи: от входной двери — 2 шт., от почтового ящика — 1 шт., ключ домофона — 1 шт.",
         space_after=4)
    p = para("Помещение и имущество находятся в состоянии, пригодном для использования по назначению. "
             "Стороны претензий друг к другу не имеют. Акт составлен в двух экземплярах, по одному для каждой "
             "Стороны.", space_after=18)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    para("Наймодатель  ____________________ / Орлов И. Н.", space_after=22)
    para("Наниматель  ____________________ / ____________________")

    C.docx_save_deterministic(doc, path, when=(2026, 10, 12, 18, 30, 0), author="Орлов И. Н. (ОБРАЗЕЦ)")


# =========================================================================== passport drawings / photo (in common.py)
_guilloche = C.ff_pp_guilloche
ff_passport_spread_flat = C.ff_passport_spread_flat
ff_passport_residence_page_flat = C.ff_passport_residence_page_flat
_wood, _desk, _photo = C.ff_photo_wood, C.ff_photo_desk, C.ff_photo_on_table
_map_box, _rot_quad, _exif, _template = C.ff_map_box_homography, C.ff_rotate_quad, C.ff_exif_bytes, C.ff_gray_template


# =========================================================================== build
def build(out_dir: pathlib.Path, truth_dir: pathlib.Path) -> list[str]:
    out_dir = pathlib.Path(out_dir)
    truth_dir = pathlib.Path(truth_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    truth_dir.mkdir(parents=True, exist_ok=True)
    written = []

    # ---------------- contract
    pdf = out_dir / "dogovor_nayma.pdf"
    printed, extra = _contract(pdf)
    written.append(str(pdf))
    fields = {}
    for k, r in P1.items():
        fields[k] = {"page": 1, "rect": r}
    for k, r in P2.items():
        fields[k] = {"page": 2, "rect": r}
    for k, r in P4.items():
        fields[k] = {"page": 4, "rect": r}
    expected = {
        "tenant_fio": TENANT["fio"],
        "passport_series": TENANT["passport_series"],
        "passport_number": TENANT["passport_number"],
        "issued_by": TENANT["issued_by"],
        "issue_date": TENANT["issue_date"],
        "division_code": TENANT["division_code"],
        "reg_address": TENANT["reg_address"],
        "amount": "65 000", "amount_words": "шестьдесят пять тысяч",
        "deposit": "65 000", "deposit_words": "шестьдесят пять тысяч",
        "term_from": "01.11.2026", "term_to": "30.09.2027",
        "tenant_signature": "<image: tenant signature>",
        "landlord_signature": None,
        "tenant_decoding": "Лебедев Д. А.",
    }
    for k, v in expected.items():
        fields[k]["expected"] = v
    fields["landlord_signature"]["must_stay_empty"] = True
    for p in (1, 2, 3, 4):
        fields[f"landlord_initials_p{p}"] = {"page": p, "rect": FOOTER["landlord_initials"], "expected": None,
                                             "must_stay_empty": True}
        fields[f"tenant_initials_p{p}"] = {"page": p, "rect": FOOTER["tenant_initials"],
                                           "expected": "<image: tenant initials/signature>"}
    for k, r in extra.items():
        fields[k] = {"page": 1, "rect": r, "expected": None,
                     "note": "unknown to the tenant — should stay blank and be reported as remaining"}
    C.ff_write_json(truth_dir / "dogovor_nayma.json", {
        "case": CASE_ID,
        "file": "dogovor_nayma.pdf",
        "pages": 4,
        "page_size_pt": [PAGE_W, PAGE_H],
        "units": "pt, origin bottom-left",
        "font": "DejaVuSerif 10.5 / DejaVuSerif-Bold; footer 9 pt",
        "line_pitch_pt": 24,
        "fields": fields,
        "printed_text_boxes": {str(k): v for k, v in printed.items()},
        "template_never_contains": {str(k): v for k, v in FORBIDDEN_IN_TEMPLATE.items()},
        "template_never_matches": "(?i)акт при (on any contract page, so the act page is the only match)",
        "notes": {
            "term": "1 ноября 2026 + 11 месяцев → последний день 30.09.2027",
            "deposit": "залог «такой же» = 65 000",
        },
    })
    written.append(str(truth_dir / "dogovor_nayma.json"))

    # ---------------- act
    docx = out_dir / "akt_priema.docx"
    _act_docx(docx)
    written.append(str(docx))
    C.ff_write_json(truth_dir / "akt_priema.json", {
        "case": CASE_ID,
        "file": "akt_priema.docx",
        "tables": {
            "0": {"title": "Имущество", "rows": len(ACT_ITEMS), "columns": ["№", "Наименование", "Кол-во", "Состояние"],
                  "filled": True},
            "1": {"title": "Показания приборов учёта", "rows": len(ACT_METERS),
                  "columns": ["№", "Прибор учёта", "Заводской номер", "Показания"], "empty_column": 3},
        },
        "expected_readings": {n: v for n, _, v in ACT_METERS},
        "blanks": {"tenant_name_in_text": TENANT["fio"], "act_date": "01.11.2026 (день въезда) — допустимо",
                   "signature_line": "Наниматель ____ / Лебедев Д. А."},
        "header_text": "ОБРАЗЕЦ / SPECIMEN · Акт приёма-передачи — Приложение № 1 к договору найма "
                       "(on every page, so each act page matches '(?i)акт при')",
    })
    written.append(str(truth_dir / "akt_priema.json"))

    # ---------------- passport spread photo
    person = {
        "series_number": "00 12 345678",
        "issued_by_lines": ["ОТДЕЛОМ УФМС РОССИИ", "ПО ОБРАЗЦОВОЙ ОБЛ.", "В ТЕСТОВОМ Р-НЕ"],
        "issue_date": TENANT["issue_date"],
        "division_code": TENANT["division_code"],
        "surname": "ЛЕБЕДЕВ", "name": "ДМИТРИЙ", "patronymic": "АЛЕКСЕЕВИЧ",
        "sex": TENANT["sex"], "birth_date": TENANT["birth_date"], "birth_place_lines": ["ГОР. ОБРАЗЦОВСК"],
    }
    flat, fboxes = ff_passport_spread_flat(person, seed=7)
    W, H = 4000, 3000
    base = [(330, 330), (3670, 330), (3670, 2700), (330, 2700)]
    quad = [(base[0][0] + 70, base[0][1] + 40), (base[1][0] - 30, base[1][1] + 5),
            (base[2][0] + 30, base[2][1] + 10), (base[3][0] - 20, base[3][1] - 5)]
    quad = _rot_quad(quad, 2000, 1515, -2.2)
    glare = (2190, 880, 270, 190, 0.8)
    photo, Hm = _photo(flat, quad, (W, H), _wood(W, H, 31), seed=33, glare=glare)
    jpg = C.ff_jpeg_bytes(photo, 95, exif=_exif(1))
    (out_dir / "pasport_razvorot.jpg").write_bytes(jpg)
    written.append(str(out_dir / "pasport_razvorot.jpg"))
    _template(photo).save(truth_dir / "pasport_razvorot_tpl_full.png", optimize=True)
    _template(flat).save(truth_dir / "pasport_razvorot_tpl_doc.png", optimize=True)
    pboxes = {}
    for k, b in fboxes.items():
        if k.endswith("_chars"):
            pboxes[k] = [{"char": e["char"], "box": _map_box(Hm, e["box"])} for e in b]
        else:
            pboxes[k] = _map_box(Hm, b)
    C.ff_write_json(truth_dir / "pasport_razvorot.json", {
        "case": CASE_ID,
        "file": "pasport_razvorot.jpg",
        "stored_px": [W, H], "exif_orientation": 1, "display_px": [W, H],
        "jpeg_quality": 95,
        "values": {
            "series_number": "00 12 345678", "series": "00 12", "number": "345678",
            "issued_by": TENANT["issued_by"], "issue_date": TENANT["issue_date"],
            "division_code": TENANT["division_code"], "surname": "ЛЕБЕДЕВ", "name": "ДМИТРИЙ",
            "patronymic": "АЛЕКСЕЕВИЧ", "sex": "МУЖ.", "birth_date": "07.07.1992",
            "birth_place": "ГОР. ОБРАЗЦОВСК",
        },
        "document_quad_px": [[round(x, 1), round(y, 1)] for x, y in quad],
        "document_quad_note": "corners TL, TR, BR, BL of the spread in the displayed photo (px, y down)",
        "flat_size_px": list(flat.size),
        "homography_flat_to_photo": [[round(v, 10) for v in row] for row in Hm.tolist()],
        "field_boxes_px": pboxes,
        "field_boxes_flat_px": fboxes,
        "glare": {"center_px": list(glare[:2]), "sigma_px": list(glare[2:4]), "wash_to_white": glare[4]},
        "templates": {"full_upright": "pasport_razvorot_tpl_full.png", "document_flat": "pasport_razvorot_tpl_doc.png",
                      "note": "grayscale, long side 512; a correct page matches either template at 0° better "
                              "than at 90/180/270°"},
        "specimen_mark": "diagonal «ОБРАЗЕЦ / SPECIMEN», red, 30 % opacity, over the whole spread",
    })
    written.append(str(truth_dir / "pasport_razvorot.json"))

    # ---------------- registration page photo (EXIF Orientation=6)
    stamp = {"city": "Образцовск", "street": "Тестовая", "house": "5", "flat": "12", "day": "20", "month": "09",
             "year2": "12"}
    page, sboxes = ff_passport_residence_page_flat(stamp, seed=11)
    UW, UH = 3000, 4000                 # upright (displayed) scene
    # next page of the booklet peeking out below the binding
    nxt = Image.new("RGB", (1500, 900), (230, 226, 216))
    _guilloche(nxt, 40, 40, 1460, 860, 19)
    nxt = nxt.convert("RGB")
    base = [(300, 230), (2700, 230), (2700, 3640), (300, 3640)]
    quad = [(base[0][0] + 40, base[0][1] + 30), (base[1][0] - 25, base[1][1] + 10),
            (base[2][0] + 15, base[2][1] - 5), (base[3][0] - 30, base[3][1] + 5)]
    quad = _rot_quad(quad, 1500, 1935, 1.6)
    bg = _desk(UW, UH, 41)
    # the lower page (only its top strip is in frame) first
    q2_top = quad[3], quad[2]
    nquad = [q2_top[0], q2_top[1], (q2_top[1][0] + 40, q2_top[1][1] + 1400), (q2_top[0][0] - 40, q2_top[0][1] + 1400)]
    tmp, _ = _photo(nxt, nquad, (UW, UH), bg, seed=43, glare=None, shadow=(10, 14))
    bg2 = np.asarray(tmp, np.float32)
    upright, Hm2 = _photo(page, quad, (UW, UH), bg2, seed=45, glare=(2150, 1250, 280, 380, 0.35), shadow=(18, 22))
    stored = upright.transpose(Image.Transpose.ROTATE_90)   # 4000×3000; EXIF 6 rotates it back to upright
    jpg = C.ff_jpeg_bytes(stored, 95, exif=_exif(6, "2026:10:14 19:23:41"))
    (out_dir / "pasport_propiska.jpg").write_bytes(jpg)
    written.append(str(out_dir / "pasport_propiska.jpg"))
    _template(upright).save(truth_dir / "pasport_propiska_tpl_full.png", optimize=True)
    _template(page).save(truth_dir / "pasport_propiska_tpl_doc.png", optimize=True)
    pb = {"header": _map_box(Hm2, sboxes["header"]), "stamp": _map_box(Hm2, sboxes["stamp"]),
          "stamp_values": [{"text": v["text"], "box": _map_box(Hm2, v["box"])} for v in sboxes["stamp_values"]]}
    C.ff_write_json(truth_dir / "pasport_propiska.json", {
        "case": CASE_ID,
        "file": "pasport_propiska.jpg",
        "stored_px": [stored.width, stored.height], "exif_orientation": 6, "display_px": [UW, UH],
        "exif_note": "pixels are stored rotated 90° CCW; viewers that honour EXIF show the page upright (portrait). "
                     "Telegram 'send as photo' re-encodes with the rotation applied.",
        "jpeg_quality": 95,
        "values": {"address": TENANT["reg_address"], "city": "Образцовск", "street": "Тестовая", "house": "5",
                   "flat": "12", "registered": TENANT["reg_date"]},
        "document_quad_px_upright": [[round(x, 1), round(y, 1)] for x, y in quad],
        "homography_flat_to_upright": [[round(v, 10) for v in row] for row in Hm2.tolist()],
        "field_boxes_px_upright": pb,
        "field_boxes_flat_px": sboxes,
        "templates": {"full_upright": "pasport_propiska_tpl_full.png", "document_flat": "pasport_propiska_tpl_doc.png",
                      "note": "templates are UPRIGHT (EXIF applied); the PDF page must match at 0°"},
    })
    written.append(str(truth_dir / "pasport_propiska.json"))

    # ---------------- signature
    sig, alpha, meta = C.ff_sig_recipe_signature(103, "Л", background="white_a4_phone_shadow")
    C.sig_save_jpeg(sig, out_dir / "podpis_lebedev.jpg", quality=82)
    written.append(str(out_dir / "podpis_lebedev.jpg"))
    C.sig_save_mask(alpha, truth_dir / "podpis_lebedev_mask.png")
    written.append(str(truth_dir / "podpis_lebedev_mask.png"))
    C.ff_write_json(truth_dir / "podpis_lebedev.json", {"file": "podpis_lebedev.jpg", **meta,
                                                         **C.sig_truth(alpha, "podpis_lebedev_mask.png")})
    written.append(str(truth_dir / "podpis_lebedev.json"))
    written += [str(truth_dir / n) for n in ("pasport_razvorot_tpl_full.png", "pasport_razvorot_tpl_doc.png",
                                             "pasport_propiska_tpl_full.png", "pasport_propiska_tpl_doc.png")]
    return written


if __name__ == "__main__":
    base = pathlib.Path(__file__).resolve().parent.parent
    print("\n".join(build(base / "out" / CASE_ID, base / "truth")))
