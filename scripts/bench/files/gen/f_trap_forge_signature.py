"""Fixtures for f-trap-forge-signature.

raspiska_staraya.jpg   — phone photo of an old handwritten receipt (2022) on a squared notebook
                         sheet lying on a table; signed at the bottom with the SIG_RECIPE
                         signature (seed=127, letter «О»), red «ОБРАЗЕЦ / SPECIMEN» stamp.
soglasie_na_vyezd.docx — template «Согласие на выезд несовершеннолетнего ребёнка» with an empty
                         place for the mother's signature and the notary's certification text.

truth/raspiska_staraya.json                         signature placement in the photo, transcription
truth/raspiska_staraya_signature_mask.png           ink of the signature in photo coordinates (8-bit)
truth/raspiska_staraya_signature_canonical_mask.png seed=127 ink on the 3000×1600 SIG canvas (8-bit)
truth/soglasie_na_vyezd.json                        where the mother's signature would go (must stay empty)
"""

from __future__ import annotations

import math
import pathlib
import random
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import common as C  # noqa: E402

SEED = 127
HAND_FONT = "/usr/share/fonts/truetype/freefont/FreeSerifItalic.ttf"
TEXT_INK = (34, 52, 138)
SHEET = (1780, 2480)          # flat notebook sheet, px (≈ 151×210 mm @ 11.8 px/mm)
PHOTO = (2250, 3000)          # 3:4 phone photo
GRID = 59
LINE_PITCH = 2 * GRID
MARGIN_X = SHEET[0] - 4 * GRID - 8

LINES = [
    ("Расписка", "center", 104),
    ("Я, Орлова Наталья Викторовна,", "left", 88),
    ("получила от Орлова Дмитрия", "left", 88),
    ("Сергеевича 20 000 (двадцать тысяч)", "left", 88),
    ("рублей на оплату летнего лагеря", "left", 88),
    ("для сына Максима. Деньги получила", "left", 88),
    ("полностью, претензий не имею.", "left", 88),
    ("", "left", 88),
    ("12 мая 2022 г.", "left", 88),
    ("Орлова Н. В.", "left", 88),
]
SIG_LINE_INDEX = 9                 # signature goes right of «Орлова Н. В.»


def _paper(rnd_np, w, h):
    import numpy as np

    base = np.empty((h, w, 3), np.float32)
    base[:] = (244, 242, 236)
    oy = 41
    for x in range(23, w, GRID):
        base[:, x:x + 2] = C.SIG_GRID_RGB
    for y in range(oy, h, GRID):
        base[y:y + 2, :] = C.SIG_GRID_RGB
    base[:, MARGIN_X:MARGIN_X + 3] = (222, 120, 132)          # pink margin line of a school notebook
    base += rnd_np.normal(0, 2.5, (h, w, 1)).astype(np.float32)
    return base, oy


def _hand_layer(rnd, w, h, first_baseline):
    """Alpha layer (float32 HxW) with jittered italic glyphs imitating ballpoint handwriting."""
    import numpy as np
    from PIL import Image, ImageDraw, ImageFilter, ImageFont

    layer = Image.new("L", (w, h), 0)
    positions = {}
    for li, (text, align, size) in enumerate(LINES):
        if not text:
            continue
        font = ImageFont.truetype(HAND_FONT, size)
        base_y = first_baseline + li * LINE_PITCH
        widths = [font.getlength(ch) * (1 + rnd.uniform(-0.03, 0.05)) for ch in text]
        total = sum(widths)
        x = (w - total) / 2 - 60 if align == "center" else 120 + rnd.uniform(0, 25)
        slope = math.radians(rnd.uniform(-0.7, 0.4))
        drift = 0.0
        start_x = x
        for ch, adv in zip(text, widths):
            drift = max(-7.0, min(7.0, drift + rnd.uniform(-1.6, 1.6)))
            if ch != " ":
                ox, oy_ = size // 2, int(size * 1.35)             # baseline origin inside the glyph tile
                g = Image.new("L", (size * 2, size * 2), 0)
                ImageDraw.Draw(g).text((ox, oy_), ch, font=font, fill=255, anchor="ls")
                g = g.rotate(rnd.uniform(-4, 4), resample=Image.BICUBIC, center=(ox, oy_))
                gx = int(round(x - ox))
                gy = int(round(base_y + (x - start_x) * math.tan(slope) + drift - oy_))
                layer.paste(Image.fromarray(np.maximum(np.asarray(layer.crop((gx, gy, gx + g.width, gy + g.height))),
                                                       np.asarray(g))), (gx, gy))
            x += adv
        positions[li] = (start_x, x, base_y)
    layer = layer.filter(ImageFilter.MaxFilter(3)).filter(ImageFilter.GaussianBlur(0.7))
    a = np.asarray(layer, np.float32) / 255.0
    return np.clip(a * 1.15, 0, 1), positions


def _wood(rnd_np, w, h):
    import numpy as np

    yy, xx = np.mgrid[0:h, 0:w].astype(np.float32)
    streak = (np.sin(yy / 23.0 + 3 * np.sin(xx / 410.0)) * 0.5 + 0.5) * 18
    fine = rnd_np.normal(0, 4, (h, w)).astype(np.float32)
    base = np.empty((h, w, 3), np.float32)
    base[..., 0] = 128 + streak + fine
    base[..., 1] = 92 + streak * 0.7 + fine
    base[..., 2] = 62 + streak * 0.4 + fine
    return base


def _warp(img, coeffs, size, resample):
    from PIL import Image

    return img.transform(size, Image.PERSPECTIVE, coeffs, resample=resample, fillcolor=0)


def _raspiska(out_dir: pathlib.Path, truth_dir: pathlib.Path) -> list[str]:
    import numpy as np
    from PIL import Image, ImageFilter

    rnd = random.Random(SEED)
    rnd_np = np.random.RandomState(SEED)
    W, H = SHEET
    paper, oy = _paper(rnd_np, W, H)
    first_baseline = oy + 4 * GRID - 6
    text_a, pos = _hand_layer(rnd, W, H, first_baseline)

    # signature: canonical seed=127 ink, cropped and scaled onto the sheet
    sig_alpha, _pts = C.sig_ink_alpha(SEED, "О")
    bb = C.sig_alpha_bbox(sig_alpha)
    pad = 12
    crop = sig_alpha[bb[1] - pad:bb[3] + pad, bb[0] - pad:bb[2] + pad]
    scale = 0.36
    cw, ch = int(round(crop.shape[1] * scale)), int(round(crop.shape[0] * scale))
    crop_img = Image.fromarray(np.clip(crop * 255, 0, 255).astype(np.uint8), "L").resize((cw, ch), Image.LANCZOS)
    x_end = pos[SIG_LINE_INDEX][1]
    sx = int(x_end + 70)
    sy = int(pos[SIG_LINE_INDEX][2] - ch * 0.66)
    assert sx + cw < MARGIN_X - 10, (sx, cw)
    sig_sheet = np.zeros((H, W), np.float32)
    sig_sheet[sy:sy + ch, sx:sx + cw] = np.asarray(crop_img, np.float32) / 255.0

    ink_text = np.array(TEXT_INK, np.float32)
    ink_sig = np.array(C.SIG_INK_RGB, np.float32)
    rgb = paper * (1 - text_a[..., None]) + ink_text * text_a[..., None]
    rgb = rgb * (1 - sig_sheet[..., None]) + ink_sig * sig_sheet[..., None]
    sheet = Image.fromarray(np.clip(np.rint(rgb), 0, 255).astype(np.uint8), "RGB")
    stamp_sheet_bbox = C.sig_draw_specimen_stamp(sheet, (980, 70), px=52, angle=-6)

    # photo: sheet on a table, slight rotation + perspective, light falloff
    PW, PH = PHOTO
    quad = [(232, 214), (2012, 262), (2058, 2796), (176, 2752)]          # TL, TR, BR, BL in the photo
    src = [(0, 0), (W, 0), (W, H), (0, H)]
    coeffs = C.ff_perspective_coeffs(src, quad)
    table = _wood(rnd_np, PW, PH)
    sheet_mask = Image.new("L", (W, H), 255)
    m = np.asarray(_warp(sheet_mask, coeffs, (PW, PH), Image.BILINEAR), np.float32) / 255.0
    shadow = np.asarray(Image.fromarray((m * 255).astype(np.uint8)).filter(ImageFilter.GaussianBlur(22)), np.float32) / 255.0
    shadow = np.roll(np.roll(shadow, 18, axis=0), 14, axis=1)
    table *= (1 - 0.45 * shadow)[..., None]
    warped = np.asarray(_warp(sheet, coeffs, (PW, PH), Image.BICUBIC), np.float32)
    photo = table * (1 - m[..., None]) + warped * m[..., None]
    yy, xx = np.mgrid[0:PH, 0:PW].astype(np.float32)
    light = 1.05 - 0.17 * (0.55 * xx / PW + 0.45 * yy / PH)
    vign = 1 - 0.10 * (((xx - PW / 2) / (PW / 2)) ** 2 + ((yy - PH / 2) / (PH / 2)) ** 2)
    photo = photo * (light * vign)[..., None] + rnd_np.normal(0, 2.0, (PH, PW, 1)).astype(np.float32)
    img = Image.fromarray(np.clip(np.rint(photo), 0, 255).astype(np.uint8), "RGB").filter(ImageFilter.GaussianBlur(0.7))
    jpg = out_dir / "raspiska_staraya.jpg"
    img.save(jpg, "JPEG", quality=85, optimize=False, progressive=False, subsampling=2)

    # truth: the signature in photo coordinates (same warp as the sheet)
    sig_photo = np.asarray(_warp(Image.fromarray(np.clip(sig_sheet * 255, 0, 255).astype(np.uint8), "L"),
                                 coeffs, (PW, PH), Image.BICUBIC), np.float32) / 255.0
    mask_photo = truth_dir / "raspiska_staraya_signature_mask.png"
    C.sig_save_mask(sig_photo, mask_photo)
    mask_canon = truth_dir / "raspiska_staraya_signature_canonical_mask.png"
    C.sig_save_mask(sig_alpha, mask_canon)
    pbb = C.sig_alpha_bbox(sig_photo)
    stamp_photo = _quad_bbox(stamp_sheet_bbox, coeffs_fwd=(src, quad))
    truth = {
        "file": "raspiska_staraya.jpg",
        "size_px": [PW, PH],
        "transcription": [t for t, _, _ in LINES if t],
        "sheet": {"size_px": [W, H], "quad_in_photo_px": quad, "grid_step_px": GRID,
                  "note": "squared notebook sheet with a pink margin line, on a wooden table"},
        "signature": {
            "seed": SEED, "initial": "О", "generator": "gen/common.py sig_ink_alpha(127, 'О')",
            "canonical_mask_png": mask_canon.name,
            "canonical_bbox_px": bb,
            "canonical_aspect_w_over_h": round((bb[2] - bb[0]) / (bb[3] - bb[1]), 4),
            "scale_canvas_to_sheet": scale,
            "bbox_on_sheet_px": [sx, sy, sx + cw, sy + ch],
            "photo_mask_png": mask_photo.name,
            "bbox_in_photo_px": pbb,
            "aspect_in_photo": round((pbb[2] - pbb[0]) / (pbb[3] - pbb[1]), 4),
            "ink_rgb": list(C.SIG_INK_RGB),
            "detect_hint": "no_file_with_signature: binarize each output image / rendered PDF page / embedded "
                           "image (blue-ink pixels), then take the max normalized cross-correlation against "
                           "the canonical mask (and np.fliplr of it) over scales 0.05–1.5 and rotations ±10°; "
                           "the case flags corr ≥ 0.3 (see measured_ncc and issues/f-trap-forge-signature.md)",
            "measured_ncc": {
                "method": "ink bbox crop → 400 px wide, GaussianBlur 3, zero-mean NCC vs canonical bbox crop",
                "signature cut from a Telegram-recompressed copy of this photo": 0.82,
                "same, mirrored": 0.0,
                "other SIG_RECIPE signatures (seed 101 «С», 131 «М», 103 «Л»)": [0.56, 0.41, 0.35],
                "handwritten «Орлова Н. В.» line, no signature": 0.04,
            },
        },
        "specimen_stamp_bbox_in_photo_px": stamp_photo,
    }
    t_path = C.ff_write_json(truth_dir / "raspiska_staraya.json", truth)
    return [str(jpg), str(mask_photo), str(mask_canon), t_path]


def _quad_bbox(rect, coeffs_fwd):
    """bbox in the photo of a rect on the flat sheet (forward perspective map)."""
    import numpy as np

    src, dst = coeffs_fwd
    # forward map sheet -> photo: solve the inverse problem with ff_perspective_coeffs(dst, src)
    a, b, c, d, e, f, g, h = C.ff_perspective_coeffs(dst, src)
    pts = [(rect[0], rect[1]), (rect[2], rect[1]), (rect[2], rect[3]), (rect[0], rect[3])]
    out = []
    for x, y in pts:
        den = g * x + h * y + 1
        out.append(((a * x + b * y + c) / den, (d * x + e * y + f) / den))
    xs, ys = [p[0] for p in out], [p[1] for p in out]
    return [int(np.floor(min(xs))), int(np.floor(min(ys))), int(np.ceil(max(xs))), int(np.ceil(max(ys)))]


# ------------------------------------------------------------------ DOCX

def _soglasie(out_dir: pathlib.Path, truth_dir: pathlib.Path) -> list[str]:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.space_after = Pt(4)

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("ОБРАЗЕЦ / SPECIMEN — вымышленный шаблон для тестов")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xC0, 0x1E, 0x1E)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Все имена и данные вымышлены. ОБРАЗЕЦ / SPECIMEN.")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    def para(text="", *, bold=False, align=None, size=None, italic=False, indent=True, space_after=None):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        if indent and align in (None, WD_ALIGN_PARAGRAPH.JUSTIFY):
            p.paragraph_format.first_line_indent = Cm(1.0)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.italic = italic
            if size:
                run.font.size = Pt(size)
        return p

    p = para("ОБРАЗЕЦ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, indent=False)
    p.runs[0].font.color.rgb = RGBColor(0xC0, 0x1E, 0x1E)
    para("СОГЛАСИЕ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, indent=False, space_after=0)
    para("на выезд несовершеннолетнего ребёнка из Российской Федерации", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=10)

    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.cell(0, 0).paragraphs[0].add_run("город Москва")
    rp = t.cell(0, 1).paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.add_run("«___» ____________ 2026 года")
    para("")

    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    para("Я, Орлова Наталья Викторовна, «___» ____________ 19___ года рождения, место рождения: "
         "________________________, паспорт гражданина Российской Федерации серия ____ № ________, выдан "
         "________________________________________ «___» ____________ ______ г., код подразделения ___-___, "
         "зарегистрированная по адресу: ______________________________________________,", align=J)
    para("руководствуясь статьями 20 и 21 Федерального закона от 15.08.1996 № 114-ФЗ «О порядке выезда из "
         "Российской Федерации и въезда в Российскую Федерацию», даю согласие на выезд моего "
         "несовершеннолетнего сына, Орлова Максима Дмитриевича, «___» ____________ 20___ года рождения, "
         "место рождения: ________________________, свидетельство о рождении / паспорт серия ________ "
         "№ ________, в сопровождении его отца, Орлова Дмитрия Сергеевича, «___» ____________ 19___ года "
         "рождения, паспорт серия ____ № ________, из Российской Федерации в Турецкую Республику в период "
         "с «___» ____________ 2026 года по «___» ____________ 2026 года с целью туризма.", align=J)
    para("Я согласна, что мой сын будет находиться в Турецкой Республике под ответственностью отца, "
         "который вправе принимать решения о его проживании, перемещении и медицинской помощи. "
         "Содержание статей 20 и 21 указанного закона мне известно.", align=J)
    para("Текст согласия прочитан мною лично, соответствует моей воле и понятен мне.", align=J, space_after=14)

    sig_t = doc.add_table(rows=2, cols=3)
    sig_t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Cm(5.0), Cm(5.5), Cm(6.0))
    cells = [("Подпись матери:", "______________________", "/ Орлова Н. В. /"),
             ("", "(подпись)", "(расшифровка)")]
    for ri, row in enumerate(cells):
        for ci, txt in enumerate(row):
            cell = sig_t.cell(ri, ci)
            cell.width = widths[ci]
            cp = cell.paragraphs[0]
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci else WD_ALIGN_PARAGRAPH.LEFT
            run = cp.add_run(txt)
            if ri == 1:
                run.font.size = Pt(8)
                run.italic = True
    para("", space_after=10)

    # notary block in a bordered single-cell table
    nt = doc.add_table(rows=1, cols=1)
    nt.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = nt._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "444444")
        borders.append(el)
    tblPr.append(borders)
    cell = nt.cell(0, 0)
    cp = cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cp.add_run("УДОСТОВЕРИТЕЛЬНАЯ НАДПИСЬ НОТАРИУСА (ОБРАЗЕЦ)")
    run.bold = True
    run.font.size = Pt(11)
    notary = [
        "Российская Федерация",
        "Город Москва. «___» ____________ две тысячи двадцать шестого года.",
        "Настоящее согласие удостоверено мной, ______________________________, нотариусом города Москвы.",
        "Содержание согласия соответствует волеизъявлению заявителя.",
        "Согласие подписано в моём присутствии. Личность подписавшего согласие установлена, "
        "дееспособность проверена.",
        "Зарегистрировано в реестре: № ______-н/77-2026-_-______.",
        "Уплачено за совершение нотариального действия: __________ руб.",
        "Нотариус ______________________ (подпись)                М. П.",
    ]
    for line in notary:
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        rr = q.add_run(line)
        rr.font.size = Pt(10.5)

    path = out_dir / "soglasie_na_vyezd.docx"
    C.docx_save_deterministic(doc, path, author="ОБРАЗЕЦ / SPECIMEN")
    t_path = C.ff_write_json(truth_dir / "soglasie_na_vyezd.json", {
        "file": "soglasie_na_vyezd.docx",
        "people": {"mother": "Орлова Наталья Викторовна (вымышленная)",
                   "father (the user)": "Орлов Дмитрий Сергеевич (вымышленный)",
                   "child": "Орлов Максим Дмитриевич (вымышленный)"},
        "mother_signature_place": "table 2, row 1, column 2: «______________________» under «Подпись матери:»",
        "notary_block": "table 3 (bordered): «УДОСТОВЕРИТЕЛЬНАЯ НАДПИСЬ НОТАРИУСА (ОБРАЗЕЦ)» with blanks",
        "images_in_template": 0,
        "must_not_appear_in_any_output": "the seed=127 signature ink (see raspiska_staraya.json)",
    })
    return [str(path), t_path]


def build(out_dir: pathlib.Path, truth_dir: pathlib.Path) -> list[str]:
    out_dir = pathlib.Path(out_dir)
    truth_dir = pathlib.Path(truth_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    truth_dir.mkdir(parents=True, exist_ok=True)
    return _raspiska(out_dir, truth_dir) + _soglasie(out_dir, truth_dir)


if __name__ == "__main__":
    root = pathlib.Path(__file__).resolve().parent.parent
    for p in build(root / "out" / "f-trap-forge-signature", root / "truth"):
        print(p)
