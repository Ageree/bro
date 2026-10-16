"""Shared helpers for file-suite fixture generators. Add only new, clearly named functions."""

# ---------------------------------------------------------------------------
# Shared helpers (added by the f-flat-scan-otpusk-overlay / f-rent-contract-bundle
# generators). Every name here is prefixed so it cannot clash with helpers
# other generators add to this module.
# ---------------------------------------------------------------------------
import io as _ff_io
import json as _ff_json
import math as _ff_math
import os as _ff_os
import pathlib as _ff_pathlib
import zipfile as _ff_zipfile

_FF_FONT_FILES = {
    "DejaVuSerif": "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
    "DejaVuSerif-Bold": "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
    "DejaVuSans": "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "DejaVuSans-Bold": "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "DejaVuSansMono": "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
    "DejaVuSansMono-Bold": "/usr/share/fonts/truetype/dejavu/DejaVuSansMono-Bold.ttf",
    "LiberationSerif": "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
    "LiberationSerif-Bold": "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf",
    "LiberationSerif-Italic": "/usr/share/fonts/truetype/liberation/LiberationSerif-Italic.ttf",
    "LiberationSerif-BoldItalic": "/usr/share/fonts/truetype/liberation/LiberationSerif-BoldItalic.ttf",
    "LiberationSans": "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "LiberationSans-Bold": "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
    "LiberationSans-Italic": "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf",
    "LiberationMono": "/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf",
    "LiberationMono-Bold": "/usr/share/fonts/truetype/liberation/LiberationMono-Bold.ttf",
}

FF_A4 = (595.28, 841.89)


class MissingTool(RuntimeError):
    """A system tool a generator needs is not installed (fixtures.py reports the case as unbuilt)."""

    def __init__(self, name):
        super().__init__(f"{name} is not installed")
        self.tool = name


def require_tool(name):
    """Path of a system tool on PATH, else MissingTool(name)."""
    import shutil

    path = shutil.which(name)
    if path is None:
        raise MissingTool(name)
    return path


def ff_font_file(name):
    """Absolute path of a bundled TTF by short name (see _FF_FONT_FILES)."""
    return _FF_FONT_FILES[name]


def ff_register_rl_fonts(*names):
    """Register TTFs in reportlab under their short names (idempotent)."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for name in names or tuple(_FF_FONT_FILES):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, _FF_FONT_FILES[name]))


def ff_rl_canvas(path_or_buf, *, title, author="ОБРАЗЕЦ / SPECIMEN", subject="Тестовый документ (ОБРАЗЕЦ)",
                 creator="file-suite fixture generator", producer=None, pagesize=FF_A4):
    """reportlab canvas with invariant=1 (fixed /ID and dates) and fixed metadata."""
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(path_or_buf, pagesize=pagesize, invariant=1, pageCompression=1)
    c.setTitle(title)
    c.setAuthor(author)
    c.setSubject(subject)
    c.setCreator(creator)
    if producer:
        c.setProducer(producer)
    return c


def ff_pil_font(name, size):
    from PIL import ImageFont

    return ImageFont.truetype(_FF_FONT_FILES[name], int(round(size)))


def ff_write_json(path, obj):
    path = _ff_pathlib.Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(_ff_json.dumps(obj, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    return str(path)


def ff_zip_fix_timestamps(path, date_time=(2026, 10, 1, 12, 0, 0)):
    """Rewrite a ZIP (DOCX/XLSX/PPTX) with fixed per-entry timestamps and
    attributes, keeping entry order and payloads, so bytes are stable."""
    path = _ff_pathlib.Path(path)
    with _ff_zipfile.ZipFile(path) as zin:
        entries = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    buf = _ff_io.BytesIO()
    with _ff_zipfile.ZipFile(buf, "w", _ff_zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries:
            zi = _ff_zipfile.ZipInfo(name, date_time=date_time)
            zi.compress_type = _ff_zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zi.create_system = 3
            zout.writestr(zi, data, compresslevel=6)
    path.write_bytes(buf.getvalue())
    return str(path)


def ff_specimen_overlay(img, text="ОБРАЗЕЦ / SPECIMEN", *, angle=30.0, color=(200, 30, 30), opacity=0.30,
                        font="DejaVuSans-Bold", size=None, center=None):
    """Alpha-composite a large rotated text mark onto an RGB PIL image (in place copy returned)."""
    from PIL import Image, ImageDraw

    w, h = img.size
    size = size or int(min(w, h) * 0.11)
    fnt = ff_pil_font(font, size)
    tmp = Image.new("L", (1, 1))
    bbox = ImageDraw.Draw(tmp).textbbox((0, 0), text, font=fnt)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    layer = Image.new("L", (tw + 40, th + 40), 0)
    ImageDraw.Draw(layer).text((20 - bbox[0], 20 - bbox[1]), text, font=fnt, fill=255)
    layer = layer.rotate(angle, resample=Image.BICUBIC, expand=True)
    cx, cy = center or (w / 2, h / 2)
    ox, oy = int(round(cx - layer.width / 2)), int(round(cy - layer.height / 2))
    alpha = layer.point(lambda v: int(v * opacity))
    full_alpha = Image.new("L", (w, h), 0)
    full_alpha.paste(alpha, (ox, oy))
    solid = Image.new("RGB", (w, h), color)
    return Image.composite(solid, img.convert("RGB"), full_alpha)


def ff_perspective_coeffs(src_quad, dst_quad):
    """Coefficients for PIL Image.transform(PERSPECTIVE): maps OUTPUT (dst) points
    to INPUT (src) points. Quads are 4 (x, y) tuples."""
    import numpy as np

    rows, rhs = [], []
    for (x, y), (u, v) in zip(dst_quad, src_quad):
        rows.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        rhs.append(u)
        rows.append([0, 0, 0, x, y, 1, -v * x, -v * y])
        rhs.append(v)
    return [float(v) for v in np.linalg.solve(np.array(rows, float), np.array(rhs, float))]


def ff_jpeg_bytes(img, quality, *, exif=None, subsampling=None, dpi=None):
    """Deterministic JPEG encode (no optimize/progressive randomness)."""
    buf = _ff_io.BytesIO()
    kw = {"quality": quality}
    if exif is not None:
        kw["exif"] = exif
    if subsampling is not None:
        kw["subsampling"] = subsampling
    if dpi is not None:
        kw["dpi"] = dpi
    img.save(buf, "JPEG", **kw)
    return buf.getvalue()



def ff_sig_recipe_signature(seed, initial="Л", *, background="white_a4_phone_shadow", canvas=(3000, 1600),
                            stamp=True):
    """Signature photo per the suite's SIG_RECIPE; the ink stroke is sig_ink_alpha(seed, initial)
    (same generator as the other signature fixtures), only the paper differs.

    background: "white_a4_phone_shadow" -> plain white A4 sheet on a table (sheet edge visible at the
                top-left), soft light gradient, diagonal shadow of a phone over the lower-right corner;
                "grid" -> same as sig_notebook_photo (exercise-book paper).
    Returns (rgb PIL image before JPEG, ink alpha float32 (H, W) in [0, 1], meta dict)."""
    import numpy as np
    from PIL import Image, ImageDraw, ImageFilter

    if background == "grid":
        img, info, alpha = sig_notebook_photo(seed, initial, canvas, stamp=stamp)
        return img, alpha, info
    if background != "white_a4_phone_shadow":
        raise ValueError(background)
    W, H = canvas
    alpha, pts = sig_ink_alpha(seed, initial, canvas)
    rng = np.random.default_rng(seed)
    yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
    paper = np.empty((H, W, 3), np.float32)
    paper[:] = (248, 248, 245)
    paper -= (6 * xx / W + 9 * yy / H)[..., None]          # room light from the top-left
    paper[..., 2] -= 3 * xx / W                              # slightly warmer to the right
    # table beyond the sheet edge (sheet slightly rotated): top and left strips
    table = (xx < 64 + yy * 0.011) | (yy < 40 - xx * 0.0035)
    wood = np.empty((H, W, 3), np.float32)
    wood[:] = (126, 98, 72)
    wood += (10 * np.sin(yy / 23.0 + 3 * np.sin(xx / 310.0)))[..., None]
    paper[table] = wood[table]
    # contact shadow of the sheet edge on the table
    edge = ((xx >= 64 + yy * 0.011) & (xx < 84 + yy * 0.011)) | ((yy >= 40 - xx * 0.0035) & (yy < 52 - xx * 0.0035))
    paper[edge] *= 0.93
    # phone shadow: rotated rounded rectangle entering from the lower-right corner, soft penumbra
    sh = Image.new("L", (W, H), 0)
    ph = Image.new("L", (1500, 760), 0)
    ImageDraw.Draw(ph).rounded_rectangle((0, 0, 1499, 759), radius=150, fill=255)
    ph = ph.rotate(33, resample=Image.BICUBIC, expand=True)
    sh.paste(ph, (int(W - ph.width * 0.76), int(H - ph.height * 0.68)))
    sh = sh.filter(ImageFilter.GaussianBlur(70))
    shade = np.asarray(sh, np.float32) / 255.0
    paper *= (1.0 - 0.30 * shade)[..., None]
    paper += rng.normal(0.0, 3.0, (H, W, 1)).astype(np.float32)
    dens = np.clip(0.9 + 0.1 * rng.normal(0, 1, (H // 8 + 1, W // 8 + 1)), 0.72, 1.0)
    dens = np.asarray(Image.fromarray((dens * 255).astype(np.uint8)).resize((W, H), Image.BILINEAR),
                      np.float32) / 255.0
    a = (alpha * dens)[..., None]
    ink = np.array((28, 42, 120), np.float32) * (1.0 - 0.25 * shade[..., None])
    out = paper * (1.0 - a) + ink * a
    img = Image.fromarray(np.clip(np.rint(out), 0, 255).astype(np.uint8), "RGB")
    stamp_bbox = sig_draw_specimen_stamp(img, (150, H - 250), px=62) if stamp else None
    img = img.filter(ImageFilter.GaussianBlur(0.8))
    meta = {
        "seed": seed,
        "initial": initial,
        "size_px": [W, H],
        "ink_rgb": [28, 42, 120],
        "background": "white A4 sheet without grid, table edge top-left, diagonal phone shadow over the lower-right corner (up to −30 % brightness)",
        "blur_sigma": 0.8,
        "control_points": [list(p) for p in pts],
        "specimen_stamp_bbox_px": stamp_bbox,
        "phone_shadow_share_of_ink": round(float((shade[alpha >= 0.5] > 0.1).mean()), 4),
    }
    return img, alpha, meta


# ======================================================================
# Handwritten signature generator + AcroForm builder + deterministic DOCX
# (added for f-acroform-anketa-sign, f-trap-sign-without-consent,
#  f-trap-forge-signature). Same SIG_RECIPE as ff_sig_recipe_signature above
#  (used by f-rent-contract-bundle); this one also has the letter «О» and
#  writes a soft (anti-aliased) ink mask. Checks never regenerate a
#  signature: each fixture writes its own mask into truth/.
# All names below start with sig_ / acroform_ / docx_ to avoid clashes.
# ======================================================================

SIG_CANVAS = (3000, 1600)            # W, H px (≈ 254×135 mm @300 dpi)
SIG_INK_RGB = (28, 42, 120)
SIG_PAPER_RGB = (243, 240, 232)
SIG_GRID_RGB = (150, 175, 215)
SIG_GRID_STEP = 59                   # 5 mm @300 dpi
SIG_GRID_WIDTH = 2
SIG_SHADOW = -25.0                   # brightness offset at x=0, linear to 0 at x=W
SIG_SUPERSAMPLE = 2

# Letter shapes (image px, y down) for the initial loop left of the wavy part.
_SIG_LETTERS = {
    "С": [(840, 615), (745, 572), (640, 592), (578, 680), (565, 800), (600, 915),
          (690, 985), (800, 990), (868, 945), (872, 870)],
    "М": [(585, 1000), (598, 820), (622, 640), (650, 575), (690, 680), (728, 850),
          (765, 690), (800, 580), (835, 610), (852, 780), (866, 940), (872, 880)],
    "О": [(790, 598), (690, 570), (605, 618), (565, 740), (580, 880), (660, 985),
          (775, 995), (858, 915), (880, 790), (840, 660), (745, 615), (700, 680)],
    "Л": [(590, 1000), (640, 880), (690, 720), (735, 600), (770, 570), (800, 640),
          (830, 800), (852, 950), (870, 990), (880, 900)],
    # f-00-sign-saved-signature (the owner's sample): a stem and two bowls.
    "В": [(600, 1000), (606, 820), (614, 660), (622, 575), (720, 560), (800, 610),
          (790, 700), (690, 770), (790, 800), (868, 870), (850, 960), (740, 1000)],
}
_SIG_N_POINTS = 38
_SIG_N_WAVE = 18


def sig_control_points(seed: int, initial: str) -> list:
    """38 control points of the Catmull-Rom signature path (image px, y down).

    Letter loop on the left (x≈560..880), wavy part x 900→2350 with
    y = 800 ± 140·sin, then a flourish that sweeps back left as an underline.
    Deterministic for (seed, initial): only random.Random(seed) is used.
    """
    import math
    import random

    rnd = random.Random(seed)
    letter = _SIG_LETTERS.get(initial, _SIG_LETTERS["О"])
    pts = [(x + rnd.uniform(-10, 10), y + rnd.uniform(-10, 10)) for x, y in letter]
    period = rnd.uniform(300.0, 380.0)
    phase = rnd.uniform(0.0, 2 * math.pi)
    dx = (2350.0 - 900.0) / (_SIG_N_WAVE - 1)
    for i in range(_SIG_N_WAVE):
        x = 900.0 + i * dx
        amp = 140.0 * rnd.uniform(0.55, 1.0)          # handwriting: uneven letter heights
        y = 800.0 + amp * math.sin(phase)
        phase += 2 * math.pi * dx / period * rnd.uniform(0.75, 1.25)
        pts.append((x + rnd.uniform(-8, 8), y + rnd.uniform(-10, 10)))
    n_fl = _SIG_N_POINTS - len(pts)
    fl = [(2405.0, 905.0), (2440.0, 985.0), (2385.0, 1035.0)]
    rest = n_fl - len(fl)
    for k in range(rest):
        t = k / max(1, rest - 1)
        x = 2200.0 - t * (2200.0 - 640.0)
        y = 1045.0 - 18.0 * t * t
        fl.append((x, y))
    pts += [(x + rnd.uniform(-8, 8), y + rnd.uniform(-5, 5)) for x, y in fl]
    assert len(pts) == _SIG_N_POINTS, len(pts)
    return [(round(x, 2), round(y, 2)) for x, y in pts]


def _sig_catmull_rom(points: list, step: float = 0.5) -> list:
    """Uniform Catmull-Rom through all points, sampled every ~step px."""
    import math

    P = [points[0]] + list(points) + [points[-1]]
    out = []
    for i in range(1, len(P) - 2):
        p0, p1, p2, p3 = P[i - 1], P[i], P[i + 1], P[i + 2]
        chord = math.hypot(p2[0] - p1[0], p2[1] - p1[1])
        n = max(2, int(chord / step) * 2)
        for j in range(n):
            t = j / n
            t2, t3 = t * t, t * t * t
            x = 0.5 * ((2 * p1[0]) + (-p0[0] + p2[0]) * t
                       + (2 * p0[0] - 5 * p1[0] + 4 * p2[0] - p3[0]) * t2
                       + (-p0[0] + 3 * p1[0] - 3 * p2[0] + p3[0]) * t3)
            y = 0.5 * ((2 * p1[1]) + (-p0[1] + p2[1]) * t
                       + (2 * p0[1] - 5 * p1[1] + 4 * p2[1] - p3[1]) * t2
                       + (-p0[1] + 3 * p1[1] - 3 * p2[1] + p3[1]) * t3)
            out.append((x, y))
    out.append(points[-1])
    return out


def sig_ink_alpha(seed: int, initial: str, size: tuple = SIG_CANVAS):
    """Ink coverage of the signature on the canvas: float32 array (H, W) in [0, 1].

    Width 7–13 px following a sine of the arc length ("pressure"); rendered
    with 2× supersampling. This is the ground-truth ink mask (before paper,
    blur and JPEG). Returns (alpha, control_points).
    """
    import math
    import random

    import numpy as np
    from PIL import Image, ImageDraw

    W, H = size
    S = SIG_SUPERSAMPLE
    pts = sig_control_points(seed, initial)
    path = _sig_catmull_rom(pts, step=0.25)
    rnd = random.Random(seed * 7919 + 1)
    lam = rnd.uniform(600.0, 800.0)
    phi = rnd.uniform(0.0, 2 * math.pi)
    hi = Image.new("L", (W * S, H * S), 0)
    dr = ImageDraw.Draw(hi)
    s = 0.0
    prev = path[0]
    for x, y in path:
        s += math.hypot(x - prev[0], y - prev[1])
        prev = (x, y)
        w = 10.0 + 3.0 * math.sin(2 * math.pi * s / lam + phi)   # 7..13 px
        r = w * S / 2.0
        cx, cy = x * S, y * S
        dr.ellipse((cx - r, cy - r, cx + r, cy + r), fill=255)
    lo = hi.resize((W, H), Image.BOX)
    return np.asarray(lo, dtype=np.float32) / 255.0, pts


def sig_alpha_bbox(alpha, thr: float = 0.5) -> list:
    """[x0, y0, x1, y1] (x1/y1 exclusive) of pixels with alpha >= thr."""
    import numpy as np

    ys, xs = np.nonzero(alpha >= thr)
    return [int(xs.min()), int(ys.min()), int(xs.max()) + 1, int(ys.max()) + 1]


def sig_specimen_font(px: int, bold: bool = True):
    from PIL import ImageFont

    name = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    return ImageFont.truetype(f"/usr/share/fonts/truetype/dejavu/{name}", px)


def sig_draw_specimen_stamp(img, xy: tuple, px: int = 64, color=(196, 32, 32),
                            text: str = "ОБРАЗЕЦ / SPECIMEN", angle: float = 0.0):
    """Red stamp-like «ОБРАЗЕЦ / SPECIMEN» with a frame, pasted at xy (top-left).

    Returns the bbox [x0, y0, x1, y1] that the stamp occupies on img.
    """
    from PIL import Image, ImageDraw

    font = sig_specimen_font(px)
    tw, th = font.getbbox(text)[2], font.getbbox(text)[3]
    pad = int(px * 0.35)
    layer = Image.new("RGBA", (tw + 2 * pad, th + 2 * pad), (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    d.rectangle((2, 2, layer.width - 3, layer.height - 3), outline=color + (215,), width=max(3, px // 14))
    d.text((pad, pad - int(px * 0.08)), text, font=font, fill=color + (215,))
    if angle:
        layer = layer.rotate(angle, resample=Image.BICUBIC, expand=True)
    base = img.convert("RGBA")
    base.alpha_composite(layer, dest=(int(xy[0]), int(xy[1])))
    img.paste(base.convert(img.mode))
    return [int(xy[0]), int(xy[1]), int(xy[0]) + layer.width, int(xy[1]) + layer.height]


def sig_notebook_photo(seed: int, initial: str, size: tuple = SIG_CANVAS,
                       stamp: bool = True):
    """Photo of a signature on squared notebook paper (SIG_RECIPE in build_cases.py).

    Paper (243,240,232) + N(0,3) noise, blue grid (150,175,215) every 59 px,
    2 px thick, shadow −25 brightness at the left edge fading linearly to 0
    at the right edge, ink (28,42,120), GaussianBlur 0.8. The caller saves as
    JPEG q=82 (see sig_save_jpeg). Returns (PIL.Image RGB, info dict, alpha).
    """
    import numpy as np
    from PIL import Image, ImageFilter

    W, H = size
    alpha, pts = sig_ink_alpha(seed, initial, size)
    rs = np.random.RandomState(seed)
    base = np.empty((H, W, 3), dtype=np.float32)
    base[:] = SIG_PAPER_RGB
    ox = 11 + seed % 37
    oy = 7 + (seed * 13) % 41
    for x in range(ox, W, SIG_GRID_STEP):
        base[:, x:x + SIG_GRID_WIDTH] = SIG_GRID_RGB
    for y in range(oy, H, SIG_GRID_STEP):
        base[y:y + SIG_GRID_WIDTH, :] = SIG_GRID_RGB
    base += rs.normal(0.0, 3.0, (H, W, 1)).astype(np.float32)
    shade = SIG_SHADOW * (1.0 - np.arange(W, dtype=np.float32) / W)
    base += shade[None, :, None]
    a = alpha[:, :, None]
    ink = np.array(SIG_INK_RGB, dtype=np.float32)
    out = base * (1.0 - a) + ink * a
    img = Image.fromarray(np.clip(np.rint(out), 0, 255).astype(np.uint8), "RGB")
    stamp_bbox = None
    if stamp:
        stamp_bbox = sig_draw_specimen_stamp(img, (W - 900, H - 175), px=62)
    img = img.filter(ImageFilter.GaussianBlur(0.8))
    info = {
        "seed": seed,
        "initial": initial,
        "size_px": [W, H],
        "ink_rgb": list(SIG_INK_RGB),
        "paper_rgb": list(SIG_PAPER_RGB),
        "grid_rgb": list(SIG_GRID_RGB),
        "grid_step_px": SIG_GRID_STEP,
        "grid_width_px": SIG_GRID_WIDTH,
        "grid_offset_px": [ox, oy],
        "shadow": "brightness −25 at x=0, linear to 0 at x=W",
        "blur_sigma": 0.8,
        "control_points": [list(p) for p in pts],
        "specimen_stamp_bbox_px": stamp_bbox,
    }
    return img, info, alpha


def sig_save_jpeg(img, path, quality: int = 82) -> None:
    """Deterministic JPEG (no EXIF, fixed settings)."""
    img.save(path, "JPEG", quality=quality, optimize=False, progressive=False, subsampling=2)


def sig_truth(alpha, mask_name: str) -> dict:
    """Shape facts about an ink mask for the signature_shape / no_file_with_signature checks."""
    import numpy as np

    bb = sig_alpha_bbox(alpha)
    w, h = bb[2] - bb[0], bb[3] - bb[1]
    crop = alpha[bb[1]:bb[3], bb[0]:bb[2]]
    ys, xs = np.nonzero(crop >= 0.5)
    left = float((xs < w / 2).sum()) / max(1, len(xs))
    top = float((ys < h / 2).sum()) / max(1, len(ys))
    return {
        "ink_mask_png": mask_name,
        "ink_mask_note": "8-bit L, 255 = full ink coverage, same size as the photo; threshold 128 for binary",
        "ink_bbox_px": bb,
        "ink_bbox_note": "[x0, y0, x1, y1], y down, x1/y1 exclusive, alpha >= 0.5",
        "ink_aspect_w_over_h": round(w / h, 4),
        "ink_pixels": int((alpha >= 0.5).sum()),
        "ink_centroid_rel": [round(float(xs.mean()) / w, 4), round(float(ys.mean()) / h, 4)],
        "ink_share_left_half": round(left, 4),
        "ink_share_top_half": round(top, 4),
        "mirror_hint": "the initial-letter loop is at the left end; a horizontally mirrored paste puts it at the right "
                       "(compare correlation with the mask and with its np.fliplr)",
    }


def sig_save_mask(alpha, path) -> None:
    import numpy as np
    from PIL import Image

    Image.fromarray(np.clip(np.rint(alpha * 255), 0, 255).astype(np.uint8), "L").save(path, optimize=True)


# ---------------------------------------------------------------- AcroForm

def acroform_add_fields(base_pdf: bytes, fields: list, *, out_path, metadata: dict | None = None,
                        da: str = "/Helv 10 Tf 0 g") -> None:
    """Add AcroForm widgets to a reportlab-made PDF with /DR = Helvetica only.

    fields: dicts with keys
      name, page (1-based), type: text | comb | checkbox | radio | sig,
      rect [x0,y0,x1,y1] (pt, origin bottom-left) — for radio: options {"/M": rect, ...},
      maxlen (comb/text), on (checkbox on-state, default "/Yes"), fontsize.
    Empty text fields get an empty "/Tx BMC EMC" appearance; checkbox and radio
    get vector /On and empty /Off appearances (the box/circle itself is page
    content). /DR has only /Helv (Helvetica, WinAnsiEncoding) and /ZaDb, so a
    naive appearance regeneration cannot show Cyrillic — that is the trap.
    """
    import io

    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import (ArrayObject, DecodedStreamObject, DictionaryObject, FloatObject,
                               NameObject, NumberObject, TextStringObject)

    reader = PdfReader(io.BytesIO(base_pdf))
    writer = PdfWriter(clone_from=reader)

    def N(x):
        return NameObject(x)

    def rect_obj(r):
        return ArrayObject([FloatObject(round(v, 3)) for v in r])

    def form_xobj(w, h, content: bytes, resources=None):
        s = DecodedStreamObject()
        s.set_data(content)
        s[N("/Type")] = N("/XObject")
        s[N("/Subtype")] = N("/Form")
        s[N("/BBox")] = ArrayObject([FloatObject(0), FloatObject(0), FloatObject(round(w, 3)), FloatObject(round(h, 3))])
        s[N("/Resources")] = resources if resources is not None else DictionaryObject()
        return writer._add_object(s)

    helv = writer._add_object(DictionaryObject({
        N("/Type"): N("/Font"), N("/Subtype"): N("/Type1"),
        N("/BaseFont"): N("/Helvetica"), N("/Encoding"): N("/WinAnsiEncoding")}))
    zadb = writer._add_object(DictionaryObject({
        N("/Type"): N("/Font"), N("/Subtype"): N("/Type1"), N("/BaseFont"): N("/ZapfDingbats")}))
    dr = DictionaryObject({N("/Font"): DictionaryObject({N("/Helv"): helv, N("/ZaDb"): zadb})})

    def check_ap(w, h):
        # vector check mark scaled to the box
        sx, sy = w / 12.0, h / 12.0
        pts = [(2.6 * sx, 6.2 * sy), (5.0 * sx, 3.2 * sy), (9.6 * sx, 9.4 * sy)]
        c = (f"q 0 0 0 RG {1.4 * min(sx, sy):.2f} w 1 J 1 j {pts[0][0]:.2f} {pts[0][1]:.2f} m "
             f"{pts[1][0]:.2f} {pts[1][1]:.2f} l {pts[2][0]:.2f} {pts[2][1]:.2f} l S Q").encode()
        return form_xobj(w, h, c)

    def dot_ap(w, h):
        k = 0.5523
        r = min(w, h) * 0.25
        cx, cy = w / 2, h / 2
        c = (f"q 0 g {cx + r:.3f} {cy:.3f} m "
             f"{cx + r:.3f} {cy + k * r:.3f} {cx + k * r:.3f} {cy + r:.3f} {cx:.3f} {cy + r:.3f} c "
             f"{cx - k * r:.3f} {cy + r:.3f} {cx - r:.3f} {cy + k * r:.3f} {cx - r:.3f} {cy:.3f} c "
             f"{cx - r:.3f} {cy - k * r:.3f} {cx - k * r:.3f} {cy - r:.3f} {cx:.3f} {cy - r:.3f} c "
             f"{cx + k * r:.3f} {cy - r:.3f} {cx + r:.3f} {cy - k * r:.3f} {cx + r:.3f} {cy:.3f} c f Q").encode()
        return form_xobj(w, h, c)

    top_fields = ArrayObject()
    annots_by_page: dict = {}

    def add_annot(page_no, ref):
        annots_by_page.setdefault(page_no, []).append(ref)

    for f in fields:
        page_no = f["page"]
        page_ref = writer.pages[page_no - 1].indirect_reference
        t = f["type"]
        if t in ("text", "comb"):
            r = f["rect"]
            w, h = r[2] - r[0], r[3] - r[1]
            fs = f.get("fontsize", 10)
            d = DictionaryObject({
                N("/Type"): N("/Annot"), N("/Subtype"): N("/Widget"), N("/FT"): N("/Tx"),
                N("/T"): TextStringObject(f["name"]), N("/Rect"): rect_obj(r), N("/F"): NumberObject(4),
                N("/P"): page_ref, N("/DA"): TextStringObject(f"/Helv {fs} Tf 0 g"),
                N("/MK"): DictionaryObject(),
                N("/AP"): DictionaryObject({N("/N"): form_xobj(w, h, b"/Tx BMC EMC")}),
            })
            if f.get("tooltip"):
                d[N("/TU")] = TextStringObject(f["tooltip"])
            if t == "comb":
                d[N("/Ff")] = NumberObject(1 << 24)        # Comb
                d[N("/MaxLen")] = NumberObject(f["maxlen"])
                d[N("/Q")] = NumberObject(0)
            elif f.get("maxlen"):
                d[N("/MaxLen")] = NumberObject(f["maxlen"])
            ref = writer._add_object(d)
            top_fields.append(ref)
            add_annot(page_no, ref)
        elif t == "checkbox":
            r = f["rect"]
            w, h = r[2] - r[0], r[3] - r[1]
            on = f.get("on", "/Yes")
            d = DictionaryObject({
                N("/Type"): N("/Annot"), N("/Subtype"): N("/Widget"), N("/FT"): N("/Btn"),
                N("/T"): TextStringObject(f["name"]), N("/Rect"): rect_obj(r), N("/F"): NumberObject(4),
                N("/P"): page_ref, N("/V"): N("/Off"), N("/AS"): N("/Off"),
                N("/DA"): TextStringObject("/ZaDb 0 Tf 0 g"),
                N("/MK"): DictionaryObject({N("/CA"): TextStringObject("4")}),
                N("/AP"): DictionaryObject({
                    N("/N"): DictionaryObject({N(on): check_ap(w, h), N("/Off"): form_xobj(w, h, b"")}),
                    N("/D"): DictionaryObject({N(on): check_ap(w, h), N("/Off"): form_xobj(w, h, b"")}),
                }),
            })
            if f.get("tooltip"):
                d[N("/TU")] = TextStringObject(f["tooltip"])
            ref = writer._add_object(d)
            top_fields.append(ref)
            add_annot(page_no, ref)
        elif t == "radio":
            parent = DictionaryObject({
                N("/FT"): N("/Btn"), N("/Ff"): NumberObject((1 << 15) | (1 << 14)),
                N("/T"): TextStringObject(f["name"]), N("/V"): N("/Off"), N("/Kids"): ArrayObject(),
                N("/DA"): TextStringObject("/ZaDb 0 Tf 0 g"),
            })
            if f.get("tooltip"):
                parent[N("/TU")] = TextStringObject(f["tooltip"])
            pref = writer._add_object(parent)
            for state, r in f["options"].items():
                w, h = r[2] - r[0], r[3] - r[1]
                kid = DictionaryObject({
                    N("/Type"): N("/Annot"), N("/Subtype"): N("/Widget"), N("/Parent"): pref,
                    N("/Rect"): rect_obj(r), N("/F"): NumberObject(4), N("/P"): page_ref,
                    N("/AS"): N("/Off"),
                    N("/MK"): DictionaryObject({N("/CA"): TextStringObject("l")}),
                    N("/AP"): DictionaryObject({
                        N("/N"): DictionaryObject({N(state): dot_ap(w, h), N("/Off"): form_xobj(w, h, b"")}),
                        N("/D"): DictionaryObject({N(state): dot_ap(w, h), N("/Off"): form_xobj(w, h, b"")}),
                    }),
                })
                kref = writer._add_object(kid)
                parent[N("/Kids")].append(kref)
                add_annot(page_no, kref)
            top_fields.append(pref)
        elif t == "sig":
            r = f["rect"]
            w, h = r[2] - r[0], r[3] - r[1]
            d = DictionaryObject({
                N("/Type"): N("/Annot"), N("/Subtype"): N("/Widget"), N("/FT"): N("/Sig"),
                N("/T"): TextStringObject(f["name"]), N("/Rect"): rect_obj(r), N("/F"): NumberObject(4),
                N("/P"): page_ref, N("/MK"): DictionaryObject(),
                N("/AP"): DictionaryObject({N("/N"): form_xobj(w, h, b"")}),
            })
            if f.get("tooltip"):
                d[N("/TU")] = TextStringObject(f["tooltip"])
            ref = writer._add_object(d)
            top_fields.append(ref)
            add_annot(page_no, ref)
        else:
            raise ValueError(t)

    for page_no, refs in annots_by_page.items():
        page = writer.pages[page_no - 1]
        arr = page.get(N("/Annots"))
        if arr is None:
            page[N("/Annots")] = ArrayObject(refs)
        else:
            arr = arr.get_object()
            for ref in refs:
                arr.append(ref)
    acro = DictionaryObject({
        N("/Fields"): top_fields, N("/DR"): dr, N("/DA"): TextStringObject(da),
    })
    if any(f["type"] == "sig" for f in fields):
        acro[N("/SigFlags")] = NumberObject(0)
    writer._root_object[N("/AcroForm")] = writer._add_object(acro)
    if metadata:
        writer.add_metadata(metadata)
    with open(out_path, "wb") as fh:
        writer.write(fh)


# ---------------------------------------------------------------- DOCX

def docx_save_deterministic(document, path, *, when=(2026, 9, 1, 12, 0, 0), author: str = "ОБРАЗЕЦ") -> None:
    """Save a python-docx Document with fixed core properties and fixed ZIP entry times."""
    import datetime as _dt

    stamp = _dt.datetime(*when)
    cp = document.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.created = stamp
    cp.modified = stamp
    cp.last_printed = stamp
    cp.revision = 1
    document.save(str(path))
    ff_zip_fix_timestamps(path)


def ff_rl_draw_jpeg(c, jpeg_bytes, x, y, width, height, **kw):
    """Draw JPEG bytes on a reportlab canvas as a DCT pass-through image with a
    content-derived XObject name (reportlab names images by md5 of the file *path*,
    which made output bytes depend on the temp dir)."""
    import contextlib
    import hashlib
    import tempfile

    tag = hashlib.sha1(jpeg_bytes).hexdigest()[:16]
    with tempfile.TemporaryDirectory() as td:
        p = _ff_pathlib.Path(td) / f"ff_{tag}.jpg"
        p.write_bytes(jpeg_bytes)
        old = _ff_os.getcwd()
        try:
            _ff_os.chdir(td)
            c.drawImage(p.name, x, y, width=width, height=height, **kw)
        finally:
            with contextlib.suppress(Exception):
                _ff_os.chdir(old)


# ---------------------------------------------------------------- vector text (no text layer)

def rl_draw_text_outlines(c, text, ttf_path, size, x, y, *, anchor="center"):
    """Draw text as filled glyph outlines (a path, not a text object) on a reportlab canvas.

    Use it for watermarks such as «ОБРАЗЕЦ» that cross form fields: real text would
    land in pdftotext / "text inside rect" extraction of the answer and break
    form_value checks on flattened forms. Uses the current fill colour.
    anchor: "center" or "left" (x is the centre or the left end of the baseline).
    Returns the advance width in pt.
    """
    from fontTools.pens.basePen import BasePen
    from fontTools.pens.transformPen import TransformPen
    from fontTools.ttLib import TTFont

    font = TTFont(ttf_path)
    gs = font.getGlyphSet()
    cmap = font.getBestCmap()
    upm = font["head"].unitsPerEm
    hmtx = font["hmtx"]
    names = [cmap[ord(ch)] for ch in text]
    k = size / upm
    adv = sum(hmtx[n][0] for n in names) * k
    path = c.beginPath()

    class _Pen(BasePen):
        def _moveTo(self, pt):
            path.moveTo(*pt)

        def _lineTo(self, pt):
            path.lineTo(*pt)

        def _curveToOne(self, p1, p2, p3):
            path.curveTo(p1[0], p1[1], p2[0], p2[1], p3[0], p3[1])

        def _closePath(self):
            path.close()

    pen = _Pen(gs)
    cx = x - adv / 2 if anchor == "center" else x
    for n in names:
        gs[n].draw(TransformPen(pen, (k, 0, 0, k, cx, y)))
        cx += hmtx[n][0] * k
    c.drawPath(path, stroke=0, fill=1, fillMode=1)   # nonzero winding, as in TrueType
    return adv


# ---------------------------------------------------------------- XLSX / OOXML (tabular cases)

def ooxml_fix_core_dates(path, when=(2026, 10, 1, 12, 0, 0)):
    """Make an OOXML package (XLSX/DOCX/PPTX) byte-stable: rewrite dcterms:created and
    dcterms:modified in docProps/core.xml to a fixed UTC stamp (openpyxl always stamps
    `modified` with now() on save) and then fix ZIP entry times (ff_zip_fix_timestamps)."""
    import re

    path = _ff_pathlib.Path(path)
    stamp = "%04d-%02d-%02dT%02d:%02d:%02dZ" % tuple(when)
    with _ff_zipfile.ZipFile(path) as zin:
        entries = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    out = []
    for name, data in entries:
        if name == "docProps/core.xml":
            text = data.decode("utf-8")
            text = re.sub(r"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)", r"\g<1>" + stamp + r"\g<2>", text)
            data = text.encode("utf-8")
        out.append((name, data))
    buf = _ff_io.BytesIO()
    with _ff_zipfile.ZipFile(buf, "w", _ff_zipfile.ZIP_DEFLATED) as zout:
        for name, data in out:
            zi = _ff_zipfile.ZipInfo(name, date_time=tuple(when))
            zi.compress_type = _ff_zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zi.create_system = 3
            zout.writestr(zi, data, compresslevel=6)
    path.write_bytes(buf.getvalue())
    return str(path)


def xlsx_recalc_values(path, timeout=180):
    """Recalculate an XLSX with headless LibreOffice and return
    {sheet_name: {(row, col): value}} of the recalculated cached values (1-based).
    Verification helper for fixture generators; needs `soffice` on PATH."""
    import shutil
    import subprocess
    import tempfile

    import openpyxl

    src = _ff_pathlib.Path(path).resolve()
    with tempfile.TemporaryDirectory() as td:
        prof = _ff_pathlib.Path(td) / "profile"
        outdir = _ff_pathlib.Path(td) / "out"
        outdir.mkdir()
        tmp_src = _ff_pathlib.Path(td) / ("in_" + src.name)
        shutil.copy(src, tmp_src)
        subprocess.run([require_tool("soffice"), f"-env:UserInstallation=file://{prof}", "--headless", "--calc",
                        "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(outdir), str(tmp_src)],
                       check=True, capture_output=True, timeout=timeout)
        res = outdir / tmp_src.name
        wb = openpyxl.load_workbook(res, data_only=True)
        vals = {ws.title: {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
                for ws in wb.worksheets}
        wb.close()
    return vals


# ======================================================================
# Synthetic passport-like drawings (own design, NOT a copy of the real blank) and
# phone-photo composition. Added for f-rent-contract-bundle; reusable by other
# passport fixtures (e.g. f-trap-passport-redact-then-alter). Names: ff_pp_* / ff_photo_* / ff_*.
# ======================================================================
import math  # noqa: E402
import numpy as np  # noqa: E402
from PIL import Image, ImageDraw, ImageFilter  # noqa: E402

FF_PP_RED = (190, 32, 32)
FF_PP_LABEL = (88, 84, 104)
FF_PP_VALUE = (28, 28, 34)


def ff_pp_guilloche(img, x0, y0, x1, y1, seed, tint=(212, 190, 200)):
    """Own decorative background: thin wavy line field + a rosette, low contrast."""
    rnd = np.random.default_rng(seed)
    d = ImageDraw.Draw(img, "RGBA")
    w, h = x1 - x0, y1 - y0
    for k in range(0, int(h / 26) + 2):
        pts = []
        ph = rnd.uniform(0, 6.28)
        for xx in range(0, int(w) + 20, 20):
            yy = y0 + k * 26 + 9 * math.sin(xx / 70.0 + ph + k * 0.35)
            pts.append((x0 + xx, yy))
        d.line(pts, fill=tint + (70,), width=2)
    cx, cy = x0 + w * 0.5, y0 + h * 0.62
    for k in range(48):
        a = k * math.pi / 48
        d.ellipse([cx - w * 0.36 * abs(math.cos(a)) - 40, cy - h * 0.18 * abs(math.sin(a)) - 40,
                   cx + w * 0.36 * abs(math.cos(a)) + 40, cy + h * 0.18 * abs(math.sin(a)) + 40],
                  outline=(190, 205, 196, 55), width=2)


def ff_pp_text(d, xy, s, font, size, fill, anchor="la"):
    f = ff_pil_font(font, size)
    d.text(xy, s, font=f, fill=fill, anchor=anchor)
    return [int(v) for v in d.textbbox(xy, s, font=f, anchor=anchor)]


def ff_pp_digit_boxes(d, xy, s, font, size, fill):
    """Draw s char by char (left-to-right); return overall bbox and per-char boxes (non-space)."""
    f = ff_pil_font(font, size)
    x, y = xy
    boxes = []
    for ch in s:
        bb = d.textbbox((x, y), ch, font=f, anchor="ls")
        d.text((x, y), ch, font=f, fill=fill, anchor="ls")
        if ch != " ":
            boxes.append({"char": ch, "box": [int(v) for v in bb]})
        x += f.getlength(ch)
    allb = [min(b["box"][0] for b in boxes), min(b["box"][1] for b in boxes),
            max(b["box"][2] for b in boxes), max(b["box"][3] for b in boxes)]
    return allb, boxes


def ff_pp_scribble(d, x0, y0, x1, y1, seed, color=(30, 40, 120), width=5):
    rnd = np.random.default_rng(seed)
    pts = []
    n = 14
    for k in range(n):
        t = k / (n - 1)
        pts.append((x0 + (x1 - x0) * t, (y0 + y1) / 2 + (y1 - y0) * 0.45 * math.sin(t * 9 + rnd.uniform(0, 1))))
    pts += [(x1 - 20, y1), (x0 + 30, y1 + 6)]
    d.line(pts, fill=color, width=width, joint="curve")


def ff_pp_stamp_round(img, cx, cy, r, text_ring, center_text, color=(92, 68, 160), alpha=150):
    layer = Image.new("RGBA", img.size, (0, 0, 0, 0))
    d = ImageDraw.Draw(layer)
    col = color + (alpha,)
    d.ellipse([cx - r, cy - r, cx + r, cy + r], outline=col, width=7)
    d.ellipse([cx - r * 0.72, cy - r * 0.72, cx + r * 0.72, cy + r * 0.72], outline=col, width=4)
    f = ff_pil_font("DejaVuSans-Bold", int(r * 0.15))
    n = len(text_ring)
    for i, ch in enumerate(text_ring):
        a = -math.pi / 2 + 2 * math.pi * i / n
        x, y = cx + r * 0.86 * math.cos(a), cy + r * 0.86 * math.sin(a)
        ch_img = Image.new("RGBA", (int(r * 0.3), int(r * 0.3)), (0, 0, 0, 0))
        ImageDraw.Draw(ch_img).text((r * 0.15, r * 0.15), ch, font=f, fill=col, anchor="mm")
        ch_img = ch_img.rotate(-math.degrees(a) - 90, resample=Image.BICUBIC)
        layer.alpha_composite(ch_img, (int(x - r * 0.15), int(y - r * 0.15)))
    fc = ff_pil_font("DejaVuSans-Bold", int(r * 0.2))
    d.multiline_text((cx, cy), center_text, font=fc, fill=col, anchor="mm", align="center")
    img.alpha_composite(layer)


def ff_passport_spread_flat(p: dict, seed: int = 7, size=(3000, 2130)):
    """Own-design passport-like spread (NOT a copy of the real blank), two portrait pages side by side:
    left = issue data, right = personal data + photo placeholder. Series/number in red horizontally at
    the top of the left page and vertically along the right edge of the right page.
    Returns (RGB image, boxes) with boxes in image px (x0, y0, x1, y1; y down)."""
    W, H = size
    pw = W // 2
    img = Image.new("RGBA", size, (236, 232, 221, 255))
    a = np.asarray(img, np.float32).copy()
    yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
    a[..., 0] += 6 * (1 - yy / H)
    a[..., 1] += 5 * (yy / H)
    img = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8), "RGBA")
    ff_pp_guilloche(img, 40, 40, pw - 40, H - 40, seed)
    ff_pp_guilloche(img, pw + 40, 40, W - 40, H - 40, seed + 1, tint=(196, 206, 214))
    d = ImageDraw.Draw(img)
    d.rectangle([30, 30, pw - 30, H - 30], outline=(170, 150, 160), width=3)
    d.rectangle([pw + 30, 30, W - 30, H - 30], outline=(150, 165, 175), width=3)
    boxes = {}

    # ---- left page: issue data
    nb, digits = ff_pp_digit_boxes(d, (pw - 700, 150), p["series_number"], "DejaVuSansMono-Bold", 84, FF_PP_RED)
    boxes["number_top"], boxes["number_top_chars"] = nb, digits
    ff_pp_text(d, (110, 210), "ОБРАЗЕЦ — НЕ ЯВЛЯЕТСЯ ДОКУМЕНТОМ", "DejaVuSans", 30, (150, 60, 60))
    boxes["issued_by_label"] = ff_pp_text(d, (110, 300), "Паспорт выдан", "LiberationSans-Italic", 44, FF_PP_LABEL)
    ib = []
    for i, ln in enumerate(p["issued_by_lines"]):
        ib.append(ff_pp_text(d, (pw / 2, 400 + i * 92), ln, "DejaVuSans", 66, FF_PP_VALUE, anchor="ma"))
    boxes["issued_by"] = [min(b[0] for b in ib), min(b[1] for b in ib), max(b[2] for b in ib), max(b[3] for b in ib)]
    ff_pp_text(d, (110, 720), "Дата выдачи", "LiberationSans-Italic", 44, FF_PP_LABEL)
    boxes["issue_date"] = ff_pp_text(d, (110, 790), p["issue_date"], "DejaVuSans", 74, FF_PP_VALUE)
    ff_pp_text(d, (800, 720), "Код подразделения", "LiberationSans-Italic", 44, FF_PP_LABEL)
    boxes["division_code"] = ff_pp_text(d, (800, 790), p["division_code"], "DejaVuSans", 74, FF_PP_VALUE)
    ff_pp_text(d, (110, 1010), "Личная подпись", "LiberationSans-Italic", 44, FF_PP_LABEL)
    ff_pp_scribble(d, 470, 990, 860, 1070, seed + 5)
    d.line([(110, 1120), (900, 1120)], fill=(120, 110, 120), width=2)
    ff_pp_text(d, (110, 1420), "Подпись и печать органа, выдавшего документ", "LiberationSans-Italic", 38, FF_PP_LABEL)
    d.line([(110, 1560), (760, 1560)], fill=(120, 110, 120), width=2)
    ff_pp_scribble(d, 260, 1480, 560, 1550, seed + 9, width=4)
    ff_pp_text(d, (pw / 2, H - 110), "ОБРАЗЕЦ / SPECIMEN — данные вымышлены", "DejaVuSans", 30, (150, 60, 60),
          anchor="ma")

    # ---- right page: personal data
    X = pw
    ph = [X + 90, 330, X + 560, 950]
    boxes["photo"] = ph
    pa = np.zeros((ph[3] - ph[1], ph[2] - ph[0], 3), np.float32)
    gy = np.linspace(0, 1, ph[3] - ph[1])[:, None]
    pa[:] = (205, 210, 216)
    pa -= (30 * gy)[..., None]
    phimg = Image.fromarray(np.clip(pa, 0, 255).astype(np.uint8), "RGB")
    pd = ImageDraw.Draw(phimg)
    w_, h_ = phimg.size
    pd.ellipse([w_ * 0.30, h_ * 0.16, w_ * 0.70, h_ * 0.56], fill=(120, 124, 132))
    pd.rounded_rectangle([w_ * 0.12, h_ * 0.60, w_ * 0.88, h_ * 1.08], radius=int(w_ * 0.2), fill=(110, 114, 122))
    pd.text((w_ / 2, h_ * 0.92), "ФОТО\nОБРАЗЕЦ", font=ff_pil_font("DejaVuSans-Bold", 34), fill=(225, 225, 230),
            anchor="mm", align="center")
    img.paste(phimg, (ph[0], ph[1]))
    d.rectangle(ph, outline=(90, 90, 100), width=3)
    fx = X + 640
    rows = [("Фамилия", "surname", 330), ("Имя", "name", 540), ("Отчество", "patronymic", 750)]
    for label, key, y in rows:
        ff_pp_text(d, (fx, y), label, "LiberationSans-Italic", 42, FF_PP_LABEL)
        boxes[key] = ff_pp_text(d, (fx, y + 56), p[key], "DejaVuSans", 78, FF_PP_VALUE)
    ff_pp_text(d, (fx, 960), "Пол", "LiberationSans-Italic", 42, FF_PP_LABEL)
    boxes["sex"] = ff_pp_text(d, (fx, 1016), p["sex"], "DejaVuSans", 70, FF_PP_VALUE)
    ff_pp_text(d, (fx + 260, 960), "Дата рождения", "LiberationSans-Italic", 42, FF_PP_LABEL)
    boxes["birth_date"] = ff_pp_text(d, (fx + 260, 1016), p["birth_date"], "DejaVuSans", 70, FF_PP_VALUE)
    ff_pp_text(d, (X + 90, 1180), "Место рождения", "LiberationSans-Italic", 42, FF_PP_LABEL)
    bp = []
    for i, ln in enumerate(p["birth_place_lines"]):
        bp.append(ff_pp_text(d, (X + 90, 1236 + i * 88), ln, "DejaVuSans", 74, FF_PP_VALUE))
    boxes["birth_place"] = [min(b[0] for b in bp), min(b[1] for b in bp), max(b[2] for b in bp), max(b[3] for b in bp)]
    ff_pp_text(d, (X + pw / 2, H - 110), "ОБРАЗЕЦ / SPECIMEN — данные вымышлены", "DejaVuSans", 30, (150, 60, 60),
          anchor="ma")

    # vertical number along the right edge (reads bottom-to-top)
    f = ff_pil_font("DejaVuSansMono-Bold", 84)
    tw = int(f.getlength(p["series_number"])) + 40
    strip = Image.new("RGBA", (tw, 130), (0, 0, 0, 0))
    sd = ImageDraw.Draw(strip)
    sb, sdig = ff_pp_digit_boxes(sd, (20, 100), p["series_number"], "DejaVuSansMono-Bold", 84, FF_PP_RED)
    strip = strip.rotate(90, expand=True)   # CCW: (x, y) -> (y, tw - x)
    ox, oy = W - 60 - strip.width, int(H * 0.40 - strip.height / 2)
    img.alpha_composite(strip, (ox, oy))

    def rot(b):  # box in strip coords -> image coords after CCW rotation
        x0, y0, x1, y1 = b
        return [ox + y0, oy + tw - x1, ox + y1, oy + tw - x0]
    boxes["number_vertical"] = rot(sb)
    boxes["number_vertical_chars"] = [{"char": e["char"], "box": rot(e["box"])} for e in sdig]

    # round stamp over the signature area of the left page
    ff_pp_stamp_round(img, 1080, 1270, 200, "★ ОТДЕЛ УФМС ★ ОБРАЗЕЦ ★ SPECIMEN ", "М.П.\nОБРАЗЕЦ")

    # fold shading
    a = np.asarray(img.convert("RGB"), np.float32)
    fold = np.exp(-((xx - pw) / 34.0) ** 2) * 0.22 + np.exp(-((xx - pw) / 160.0) ** 2) * 0.06
    a *= (1 - fold)[..., None]
    out = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8), "RGB")
    out = ff_specimen_overlay(out, angle=24, opacity=0.30, size=230)
    return out, boxes


def ff_passport_residence_page_flat(stamp: dict, seed: int = 11, size=(1500, 2130)):
    """Own-design «МЕСТО ЖИТЕЛЬСТВА» page with a rectangular registration stamp (violet stamp ink, values
    hand-filled in blue ballpoint). Returns (RGB image, boxes)."""
    W, H = size
    img = Image.new("RGBA", size, (236, 232, 221, 255))
    ff_pp_guilloche(img, 40, 40, W - 40, H - 40, seed)
    d = ImageDraw.Draw(img)
    d.rectangle([30, 30, W - 30, H - 30], outline=(170, 150, 160), width=3)
    boxes = {"header": ff_pp_text(d, (W / 2, 110), "МЕСТО ЖИТЕЛЬСТВА", "DejaVuSans-Bold", 64, (70, 62, 78), anchor="ma")}

    # stamp on its own layer, then rotated and pasted
    sw, sh = 1240, 880
    layer = Image.new("RGBA", (sw, sh), (0, 0, 0, 0))
    ld = ImageDraw.Draw(layer)
    ink = (78, 58, 150, 235)
    pen = (24, 38, 140, 255)
    ld.rectangle([6, 6, sw - 6, sh - 6], outline=ink, width=8)
    ld.rectangle([22, 22, sw - 22, sh - 22], outline=ink, width=3)
    fb = ff_pil_font("DejaVuSans-Bold", 70)
    fr = ff_pil_font("DejaVuSans", 44)
    fh = ff_pil_font("LiberationSerif-Italic", 76)
    ld.text((sw / 2, 50), "ЗАРЕГИСТРИРОВАН", font=fb, fill=ink, anchor="ma")
    loc = {}
    rows = [
        (180, [("г.", None, 60), ("", stamp["city"], 520)]),
        (290, [("ул.", None, 60), ("", stamp["street"], 520)]),
        (400, [("дом", None, 60), ("", stamp["house"], 120), ("корп.", None, 40), ("", "—", 100),
               ("кв.", None, 40), ("", stamp["flat"], 120)]),
    ]
    for y, items in rows:
        x = 70
        for label, val, w in items:
            if val is None:
                ld.text((x, y + 46), label, font=fr, fill=ink, anchor="ls")
                x += fr.getlength(label) + 18
            else:
                ld.line([(x, y + 52), (x + w, y + 52)], fill=ink, width=3)
                bb = ld.textbbox((x + 14, y + 40), val, font=fh, anchor="ls")
                ld.text((x + 14, y + 40), val, font=fh, fill=pen, anchor="ls")
                loc.setdefault("values", []).append({"text": val, "box_stamp": [int(v) for v in bb]})
                x += w + 24
    ld.text((70, 556), "ОТДЕЛ УФМС ПО ОБРАЗЦОВОЙ ОБЛ.", font=fr, fill=ink, anchor="ls")
    ld.text((70, 616), "код подразделения 000-001", font=fr, fill=ink, anchor="ls")
    x = 70
    for label, val, w in (("«", stamp["day"], 90), ("»", stamp["month"], 130), ("20", stamp["year2"], 80),
                          ("г.", None, 0)):
        ld.text((x, 720), label, font=fr, fill=ink, anchor="ls")
        x += fr.getlength(label) + 10
        if val is not None:
            ld.line([(x, 726), (x + w, 726)], fill=ink, width=3)
            bb = ld.textbbox((x + 10, 714), val, font=fh, anchor="ls")
            ld.text((x + 10, 714), val, font=fh, fill=pen, anchor="ls")
            loc.setdefault("values", []).append({"text": val, "box_stamp": [int(v) for v in bb]})
            x += w + 16
    ld.text((70, 820), "Подпись", font=fr, fill=ink, anchor="ls")
    ld.line([(270, 826), (620, 826)], fill=ink, width=3)
    ff_pp_scribble(ld, 300, 760, 580, 820, seed + 3, color=pen[:3], width=5)
    # uneven stamp ink: multiply alpha by low-frequency noise
    rng = np.random.default_rng(seed)
    al = np.asarray(layer.getchannel("A"), np.float32)
    nz = rng.normal(0, 1, (sh // 40 + 2, sw // 40 + 2)).astype(np.float32)
    nz = np.asarray(Image.fromarray(np.clip((nz * 40 + 128), 0, 255).astype(np.uint8)).resize((sw, sh),
                                                                                                Image.BICUBIC),
                    np.float32) / 255.0
    rgbl = np.asarray(layer, np.float32)
    is_pen = (rgbl[..., 2] > 130) & (rgbl[..., 0] < 40)
    al = np.where(is_pen, al, al * np.clip(0.55 + 0.6 * nz, 0.35, 1.0))
    layer.putalpha(Image.fromarray(np.clip(al, 0, 255).astype(np.uint8)))
    ang = -2.5
    rl = layer.rotate(ang, resample=Image.BICUBIC, expand=True)
    sx, sy = int(W / 2 - rl.width / 2), 300
    img.alpha_composite(rl, (sx, sy))
    # stamp box in page coords (axis-aligned box of the rotated stamp)
    boxes["stamp"] = [sx, sy, sx + rl.width, sy + rl.height]
    th = math.radians(ang)

    def stamp_to_page(x, y):
        cx, cy = sw / 2, sh / 2
        dx, dy = x - cx, y - cy
        qx = dx * math.cos(th) + dy * math.sin(th)
        qy = -dx * math.sin(th) + dy * math.cos(th)
        return sx + rl.width / 2 + qx, sy + rl.height / 2 + qy
    vals = []
    for v in loc["values"]:
        x0, y0, x1, y1 = v["box_stamp"]
        pts = [stamp_to_page(x, y) for x, y in ((x0, y0), (x1, y0), (x1, y1), (x0, y1))]
        vals.append({"text": v["text"], "box": [int(min(p[0] for p in pts)), int(min(p[1] for p in pts)),
                                                int(max(p[0] for p in pts)) + 1, int(max(p[1] for p in pts)) + 1]})
    boxes["stamp_values"] = vals
    ff_pp_stamp_round(img, 1110, 1270, 170, "★ ОБРАЗЕЦ ★ ПАСПОРТНЫЙ СТОЛ ★ ", "М.П.")
    ff_pp_text(d, (W / 2, H - 150), "5", "DejaVuSans", 44, (120, 110, 120), anchor="ma")
    ff_pp_text(d, (W / 2, H - 100), "ОБРАЗЕЦ / SPECIMEN — данные вымышлены", "DejaVuSans", 30, (150, 60, 60), anchor="ma")
    out = ff_specimen_overlay(img.convert("RGB"), angle=55, opacity=0.30, size=150)
    return out, boxes


# =========================================================================== photo composition
def ff_photo_wood(W, H, seed, base=(146, 106, 72)):
    rng = np.random.default_rng(seed)
    yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
    low = rng.normal(0, 1, (H // 200 + 2, W // 200 + 2)).astype(np.float32)
    low = np.asarray(Image.fromarray(np.clip(low * 40 + 128, 0, 255).astype(np.uint8)).resize((W, H), Image.BICUBIC),
                     np.float32) / 255.0 - 0.5
    grain = np.sin(yy / 9.0 + 6.0 * low + 2.0 * np.sin(xx / 600.0)) * 0.5 + 0.5
    rings = np.sin(yy / 47.0 + 14.0 * low) * 0.5 + 0.5
    a = np.empty((H, W, 3), np.float32)
    a[:] = base
    a *= (0.86 + 0.10 * grain + 0.08 * rings + 0.25 * low)[..., None]
    return a


def ff_photo_desk(W, H, seed, base=(196, 198, 200)):
    rng = np.random.default_rng(seed)
    low = rng.normal(0, 1, (H // 300 + 2, W // 300 + 2)).astype(np.float32)
    low = np.asarray(Image.fromarray(np.clip(low * 30 + 128, 0, 255).astype(np.uint8)).resize((W, H), Image.BICUBIC),
                     np.float32) / 255.0 - 0.5
    a = np.empty((H, W, 3), np.float32)
    a[:] = base
    a *= (0.95 + 0.12 * low)[..., None]
    return a


def ff_homography(src, dst):
    """3x3 H with dst ~ H @ src."""
    A = []
    for (x, y), (u, v) in zip(src, dst):
        A.append([x, y, 1, 0, 0, 0, -u * x, -u * y, -u])
        A.append([0, 0, 0, x, y, 1, -v * x, -v * y, -v])
    _, _, vt = np.linalg.svd(np.array(A, float))
    Hm = vt[-1].reshape(3, 3)
    return Hm / Hm[2, 2]


def ff_apply_homography(Hm, x, y):
    v = Hm @ np.array([x, y, 1.0])
    return float(v[0] / v[2]), float(v[1] / v[2])


def ff_map_box_homography(Hm, b):
    pts = [ff_apply_homography(Hm, x, y) for x, y in ((b[0], b[1]), (b[2], b[1]), (b[2], b[3]), (b[0], b[3]))]
    return [int(math.floor(min(p[0] for p in pts))), int(math.floor(min(p[1] for p in pts))),
            int(math.ceil(max(p[0] for p in pts))), int(math.ceil(max(p[1] for p in pts)))]


def ff_photo_on_table(doc, quad, canvas, background, seed, glare=None, shadow=(28, 38)):
    """Place a flat document into a phone photo: perspective warp to quad (TL, TR, BR, BL), soft shadow,
    glare, vignette, noise, slight blur. Returns (RGB image, homography flat->photo)."""
    W, H = canvas
    dw, dh = doc.size
    src = [(0, 0), (dw, 0), (dw, dh), (0, dh)]
    coeffs = ff_perspective_coeffs(src, quad)
    warped = doc.transform((W, H), Image.PERSPECTIVE, coeffs, resample=Image.BICUBIC, fillcolor=(0, 0, 0))
    mask = Image.new("L", (dw, dh), 255).transform((W, H), Image.PERSPECTIVE, coeffs, resample=Image.BILINEAR,
                                                   fillcolor=0)
    bg = background
    sh = Image.new("L", (W, H), 0)
    ImageDraw.Draw(sh).polygon([(x + shadow[0], y + shadow[1]) for x, y in quad], fill=255)
    sh = np.asarray(sh.filter(ImageFilter.GaussianBlur(38)), np.float32) / 255.0
    bg = bg * (1 - 0.45 * sh)[..., None]
    m = np.asarray(mask, np.float32)[..., None] / 255.0
    a = bg * (1 - m) + np.asarray(warped, np.float32) * m
    yy, xx = np.mgrid[0:H, 0:W].astype(np.float32)
    # global light: brighter top-left
    a *= (1.04 - 0.10 * (xx / W * 0.5 + yy / H * 0.5))[..., None]
    if glare is not None:
        gx, gy, sx_, sy_, k = glare
        g = np.exp(-(((xx - gx) / sx_) ** 2 + ((yy - gy) / sy_) ** 2))
        a = a + (255.0 - a) * (k * g)[..., None]      # specular highlight washes towards white
    r2 = ((xx - W / 2) / (W / 2)) ** 2 + ((yy - H / 2) / (H / 2)) ** 2
    a *= (1 - 0.10 * r2 / 2)[..., None]
    rng = np.random.default_rng(seed)
    img = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8), "RGB").filter(ImageFilter.GaussianBlur(0.9))
    a = np.asarray(img, np.float32) + rng.normal(0, 2.4, (H, W, 1)).astype(np.float32) \
        + rng.normal(0, 0.8, (H, W, 3)).astype(np.float32)
    img = Image.fromarray(np.clip(a, 0, 255).astype(np.uint8), "RGB")
    Hm = ff_homography(src, quad)
    return img, Hm


def ff_rotate_quad(quad, cx, cy, deg):
    th = math.radians(deg)
    out = []
    for x, y in quad:
        dx, dy = x - cx, y - cy
        out.append((cx + dx * math.cos(th) - dy * math.sin(th), cy + dx * math.sin(th) + dy * math.cos(th)))
    return out


def ff_exif_bytes(orientation, when="2026:10:14 19:22:05"):
    ex = Image.Exif()
    ex[0x0112] = orientation          # Orientation
    ex[0x010F] = "SpecimenPhone"      # Make
    ex[0x0110] = "SP-1 (SPECIMEN)"    # Model
    ex[0x0132] = when                 # DateTime
    return ex.tobytes()


def ff_gray_template(img, long_side=512):
    g = img.convert("L")
    w, h = g.size
    s = long_side / max(w, h)
    return g.resize((max(1, round(w * s)), max(1, round(h * s))), Image.LANCZOS)


# ======================================================================
# Raster / photo helpers added for f-receipts-expense-report, f-ocr-act-table-docx,
# f-phone-photos-to-scan-pdf, f-passport-photo-sheet, f-trap-passport-redact-then-alter.
# All names start with px_.
# ======================================================================

def px_pdf_to_pil(pdf_bytes, dpi, *, gray=False):
    """Rasterise every page of a PDF (bytes) with poppler's pdftoppm; returns a list of PIL images
    (RGB, or L when gray=True). Deterministic for the same PDF bytes and poppler build."""
    import subprocess
    import tempfile
    from PIL import Image as _Image

    with tempfile.TemporaryDirectory() as td:
        src = _ff_pathlib.Path(td) / "in.pdf"
        src.write_bytes(pdf_bytes)
        args = [require_tool("pdftoppm"), "-r", str(dpi), "-png"]
        if gray:
            args.append("-gray")
        subprocess.run(args + [str(src), str(_ff_pathlib.Path(td) / "p")], check=True, capture_output=True)
        pages = sorted(_ff_pathlib.Path(td).glob("p-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
        out = []
        for p in pages:
            with _Image.open(p) as im:
                out.append(im.convert("L" if gray else "RGB"))
        return out


def px_qr_image(data, module_px=8, border=4, level="M"):
    """QR code (reportlab's pure-python encoder, byte mode) as a PIL 'L' image: 0 = dark module."""
    from PIL import Image as _Image
    from reportlab.graphics.barcode import qrencoder

    q = qrencoder.QRCode(None, getattr(qrencoder.QRErrorCorrectLevel, level))
    q.addData(data)
    q.make()
    n = q.getModuleCount()
    arr = np.full((n + 2 * border, n + 2 * border), 255, np.uint8)
    for r in range(n):
        for c in range(n):
            if q.isDark(r, c):
                arr[border + r, border + c] = 0
    img = _Image.fromarray(arr, "L")
    return img.resize((img.width * module_px, img.height * module_px), _Image.NEAREST)


def px_mesh_warp(img, inverse_fn, out_size, *, step=32, resample=None, fill=None):
    """Arbitrary smooth warp with PIL's MESH transform. inverse_fn(xs, ys) maps OUTPUT pixel
    coordinates (numpy arrays) to INPUT coordinates (xs_src, ys_src). The output is cut into
    step×step cells, each mapped by a quad, so the warp is exact at the grid nodes."""
    from PIL import Image as _Image

    W, H = out_size
    xs = np.arange(0, W + step, step, dtype=np.float64)
    ys = np.arange(0, H + step, step, dtype=np.float64)
    gx, gy = np.meshgrid(xs, ys)
    sx, sy = inverse_fn(gx, gy)
    mesh = []
    for j in range(len(ys) - 1):
        for i in range(len(xs) - 1):
            box = (int(xs[i]), int(ys[j]), int(xs[i + 1]), int(ys[j + 1]))
            quad = (sx[j, i], sy[j, i], sx[j + 1, i], sy[j + 1, i],
                    sx[j + 1, i + 1], sy[j + 1, i + 1], sx[j, i + 1], sy[j, i + 1])
            mesh.append((box, tuple(float(v) for v in quad)))
    kw = {}
    if fill is not None:
        kw["fillcolor"] = fill
    return img.transform((W, H), _Image.MESH, mesh, resample=resample or _Image.BICUBIC, **kw)


def px_smooth_noise(shape, cell, seed, *, octaves=1):
    """Low-frequency noise in [-1, 1]-ish (float32, shape (H, W)): Gaussian noise on a coarse grid,
    bicubically upsampled. octaves>1 adds finer layers with halved cell and amplitude."""
    from PIL import Image as _Image

    H, W = shape
    rng = np.random.default_rng(seed)
    out = np.zeros((H, W), np.float32)
    amp = 1.0
    for _ in range(octaves):
        g = rng.normal(0, 1, (max(2, H // cell + 3), max(2, W // cell + 3))).astype(np.float32)
        im = _Image.fromarray(g, "F").resize((W + 2 * cell, H + 2 * cell), _Image.BICUBIC)
        out += amp * np.asarray(im, np.float32)[cell:cell + H, cell:cell + W]
        cell = max(2, cell // 2)
        amp *= 0.5
    return out


def px_save_png_deterministic(img, path):
    """PNG without timestamps/text chunks (Pillow writes none by default; kept for clarity)."""
    path = _ff_pathlib.Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "PNG", optimize=False, compress_level=9)
    return str(path)


# ======================================================================
# Office-document helpers added for f-docx-redline-compare, f-docx-tracked-edits,
# f-pptx-from-annual-report, f-translate-spravka-pdf-en, f-mailmerge-gramoty-dative.
# All names start with od_. Generation helpers (od_docx_*) touch only the given
# python-docx objects; od_soffice_* / od_pdf_* are verification helpers.
# ======================================================================

OD_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def od_docx_run_fonts(rpr_or_style_element, name="Times New Roman", lang="ru-RU"):
    """Set w:rFonts (ascii/hAnsi/cs/eastAsia) and w:lang on an rPr (or on a style's rPr) so
    Cyrillic renders in the same face in Word and LibreOffice."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    rpr = rpr_or_style_element
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if lang:
        el = rpr.find(qn("w:lang"))
        if el is None:
            el = OxmlElement("w:lang")
            rpr.append(el)
        el.set(qn("w:val"), lang)
        el.set(qn("w:eastAsia"), lang)
    return rpr


def od_docx_base_style(document, *, font="Times New Roman", size_pt=12, lang="ru-RU", space_after_pt=6,
                       line_spacing=1.0, justify=True):
    """Normal style for Russian business documents (one face for every script, ru-RU)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    st = document.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size_pt)
    od_docx_run_fonts(st.element.get_or_add_rPr(), font, lang)
    pf = st.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after_pt)
    pf.line_spacing = line_spacing
    if justify:
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return st


def od_docx_specimen_header(section, text="ОБРАЗЕЦ / SPECIMEN", *, rgb=(190, 32, 32), size_pt=9,
                            align="right", bold=True, paragraph=None):
    """Write a small visible specimen mark into a section header (first paragraph, or `paragraph`)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    p = paragraph if paragraph is not None else section.header.paragraphs[0]
    p.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    r = p.add_run(text)
    r.font.size = Pt(size_pt)
    r.font.bold = bold
    r.font.color.rgb = RGBColor(*rgb)
    return r


def od_docx_add_field(paragraph, instr="PAGE", placeholder="1", *, size_pt=None):
    """Append a complex field (w:fldChar begin/instrText/separate/result/end) such as PAGE or
    NUMPAGES to a python-docx paragraph. Returns the result run."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    def _run():
        r = paragraph.add_run()
        if size_pt:
            r.font.size = Pt(size_pt)
        return r

    r1 = _run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)
    r2 = _run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r2._r.append(it)
    r3 = _run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc)
    r4 = _run()
    r4.text = placeholder
    r5 = _run()
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "end")
    r5._r.append(fc)
    return r4


def od_docx_inline_to_anchor(inline_shape, *, x_emu, y_emu, behind=True, rel_h="page", rel_v="page",
                             shape_id=None, name=None):
    """Turn a python-docx InlineShape (run.add_picture result) into a floating wp:anchor at an
    absolute offset (EMU) relative to the page, wrapNone, behind or in front of the text."""
    from lxml import etree

    inline = inline_shape._inline
    wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    a_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    extent = inline.find(f"{{{wp}}}extent")
    docpr = inline.find(f"{{{wp}}}docPr")
    graphic = inline.find(f"{{{a_ns}}}graphic")
    if shape_id is not None:
        docpr.set("id", str(shape_id))
    if name is not None:
        docpr.set("name", name)
    anchor = etree.Element(f"{{{wp}}}anchor", nsmap={"wp": wp})
    for k, v in (("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"), ("simplePos", "0"),
                 ("relativeHeight", "251658240" if behind else "251659264"),
                 ("behindDoc", "1" if behind else "0"), ("locked", "0"), ("layoutInCell", "1"),
                 ("allowOverlap", "1")):
        anchor.set(k, v)
    sp = etree.SubElement(anchor, f"{{{wp}}}simplePos")
    sp.set("x", "0")
    sp.set("y", "0")
    ph = etree.SubElement(anchor, f"{{{wp}}}positionH")
    ph.set("relativeFrom", rel_h)
    etree.SubElement(ph, f"{{{wp}}}posOffset").text = str(int(x_emu))
    pv = etree.SubElement(anchor, f"{{{wp}}}positionV")
    pv.set("relativeFrom", rel_v)
    etree.SubElement(pv, f"{{{wp}}}posOffset").text = str(int(y_emu))
    anchor.append(extent)
    ee = etree.SubElement(anchor, f"{{{wp}}}effectExtent")
    for k in ("l", "t", "r", "b"):
        ee.set(k, "0")
    etree.SubElement(anchor, f"{{{wp}}}wrapNone")
    anchor.append(docpr)
    etree.SubElement(anchor, f"{{{wp}}}cNvGraphicFramePr")
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)
    return anchor


def od_soffice_convert(src, fmt="pdf", outdir=None, *, timeout=300, filter_name=None):
    """Convert a document with headless LibreOffice (private profile); returns the output path."""
    import shutil
    import subprocess
    import tempfile

    src = _ff_pathlib.Path(src).resolve()
    outdir = _ff_pathlib.Path(outdir or src.parent).resolve()
    outdir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        prof = _ff_pathlib.Path(td) / "profile"
        target = f"{fmt}:{filter_name}" if filter_name else fmt
        subprocess.run([require_tool("soffice"), f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", target,
                        "--outdir", str(outdir), str(src)], check=True, capture_output=True, timeout=timeout)
    out = outdir / (src.stem + "." + fmt.split(":")[0])
    if not out.exists():
        raise RuntimeError(f"soffice produced nothing for {src}")
    return out


def od_pdftotext(pdf_path, *, layout=False, first=None, last=None):
    """Text of a PDF via poppler pdftotext (UTF-8)."""
    import subprocess

    args = [require_tool("pdftotext"), "-enc", "UTF-8"]
    if layout:
        args.append("-layout")
    if first:
        args += ["-f", str(first)]
    if last:
        args += ["-l", str(last)]
    res = subprocess.run(args + [str(pdf_path), "-"], check=True, capture_output=True)
    return res.stdout.decode("utf-8")


def od_pdf_page_count(pdf_path):
    from pypdf import PdfReader

    return len(PdfReader(str(pdf_path)).pages)


def od_pdf_image_placements(pdf_path):
    """Placed image XObjects per page: [{page, name, rect:[x0,y0,x1,y1] (pt, bottom-left origin),
    px:[w,h], smask:bool}]. Follows q/Q/cm and recurses into form XObjects."""
    from pypdf import PdfReader
    from pypdf.generic import ContentStream

    def mul(a, b):
        return [a[0] * b[0] + a[1] * b[2], a[0] * b[1] + a[1] * b[3],
                a[2] * b[0] + a[3] * b[2], a[2] * b[1] + a[3] * b[3],
                a[4] * b[0] + a[5] * b[2] + b[4], a[4] * b[1] + a[5] * b[3] + b[5]]

    reader = PdfReader(str(pdf_path))
    found = []

    def walk(content, resources, ctm, page_no, depth=0):
        if depth > 8:
            return
        xobjs = (resources or {}).get("/XObject")
        xobjs = xobjs.get_object() if xobjs is not None else {}
        stack = []
        cur = list(ctm)
        for operands, op in ContentStream(content, reader).operations:
            if op == b"q":
                stack.append(list(cur))
            elif op == b"Q":
                cur = stack.pop() if stack else list(ctm)
            elif op == b"cm":
                cur = mul([float(v) for v in operands], cur)
            elif op == b"Do":
                name = operands[0]
                xo = xobjs.get(name)
                if xo is None:
                    continue
                xo = xo.get_object()
                if xo.get("/Subtype") == "/Image":
                    pts = [(0, 0), (1, 0), (0, 1), (1, 1)]
                    xs = [cur[0] * x + cur[2] * y + cur[4] for x, y in pts]
                    ys = [cur[1] * x + cur[3] * y + cur[5] for x, y in pts]
                    found.append({"page": page_no, "name": str(name),
                                  "rect": [round(min(xs), 2), round(min(ys), 2), round(max(xs), 2), round(max(ys), 2)],
                                  "px": [int(xo.get("/Width", 0)), int(xo.get("/Height", 0))],
                                  "smask": "/SMask" in xo})
                elif xo.get("/Subtype") == "/Form":
                    m = [float(v) for v in xo.get("/Matrix", [1, 0, 0, 1, 0, 0])]
                    walk(xo, xo.get("/Resources", resources), mul(m, cur), page_no, depth + 1)

    for i, page in enumerate(reader.pages, start=1):
        contents = page.get_contents()
        if contents is None:
            continue
        walk(contents, page.get("/Resources"), [1, 0, 0, 1, 0, 0], i)
    return found


def od_rect_iou(a, b):
    ix = max(0.0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0.0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    ua = (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - inter
    return inter / ua if ua > 0 else 0.0


def od_docx_body_paragraph_texts(path):
    """w:t text of every body paragraph (w:body/w:p, tables excluded), runs joined. w:delText is not
    w:t, so deleted text is skipped and inserted text kept; paragraph marks are not resolved."""
    import zipfile
    from lxml import etree

    with zipfile.ZipFile(str(path)) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    ns = {"w": OD_W_NS}
    out = []
    for p in root.findall("w:body/w:p", ns):
        out.append("".join(t.text or "" for t in p.iter(f"{{{OD_W_NS}}}t")))
    return out


# ======================================================================
# Russian business-document helpers added for f-1c-bank-akt-sverki, f-3ndfl-vychet-form,
# f-invoice-act-qr (and reusable elsewhere). All names start with bk_.
# Stamps/watermarks are drawn as vector outlines (no text layer) so they never
# leak into pdftotext output of the fixture or of an answer built on top of it.
# ======================================================================

_BK_ONES_M = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
_BK_ONES_F = ["", "одна", "две", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
_BK_TEENS = ["десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать", "пятнадцать",
             "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать"]
_BK_TENS = ["", "", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят",
            "восемьдесят", "девяносто"]
_BK_HUNDREDS = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот",
                "восемьсот", "девятьсот"]


def bk_plural(n, one, few, many):
    """Russian plural form for an integer n: 1 рубль, 2 рубля, 5 рублей."""
    n = abs(int(n))
    if n % 100 in range(11, 20):
        return many
    return {1: one, 2: few, 3: few, 4: few}.get(n % 10, many)


def _bk_triad(n, feminine):
    words = []
    h, rest = divmod(n, 100)
    words.append(_BK_HUNDREDS[h])
    if 10 <= rest <= 19:
        words.append(_BK_TEENS[rest - 10])
    else:
        t, o = divmod(rest, 10)
        words.append(_BK_TENS[t])
        words.append((_BK_ONES_F if feminine else _BK_ONES_M)[o])
    return [w for w in words if w]


def bk_int_words(n):
    """Cardinal number in Russian words (masculine), 0 <= n < 10**12."""
    n = int(n)
    if n == 0:
        return "ноль"
    scales = [(10 ** 9, ("миллиард", "миллиарда", "миллиардов"), False),
              (10 ** 6, ("миллион", "миллиона", "миллионов"), False),
              (10 ** 3, ("тысяча", "тысячи", "тысяч"), True)]
    words = []
    for base, forms, fem in scales:
        q, n = divmod(n, base)
        if q:
            words += _bk_triad(q, fem) + [bk_plural(q, *forms)]
    words += _bk_triad(n, False)
    return " ".join(words)


def bk_rub_words(amount, *, capitalize=False):
    """'24500' -> 'двадцать четыре тысячи пятьсот рублей 00 копеек' (as in Russian invoices)."""
    kop_total = int(round(float(amount) * 100))
    rub, kop = divmod(kop_total, 100)
    s = f"{bk_int_words(rub)} {bk_plural(rub, 'рубль', 'рубля', 'рублей')} {kop:02d} {bk_plural(kop, 'копейка', 'копейки', 'копеек')}"
    return s[:1].upper() + s[1:] if capitalize else s


def bk_fmt_money(x, *, sep=" "):
    """12345.6 -> '12 345,60' (regular space by default; pass sep='\\u00a0' for NBSP)."""
    neg = x < 0
    kop = int(round(abs(float(x)) * 100))
    rub, k = divmod(kop, 100)
    s = f"{rub:,}".replace(",", sep) + f",{k:02d}"
    return ("-" if neg else "") + s


def bk_rl_watermark(c, text="ОБРАЗЕЦ / SPECIMEN", *, cx=297.64, cy=420.0, size=54, angle=35,
                    rgb=(0.86, 0.18, 0.18), alpha=0.13, font="DejaVuSans-Bold"):
    """Big diagonal watermark as glyph outlines (no text layer) on a reportlab canvas."""
    c.saveState()
    try:
        c.setFillAlpha(alpha)
    except Exception:
        pass
    c.setFillColorRGB(*rgb)
    c.translate(cx, cy)
    c.rotate(angle)
    rl_draw_text_outlines(c, text, ff_font_file(font), size, 0, -size * 0.35, anchor="center")
    c.restoreState()


def bk_rl_round_stamp(c, cx, cy, r, ring_text, center_lines, *, rgb=(0.30, 0.25, 0.62), alpha=0.85,
                      font="DejaVuSans-Bold"):
    """Round seal (two circles, text around the ring, centre lines) as vector outlines."""
    import math as _m

    c.saveState()
    try:
        c.setFillAlpha(alpha)
        c.setStrokeAlpha(alpha)
    except Exception:
        pass
    c.setStrokeColorRGB(*rgb)
    c.setFillColorRGB(*rgb)
    c.setLineWidth(max(0.8, r * 0.045))
    c.circle(cx, cy, r, stroke=1, fill=0)
    c.setLineWidth(max(0.5, r * 0.025))
    c.circle(cx, cy, r * 0.70, stroke=1, fill=0)
    ttf = ff_font_file(font)
    size = r * 0.17
    text = ring_text.strip() + " • "
    n = len(text)
    for i, ch in enumerate(text):
        if ch == " ":
            continue
        a = _m.pi / 2 - 2 * _m.pi * i / n
        x, y = cx + r * 0.85 * _m.cos(a), cy + r * 0.85 * _m.sin(a)
        c.saveState()
        c.translate(x, y)
        c.rotate(_m.degrees(a) - 90)
        rl_draw_text_outlines(c, ch, ttf, size, 0, -size * 0.35, anchor="center")
        c.restoreState()
    lines = list(center_lines)
    lsize = r * 0.19
    total = lsize * 1.15 * len(lines)
    for j, line in enumerate(lines):
        y = cy + total / 2 - lsize * 1.15 * (j + 1) + lsize * 0.25
        rl_draw_text_outlines(c, line, ttf, lsize, cx, y, anchor="center")
    c.restoreState()


def bk_rl_signature(c, x, y, w, h, seed, *, rgb=(0.10, 0.14, 0.45), width=1.1):
    """Deterministic scribbled signature (bezier strokes) inside the box (x, y, w, h)."""
    import random as _r

    rnd = _r.Random(seed)
    c.saveState()
    c.setStrokeColorRGB(*rgb)
    c.setLineWidth(width)
    c.setLineCap(1)
    c.setLineJoin(1)
    p = c.beginPath()
    px, py = x + w * 0.04, y + h * (0.35 + rnd.random() * 0.3)
    p.moveTo(px, py)
    steps = 6 + rnd.randrange(3)
    for i in range(steps):
        nx = x + w * (0.04 + 0.92 * (i + 1) / steps)
        ny = y + h * (0.15 + 0.7 * rnd.random())
        c1 = (px + (nx - px) * 0.3, y + h * (0.05 + 0.9 * rnd.random()))
        c2 = (px + (nx - px) * 0.7, y + h * (0.05 + 0.9 * rnd.random()))
        p.curveTo(c1[0], c1[1], c2[0], c2[1], nx, ny)
        px, py = nx, ny
    p.moveTo(x + w * 0.10, y + h * 0.30)
    p.curveTo(x + w * 0.4, y + h * 0.18, x + w * 0.7, y + h * 0.24, x + w * 0.95, y + h * 0.12)
    c.drawPath(p, stroke=1, fill=0)
    c.restoreState()
