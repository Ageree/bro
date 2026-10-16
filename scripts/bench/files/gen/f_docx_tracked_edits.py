"""Fixture for f-docx-tracked-edits.

dogovor_postavki.docx         «ДОГОВОР ПОСТАВКИ № 27/26 (ОБРАЗЕЦ)», ~5 pages A4. Sections 1–9 and their clauses are
                              one multi-level auto-numbered list: numbering.xml abstractNum 20 (multiLevelType
                              multilevel, lvl0 «%1.», lvl1 «%1.%2.», suffix space) → w:num 20; every numbered paragraph
                              carries a direct w:numPr (numId 20, ilvl 0 for section headings, ilvl 1 for clauses).
                              4.3 «Покупатель оплачивает Товар в течение 10 (десяти) календарных дней с даты поставки.»
                              Section 6 = 6.1–6.4 (6.2 penalties «Сторона, нарушившая сроки поставки…», 6.3 «Уплата
                              неустойки не освобождает Стороны от исполнения обязательств.»). Header: logo PNG
                              «ОБРАЗЕЦ» + «ООО «Образец Поставка»» + «ОБРАЗЕЦ / SPECIMEN»; footer: «Страница {PAGE}».
                              At the end a 2-column requisites table (Поставщик = us, Покупатель = counterparty) with
                              one requisite per paragraph; our bank: АО «Банк Старый Образец», р/с …0111, БИК 044500111,
                              к/с …0111. No word «конфиденциальн» anywhere (the clause is to be added).
kartochka_predpriyatiya.pdf   one-page text PDF «Карточка предприятия» of ООО «Образец Поставка»: same ИНН/КПП/ОГРН/
                              address, new bank АО «Банк Новый Образец», р/с …0222, БИК 044500222, к/с …0222; 2-column
                              table; diagonal «ОБРАЗЕЦ / SPECIMEN» drawn as outlines (no text layer).

truth/dogovor_postavki.json   numbering ids, clause texts and positions, requisites old/new, rendered clause labels.
"""

from __future__ import annotations

import copy
import io
import pathlib
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw

HERE = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))
import common as C  # noqa: E402

CASE = "f-docx-tracked-edits"
DOCX_NAME = "dogovor_postavki.docx"
CARD_NAME = "kartochka_predpriyatiya.pdf"
WHEN = (2026, 10, 1, 9, 30, 0)
NUM_ID = 20
ABSTRACT_ID = 20

TITLE = "ДОГОВОР ПОСТАВКИ № 27/26 (ОБРАЗЕЦ)"
PLACE_DATE = ("г. Москва", "«01» октября 2026 г.")
PREAMBLE = (
    "Общество с ограниченной ответственностью «Образец Поставка» (ООО «Образец Поставка»), именуемое в дальнейшем "
    "«Поставщик», в лице генерального директора Образцова Ивана Ивановича, действующего на основании Устава, с одной "
    "стороны, и Общество с ограниченной ответственностью «Тест-Покупатель» (ООО «Тест-Покупатель»), именуемое в "
    "дальнейшем «Покупатель», в лице генерального директора Тестова Тимура Тимуровича, действующего на основании "
    "Устава, с другой стороны, совместно именуемые «Стороны», а по отдельности — «Сторона», заключили настоящий "
    "договор (далее — Договор) о нижеследующем:"
)

CLAUSE_43 = "Покупатель оплачивает Товар в течение 10 (десяти) календарных дней с даты поставки."
CLAUSE_62 = ("Сторона, нарушившая сроки поставки, уплачивает пени в размере 0,1 % стоимости непоставленного Товара за "
             "каждый день просрочки.")
CLAUSE_63 = "Уплата неустойки не освобождает Стороны от исполнения обязательств."

SECTIONS = [
    ("ПРЕДМЕТ ДОГОВОРА", [
        "Поставщик обязуется передать в собственность Покупателя офисную мебель и сопутствующие товары (далее — Товар), "
        "а Покупатель обязуется принять и оплатить Товар на условиях Договора.",
        "Наименование, ассортимент, количество, цена и сроки поставки каждой партии Товара согласовываются Сторонами в "
        "спецификациях, которые являются неотъемлемыми частями Договора.",
        "Поставщик гарантирует, что Товар принадлежит ему на праве собственности, не заложен, не находится под арестом "
        "и не обременён правами третьих лиц.",
        "Поставка Товара осуществляется партиями на основании заявок Покупателя, направленных по электронной почте не "
        "позднее чем за 5 (пять) рабочих дней до желаемой даты поставки.",
        "Спецификация считается согласованной с момента её подписания уполномоченными представителями обеих Сторон. "
        "Заявка, не подтверждённая Поставщиком в течение 2 (двух) рабочих дней, считается отклонённой.",
    ]),
    ("КАЧЕСТВО И КОМПЛЕКТНОСТЬ ТОВАРА", [
        "Качество Товара должно соответствовать техническим регламентам, государственным стандартам и иным "
        "обязательным требованиям, действующим в Российской Федерации, а также условиям спецификаций.",
        "Товар поставляется в упаковке изготовителя, обеспечивающей его сохранность при транспортировке и хранении. "
        "Каждое место маркируется с указанием наименования Товара, количества и номера спецификации.",
        "Вместе с Товаром Поставщик передаёт Покупателю товарную накладную или универсальный передаточный документ, "
        "счёт, а также сертификаты или декларации соответствия, если они обязательны для данного Товара.",
        "Гарантийный срок на Товар составляет 12 (двенадцать) месяцев с даты поставки, если больший срок не установлен "
        "изготовителем. В течение гарантийного срока Поставщик за свой счёт устраняет недостатки или заменяет Товар.",
        "Товар, требующий сборки, поставляется с инструкцией по сборке и эксплуатации на русском языке. По заявке "
        "Покупателя Поставщик выполняет сборку и расстановку Товара на объекте Покупателя по ценам, указанным в "
        "спецификации.",
        "Комплектность Товара должна соответствовать спецификации и документации изготовителя. Некомплектный Товар "
        "считается непоставленным до его доукомплектования.",
    ]),
    ("ПОРЯДОК И СРОКИ ПОСТАВКИ", [
        "Поставка осуществляется силами и за счёт Поставщика по адресу склада Покупателя, указанному в заявке, в "
        "пределах города Москвы и Московской области.",
        "Срок поставки партии Товара составляет не более 10 (десяти) рабочих дней с даты согласования спецификации, "
        "если иной срок не указан в спецификации.",
        "Поставщик уведомляет Покупателя о дате и времени доставки не позднее чем за 1 (один) рабочий день до поставки.",
        "Датой поставки считается дата подписания Покупателем товарной накладной или универсального передаточного "
        "документа.",
        "Право собственности на Товар и риск его случайной гибели переходят к Покупателю с даты поставки.",
        "Поставщик вправе с согласия Покупателя досрочно поставить Товар или поставить его по частям в пределах "
        "согласованного срока.",
        "Разгрузка Товара на складе Покупателя производится силами Поставщика. Покупатель обеспечивает беспрепятственный "
        "доступ транспорта Поставщика к месту разгрузки в рабочее время с 9:00 до 18:00.",
        "Если Покупатель не может принять Товар в согласованную дату, он уведомляет об этом Поставщика не позднее чем "
        "за 1 (один) рабочий день. Расходы на повторную доставку по вине Покупателя оплачиваются Покупателем по "
        "тарифам Поставщика.",
    ]),
    ("ЦЕНА ТОВАРА И ПОРЯДОК РАСЧЁТОВ", [
        "Цена Товара указывается в спецификациях в рублях Российской Федерации и включает НДС по ставке, установленной "
        "законодательством, стоимость упаковки, маркировки, погрузки и доставки до склада Покупателя.",
        "Цена согласованной партии Товара не может быть изменена в одностороннем порядке после подписания спецификации.",
        CLAUSE_43,
        "Оплата производится безналичным перечислением денежных средств на расчётный счёт Поставщика, указанный в "
        "реквизитах Сторон. Обязательство Покупателя по оплате считается исполненным с даты зачисления денежных "
        "средств на корреспондентский счёт банка Поставщика.",
        "Стороны ежеквартально, не позднее 15-го числа месяца, следующего за кварталом, проводят сверку взаимных "
        "расчётов и подписывают акт сверки.",
        "Поставщик выставляет счёт на оплату партии Товара одновременно с её отгрузкой. Счёт, товарная накладная и "
        "счёт-фактура (или универсальный передаточный документ) направляются Покупателю через систему электронного "
        "документооборота.",
        "Расходы на банковские комиссии несёт Сторона, производящая платёж. Проценты на сумму отсрочки платежа "
        "(коммерческий кредит) не начисляются.",
    ]),
    ("ПРИЁМКА ТОВАРА", [
        "Приёмка Товара по количеству мест и внешнему виду упаковки производится Покупателем в момент получения Товара "
        "от Поставщика.",
        "Приёмка Товара по количеству единиц и качеству производится в течение 5 (пяти) рабочих дней с даты поставки. "
        "При выявлении недостачи или недостатков Покупатель составляет акт и направляет его Поставщику.",
        "Поставщик обязан в течение 10 (десяти) рабочих дней с даты получения акта восполнить недостачу или заменить "
        "Товар ненадлежащего качества.",
        "Скрытые недостатки, которые не могли быть обнаружены при обычной приёмке, могут быть заявлены Покупателем в "
        "течение гарантийного срока.",
        "Возврат Товара ненадлежащего качества осуществляется силами и за счёт Поставщика.",
        "Для участия в составлении акта о недостатках Поставщик вправе направить своего представителя. Если "
        "представитель не прибыл в течение 3 (трёх) рабочих дней с даты уведомления, акт составляется Покупателем "
        "в одностороннем порядке и является обязательным для Сторон.",
    ]),
    ("ОТВЕТСТВЕННОСТЬ СТОРОН", [
        "За неисполнение или ненадлежащее исполнение обязательств по Договору Стороны несут ответственность в "
        "соответствии с законодательством Российской Федерации и Договором.",
        CLAUSE_62,
        CLAUSE_63,
        "Сторона обязана возместить другой Стороне реальный ущерб, причинённый неисполнением или ненадлежащим "
        "исполнением Договора. Упущенная выгода возмещению не подлежит.",
    ]),
    ("ОБСТОЯТЕЛЬСТВА НЕПРЕОДОЛИМОЙ СИЛЫ", [
        "Стороны освобождаются от ответственности за частичное или полное неисполнение обязательств по Договору, если "
        "оно явилось следствием обстоятельств непреодолимой силы, возникших после заключения Договора.",
        "Сторона, для которой исполнение обязательств стало невозможным, обязана в течение 5 (пяти) рабочих дней "
        "письменно уведомить другую Сторону о наступлении и прекращении таких обстоятельств.",
        "Если обстоятельства непреодолимой силы продолжаются более 2 (двух) месяцев, каждая из Сторон вправе "
        "отказаться от исполнения Договора, предварительно уведомив другую Сторону.",
        "Надлежащим доказательством наличия обстоятельств непреодолимой силы и их продолжительности служат документы, "
        "выданные компетентными органами или Торгово-промышленной палатой Российской Федерации.",
    ]),
    ("ПОРЯДОК РАЗРЕШЕНИЯ СПОРОВ", [
        "Споры и разногласия, возникающие из Договора, Стороны разрешают путём переговоров.",
        "До обращения в суд обязателен претензионный порядок. Срок рассмотрения претензии — 15 (пятнадцать) "
        "календарных дней с даты её получения.",
        "При недостижении согласия спор передаётся на рассмотрение в Арбитражный суд города Москвы.",
        "Претензии и уведомления направляются заказным письмом с уведомлением о вручении по адресу Стороны, "
        "указанному в Договоре, либо по электронной почте с последующим направлением оригинала.",
    ]),
    ("СРОК ДЕЙСТВИЯ И ЗАКЛЮЧИТЕЛЬНЫЕ ПОЛОЖЕНИЯ", [
        "Договор вступает в силу с даты его подписания Сторонами и действует до 31 декабря 2027 года, а в части "
        "расчётов — до полного исполнения Сторонами своих обязательств.",
        "Договор может быть расторгнут по соглашению Сторон или по иным основаниям, предусмотренным законодательством "
        "Российской Федерации.",
        "Все изменения и дополнения к Договору действительны, если они совершены в письменной форме и подписаны "
        "уполномоченными представителями Сторон.",
        "Сторона обязана в течение 3 (трёх) рабочих дней письменно уведомить другую Сторону об изменении своего "
        "наименования, адреса или банковских реквизитов. До получения такого уведомления исполнение по прежним "
        "реквизитам считается надлежащим.",
        "Документы, подписанные и переданные по электронной почте или через систему электронного документооборота, "
        "имеют юридическую силу до обмена оригиналами.",
        "Ни одна из Сторон не вправе передавать свои права и обязанности по Договору третьим лицам без "
        "письменного согласия другой Стороны.",
        "Во всём, что не предусмотрено Договором, Стороны руководствуются законодательством Российской Федерации.",
        "Договор составлен в двух экземплярах, имеющих одинаковую юридическую силу, по одному для каждой из Сторон.",
    ]),
]

SUPPLIER_OLD = [
    ("name", "ООО «Образец Поставка»"),
    ("addr", "Юридический адрес: 101000, г. Москва, ул. Образцовая, д. 5, офис 12"),
    ("inn", "ИНН 7700000086"),
    ("kpp", "КПП 770001001"),
    ("ogrn", "ОГРН 1000000000000"),
    ("rs", "р/с 40702810000000000111"),
    ("bank", "в АО «Банк Старый Образец», г. Москва"),
    ("bik", "БИК 044500111"),
    ("ks", "к/с 30101810000000000111"),
    ("tel", "Тел.: +7 900 000-00-86"),
    ("mail", "E-mail: info@obrazec-postavka.example"),
]
BUYER = [
    "ООО «Тест-Покупатель»",
    "Юридический адрес: 190000, г. Санкт-Петербург, ул. Тестовая, д. 7, лит. А",
    "ИНН 7800000107",
    "КПП 780101001",
    "ОГРН 1000000000107",
    "р/с 40702810000000000333",
    "в АО «Банк Тестовый», г. Санкт-Петербург",
    "БИК 044000333",
    "к/с 30101810000000000333",
    "Тел.: +7 900 000-01-07",
    "E-mail: zakupki@test-pokupatel.example",
]
NEW_BANK = {"rs": "40702810000000000222", "bank": "АО «Банк Новый Образец»", "bik": "044500222",
            "ks": "30101810000000000222"}
OLD_BANK = {"rs": "40702810000000000111", "bank": "АО «Банк Старый Образец»", "bik": "044500111",
            "ks": "30101810000000000111"}

CARD_ROWS = [
    ("Полное наименование", "Общество с ограниченной ответственностью «Образец Поставка»"),
    ("Сокращённое наименование", "ООО «Образец Поставка»"),
    ("ИНН", "7700000086"),
    ("КПП", "770001001"),
    ("ОГРН", "1000000000000"),
    ("Юридический адрес", "101000, г. Москва, ул. Образцовая, д. 5, офис 12"),
    ("Почтовый адрес", "101000, г. Москва, ул. Образцовая, д. 5, офис 12"),
    ("Расчётный счёт", NEW_BANK["rs"]),
    ("Банк", NEW_BANK["bank"] + ", г. Москва"),
    ("БИК", NEW_BANK["bik"]),
    ("Корреспондентский счёт", NEW_BANK["ks"]),
    ("Генеральный директор", "Образцов Иван Иванович, действует на основании Устава"),
    ("Телефон", "+7 900 000-00-86"),
    ("E-mail", "info@obrazec-postavka.example"),
]


# ------------------------------------------------------------------ numbering

def _add_numbering(doc):
    numbering = doc.part.numbering_part.element
    W = C.OD_W_NS
    an = etree.SubElement(numbering, qn("w:abstractNum"))
    an.set(qn("w:abstractNumId"), str(ABSTRACT_ID))
    mlt = etree.SubElement(an, qn("w:multiLevelType"))
    mlt.set(qn("w:val"), "multilevel")
    for ilvl in range(9):
        lvl = etree.SubElement(an, qn("w:lvl"))
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, val in (("w:start", "1"), ("w:numFmt", "decimal")):
            e = etree.SubElement(lvl, qn(tag))
            e.set(qn("w:val"), val)
        sfx = etree.SubElement(lvl, qn("w:suff"))
        sfx.set(qn("w:val"), "space")
        txt = etree.SubElement(lvl, qn("w:lvlText"))
        txt.set(qn("w:val"), "".join(f"%{k + 1}." for k in range(ilvl + 1)))
        jc = etree.SubElement(lvl, qn("w:lvlJc"))
        jc.set(qn("w:val"), "left")
        ppr = etree.SubElement(lvl, qn("w:pPr"))
        ind = etree.SubElement(ppr, qn("w:ind"))
        ind.set(qn("w:left"), "0")
        ind.set(qn("w:firstLine"), "0" if ilvl == 0 else "709")
        if ilvl == 0:
            rpr = etree.SubElement(lvl, qn("w:rPr"))
            etree.SubElement(rpr, qn("w:b"))
    # schema order: all w:abstractNum before the first w:num
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        numbering.remove(an)
        first_num.addprevious(an)
    num = etree.SubElement(numbering, qn("w:num"))
    num.set(qn("w:numId"), str(NUM_ID))
    ref = etree.SubElement(num, qn("w:abstractNumId"))
    ref.set(qn("w:val"), str(ABSTRACT_ID))
    _ = W


def _set_numpr(paragraph, ilvl):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), str(ilvl))
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), str(NUM_ID))
    numpr.append(il)
    numpr.append(ni)
    # w:numPr goes after w:pStyle/w:keepNext/w:keepLines/w:pageBreakBefore/w:framePr/w:widowControl
    successors = ("w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
                  "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                  "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                  "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                  "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
    for child in ppr:
        if child.tag in {qn(s) for s in successors}:
            child.addprevious(numpr)
            break
    else:
        ppr.append(numpr)


# ------------------------------------------------------------------ images

def _logo_png() -> bytes:
    W, H = 700, 260
    img = Image.new("RGBA", (W, H), (255, 255, 255, 0))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((8, 8, 252, 252), radius=46, fill=(28, 92, 150, 255))
    d.rounded_rectangle((26, 26, 234, 234), radius=34, outline=(255, 255, 255, 255), width=6)
    f = C.ff_pil_font("DejaVuSans-Bold", 112)
    d.text((130, 124), "ОП", font=f, fill=(255, 255, 255, 255), anchor="mm")
    f2 = C.ff_pil_font("DejaVuSans-Bold", 62)
    d.text((270, 92), "ОБРАЗЕЦ", font=f2, fill=(190, 32, 32, 255), anchor="lm")
    f3 = C.ff_pil_font("DejaVuSans", 40)
    d.text((272, 168), "поставка", font=f3, fill=(28, 92, 150, 255), anchor="lm")
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=False, compress_level=9)
    return buf.getvalue()


# ------------------------------------------------------------------ docx

def _cell_lines(cell, lines, *, bold_first=False, size=10):
    cell.paragraphs[0].text = ""
    for k, line in enumerate(lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(line)
        r.font.size = Pt(size)
        if bold_first and k == 0:
            r.bold = True


def _build_docx(path: pathlib.Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)
    C.od_docx_base_style(doc, font="Times New Roman", size_pt=12, space_after_pt=4, line_spacing=1.1)
    _add_numbering(doc)

    # header: logo + company + specimen mark, bottom rule
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.9), WD_TAB_ALIGNMENT.RIGHT)
    for twips in (4680, 9360):          # the default "Header" style stops (centre, right) would catch the tab
        hp.paragraph_format.tab_stops.add_tab_stop(twips * 635, WD_TAB_ALIGNMENT.CLEAR)
    hp.add_run().add_picture(io.BytesIO(_logo_png()), height=Cm(1.1))
    r = hp.add_run("   ООО «Образец Поставка»")
    r.font.size = Pt(10)
    r.bold = True
    r.font.color.rgb = RGBColor(28, 92, 150)
    hp.add_run("\t")
    C.od_docx_specimen_header(sec, paragraph=hp, align="left")
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "4"), ("w:color", "1C5C96")):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    hppr = hp._p.get_or_add_pPr()
    after_bdr = {qn(t) for t in ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:spacing", "w:ind", "w:jc", "w:rPr")}
    nxt = next((ch for ch in hppr if ch.tag in after_bdr), None)
    if nxt is not None:
        nxt.addprevious(pbdr)
    else:
        hppr.append(pbdr)

    # footer: page number field
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = fp.add_run("Страница ")
    rr.font.size = Pt(9)
    C.od_docx_add_field(fp, "PAGE", "1", size_pt=9)

    # title block
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(TITLE)
    r.bold = True
    r.font.size = Pt(13)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
    p.paragraph_format.space_after = Pt(10)
    p.add_run(f"{PLACE_DATE[0]}\t{PLACE_DATE[1]}")
    p = doc.add_paragraph(PREAMBLE)
    p.paragraph_format.first_line_indent = Cm(1.25)

    layout = []
    for s_no, (head, clauses) in enumerate(SECTIONS, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = True
        _set_numpr(p, 0)
        r = p.add_run(head)
        r.bold = True
        layout.append({"label": f"{s_no}.", "ilvl": 0, "text": head})
        for c_no, text in enumerate(clauses, start=1):
            p = doc.add_paragraph()
            _set_numpr(p, 1)
            p.add_run(text)
            layout.append({"label": f"{s_no}.{c_no}.", "ilvl": 1, "text": text})

    # requisites (unnumbered, as is usual for the last block)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.keep_with_next = True
    r = p.add_run("АДРЕСА, БАНКОВСКИЕ РЕКВИЗИТЫ И ПОДПИСИ СТОРОН")
    r.bold = True
    t = doc.add_table(rows=3, cols=2)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in t.rows:
        row.cells[0].width = Cm(8.5)
        row.cells[1].width = Cm(8.5)
    _cell_lines(t.cell(0, 0), ["ПОСТАВЩИК"], bold_first=True, size=11)
    _cell_lines(t.cell(0, 1), ["ПОКУПАТЕЛЬ"], bold_first=True, size=11)
    _cell_lines(t.cell(1, 0), [v for _k, v in SUPPLIER_OLD], bold_first=True)
    _cell_lines(t.cell(1, 1), BUYER, bold_first=True)
    _cell_lines(t.cell(2, 0), ["Генеральный директор", "", "_______________ / И. И. Образцов /", "М. П."])
    _cell_lines(t.cell(2, 1), ["Генеральный директор", "", "_______________ / Т. Т. Тестов /", "М. П."])
    for ri, row in enumerate(t.rows):      # keep the requisites table on one page
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if ri < len(t.rows) - 1:
            for cell in row.cells:
                for cp in cell.paragraphs:
                    cp.paragraph_format.keep_with_next = True
    C.docx_save_deterministic(doc, path, when=WHEN, author="ООО «Образец Поставка» (ОБРАЗЕЦ)")
    C.ooxml_fix_core_dates(path, when=WHEN)
    return layout


# ------------------------------------------------------------------ card pdf

def _build_card(path: pathlib.Path):
    from reportlab.lib.colors import Color

    C.ff_register_rl_fonts("LiberationSans", "LiberationSans-Bold")
    W, H = C.FF_A4
    c = C.ff_rl_canvas(str(path), title="Карточка предприятия ООО «Образец Поставка» (ОБРАЗЕЦ)",
                       subject="Реквизиты организации (ОБРАЗЕЦ, вымышленные данные)")
    # diagonal specimen mark as outlines (no text layer)
    c.saveState()
    c.setFillColor(Color(0.78, 0.12, 0.12, alpha=0.13))
    c.translate(W / 2, H / 2 - 40)
    c.rotate(35)
    C.rl_draw_text_outlines(c, "ОБРАЗЕЦ / SPECIMEN", C.ff_font_file("DejaVuSans-Bold"), 54, 0, 0)
    c.restoreState()

    c.setFillColorRGB(0.11, 0.36, 0.59)
    c.rect(56, H - 92, W - 112, 36, stroke=0, fill=1)
    c.setFillColorRGB(1, 1, 1)
    c.setFont("LiberationSans-Bold", 15)
    c.drawString(68, H - 80, "КАРТОЧКА ПРЕДПРИЯТИЯ")
    c.setFont("LiberationSans-Bold", 10)
    c.drawRightString(W - 68, H - 80, "ОБРАЗЕЦ / SPECIMEN")
    c.setFillColorRGB(0, 0, 0)
    c.setFont("LiberationSans", 10.5)
    c.drawString(56, H - 116, "Сведения для заключения договоров и проведения расчётов")
    c.drawString(56, H - 131, "Реквизиты действуют с 01.10.2026 (смена обслуживающего банка).")

    x0, x1, x2 = 56, 236, W - 56
    y = H - 156
    row_h = 24
    c.setLineWidth(0.6)
    c.setStrokeColorRGB(0.45, 0.45, 0.5)
    rows_layout = []
    for k, (label, value) in enumerate(CARD_ROWS):
        y_top, y_bot = y, y - row_h
        if k % 2 == 0:
            c.setFillColorRGB(0.95, 0.96, 0.98)
            c.rect(x0, y_bot, x2 - x0, row_h, stroke=0, fill=1)
            c.setFillColorRGB(0, 0, 0)
        c.rect(x0, y_bot, x2 - x0, row_h, stroke=1, fill=0)
        c.line(x1, y_bot, x1, y_top)
        c.setFont("LiberationSans-Bold", 10)
        c.drawString(x0 + 6, y_bot + 8, label)
        size = 10.5
        while c.stringWidth(value, "LiberationSans", size) > (x2 - x1 - 12) and size > 7.5:
            size -= 0.25
        c.setFont("LiberationSans", size)
        c.drawString(x1 + 6, y_bot + 8, value)
        rows_layout.append({"label": label, "value": value, "rect_pt": [x0, y_bot, x2, y_top]})
        y = y_bot
    c.setFont("LiberationSans", 10)
    c.drawString(56, y - 30, "Генеральный директор  _______________  И. И. Образцов")
    c.setFont("LiberationSans", 8)
    c.setFillColorRGB(0.4, 0.4, 0.45)
    c.drawString(56, 40, "ОБРАЗЕЦ / SPECIMEN — все сведения вымышлены, документ создан для тестирования.")
    c.showPage()
    c.save()
    return rows_layout


def build(out_dir: pathlib.Path, truth_dir: pathlib.Path) -> list[str]:
    out_dir = pathlib.Path(out_dir)
    truth_dir = pathlib.Path(truth_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / DOCX_NAME
    card_path = out_dir / CARD_NAME
    layout = _build_docx(docx_path)
    card_rows = _build_card(card_path)

    full_text = " ".join(x["text"] for x in layout)
    assert "конфиденциальн" not in full_text.lower() and "конфиденциальн" not in PREAMBLE.lower()
    assert sum(CLAUSE_62[:30] in x["text"] for x in layout) == 1
    assert sum(x["text"].startswith("Уплата неустойки") for x in layout) == 1
    lab = {x["text"]: x["label"] for x in layout}
    assert lab[CLAUSE_43] == "4.3." and lab[CLAUSE_62] == "6.2." and lab[CLAUSE_63] == "6.3."

    truth = {
        "case": CASE,
        "files": {"docx": DOCX_NAME, "card": CARD_NAME},
        "numbering": {"numId": NUM_ID, "abstractNumId": ABSTRACT_ID, "lvlText": ["%1.", "%1.%2."],
                      "suffix": "space", "headings_ilvl": 0, "clauses_ilvl": 1,
                      "note": "w:numPr is set directly on every numbered paragraph (not via style)"},
        "clauses": layout,
        "edits_requested": {
            "payment_term": {"label": "4.3.", "old": CLAUSE_43,
                             "old_fragment": "10 (десяти) календарных", "new_fragment": "5 (пяти) банковских"},
            "new_clause_after": {"label": "6.2.", "text": CLAUSE_62, "new_label": "6.3.",
                                 "renumbered": {"6.3.": "6.4.", "6.4.": "6.5."},
                                 "renumbered_texts": {"6.4.": CLAUSE_63}},
            "requisites": {"old": OLD_BANK, "new": NEW_BANK, "untouched": ["7700000086", "770001001", "1000000000000"],
                           "supplier_cell_lines_old": [v for _k, v in SUPPLIER_OLD]},
        },
        "header_footer": {"header": "logo PNG (word/media/image1.png) + «ООО «Образец Поставка»» + «ОБРАЗЕЦ / SPECIMEN»",
                          "footer": "«Страница {PAGE}»"},
        "card_rows": card_rows,
    }
    tp = truth_dir / "dogovor_postavki.json"
    C.ff_write_json(tp, truth)
    return [str(docx_path), str(card_path), str(tp)]


if __name__ == "__main__":
    root = HERE.parent
    for p in build(root / "out" / CASE, root / "truth"):
        print(p)
