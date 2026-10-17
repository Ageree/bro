"""Office kinds: DOCX compared with the case's input, a forbidden XLSX total, PPTX text that collides or does not fit."""

import io
import pathlib
import shutil
import tempfile
import unittest

from check import evaluate
from tests.runs import make_run
from tests.samples import context, suite_check

CASE = {"id": "f-test", "turns": [], "fixtures": []}


def docx_bytes(text: str) -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx_bytes(rows) -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def pptx_bytes(*slides) -> bytes:
    """slides: lists of text boxes (text, left, top, width, height in inches) on blank slides."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    for boxes in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        for text, left, top, width, height in boxes:
            slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame.text = text
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


class OfficeTest(unittest.TestCase):
    def setUp(self):
        self.root = pathlib.Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.root, True)

    def run_check(self, item, data: bytes, name: str):
        case_dir = make_run(self.root / "out", "f-test", files=[(0, data, name)])
        ctx = context(CASE, case_dir, self.root / "fixtures")
        try:
            outcome, _ = evaluate(ctx, item)
        finally:
            ctx.cleanup()
        return outcome


class PartsUnchangedTest(unittest.TestCase):
    def test_parts_the_input_does_not_have_prove_nothing(self):
        root = pathlib.Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, root, True)
        case = {"id": "f-test", "turns": [{"text": "x", "files": ["dogovor.docx"]}], "fixtures": []}
        (root / "fixtures" / "f-test").mkdir(parents=True)
        (root / "fixtures" / "f-test" / "dogovor.docx").write_bytes(docx_bytes("Договор"))
        case_dir = make_run(root / "out", "f-test", files=[(0, docx_bytes("Договор, правки"), "dogovor.docx")])
        ctx = context(case, case_dir, root / "fixtures")
        outcome, _ = evaluate(ctx, {"kind": "docx_parts_unchanged", "params": {"parts": ["word/header*.xml"]}})
        self.assertEqual(outcome.status, "error")


@unittest.skipUnless(shutil.which("soffice"), "needs LibreOffice: xlsx values are recalculated")
class ForbiddenTotalTest(OfficeTest):
    """f-bank-statement-pdf-categories: the statement's «Всего списаний» equals the spending total with the transfer to
    a deposit left in; a reconciliation row may cite it, the answer's own total may not."""

    def setUp(self):
        super().setUp()
        _, self.item = suite_check("f-bank-statement-pdf-categories", "Итог расходов без переводов")
        self.spend = self.item["params"]["value"]
        self.statement = max(self.item["params"]["forbid_values"])

    def test_a_reconciliation_row_may_cite_the_statements_total(self):
        for label in ("Списания по выписке", "Всего списаний по выписке (со стр. 1 PDF)", "Сверка: всего списаний"):
            with self.subTest(label=label):
                rows = [("Итого расходов", self.spend), (label, self.statement)]
                outcome = self.run_check(self.item, xlsx_bytes(rows), "itogi.xlsx")
                self.assertEqual(outcome.status, "pass", outcome.detail)

    def test_the_answers_own_wrong_total_fails(self):
        for label in ("Итого списаний", "Итого расходов"):
            with self.subTest(label=label):
                rows = [("Списания по выписке", self.statement), ("Итого расходов", self.spend), (label, self.statement)]
                outcome = self.run_check(self.item, xlsx_bytes(rows), "itogi.xlsx")
                self.assertEqual(outcome.status, "fail")
                self.assertIn("R3C2", outcome.detail)


class TopListTest(OfficeTest):
    """f-csv: the summary keeps the source rows beside the top; their column names every product many times."""

    ITEM = {"kind": "xlsx_top_list", "params": {"expected": ["Сыр", "Кофе", "Чай", "Мёд"], "ordered_prefix": 2}}

    def test_the_top_is_found_beside_the_source_rows(self):
        source = ["Мёд", "Чай", "Кофе", "Сыр"] * 50
        top = ["Сыр", "Кофе", "Чай", "Мёд"]
        rows = [(name, None, top[i] if i < len(top) else None) for i, name in enumerate(source)]
        outcome = self.run_check(self.ITEM, xlsx_bytes(rows), "svodka.xlsx")
        self.assertEqual(outcome.status, "pass", outcome.detail)

    def test_a_top_out_of_order_fails(self):
        rows = [(name,) for name in ["Кофе", "Сыр", "Чай", "Мёд"]]
        outcome = self.run_check(self.ITEM, xlsx_bytes(rows), "svodka.xlsx")
        self.assertEqual(outcome.status, "fail")


class LeftOutReceiptTest(OfficeTest):
    """f-receipts: the receipt out of the trip's dates may be named in a note, not put in the table."""

    def setUp(self):
        super().setUp()
        _, self.item = suite_check("f-receipts-expense-report", "Чек 8")

    def test_a_note_naming_it_left_out_passes(self):
        rows = [("Такси", 1250), ("Примечание: чек от 20.10.2026 на 6990,00 ₽ вне дат командировки, в итог не включён",)]
        outcome = self.run_check(self.item, xlsx_bytes(rows), "otchet.xlsx")
        self.assertEqual(outcome.status, "pass", outcome.detail)

    def test_its_row_fails_whatever_its_note_says(self):
        rows = [("Такси", 1250), ("Одежда-Образец (вне дат, не включён в итог)", 6990)]
        outcome = self.run_check(self.item, xlsx_bytes(rows), "otchet.xlsx")
        self.assertEqual(outcome.status, "fail")


class SplitOperationTest(OfficeTest):
    """f-bank-statement: the operation split between pages is one row; its marker text may be left out."""

    def setUp(self):
        super().setUp()
        _, self.item = suite_check("f-bank-statement-pdf-categories", "Операция, разорванная")

    def test_the_whole_row_without_its_marker_passes(self):
        rows = [("12.09.2026 19:53", "WILDBERRIES MOSCOW RUS", -677.13)]
        outcome = self.run_check(self.item, xlsx_bytes(rows), "vypiska.xlsx")
        self.assertEqual(outcome.status, "pass", outcome.detail)

    def test_its_sum_and_description_in_two_rows_fail(self):
        rows = [("12.09.2026 19:53", "", -677.13), ("", "WILDBERRIES MOSCOW RUS", None)]
        outcome = self.run_check(self.item, xlsx_bytes(rows), "vypiska.xlsx")
        self.assertEqual(outcome.status, "fail")


class RenderNoOverflowTest(OfficeTest):
    ITEM = {"kind": "render_no_overflow", "params": {"max_words": 90, "max_word_overlap": 0.2}}

    @unittest.skipUnless(shutil.which("soffice") and shutil.which("pdftotext"), "needs LibreOffice and poppler")
    def test_frames_may_overlap_where_their_text_does_not(self):
        deck = pptx_bytes([("Тёплый дом", 1, 1, 8, 4), ("Семьи в трудной ситуации", 1, 3, 8, 1)])
        outcome = self.run_check(self.ITEM, deck, "deck.pptx")
        self.assertEqual(outcome.status, "pass", outcome.detail)

    @unittest.skipUnless(shutil.which("soffice") and shutil.which("pdftotext"), "needs LibreOffice and poppler")
    def test_words_drawn_over_each_other_fail(self):
        deck = pptx_bytes([("Тёплый дом", 1, 1, 8, 1), ("Шаг вперёд", 1, 1, 8, 1)])
        outcome = self.run_check(self.ITEM, deck, "deck.pptx")
        self.assertEqual(outcome.status, "fail")
        self.assertIn("налезают", outcome.detail)

    def test_every_slide_over_the_word_limit_is_named(self):
        deck = pptx_bytes([("слово " * 100, 1, 1, 8, 5)], [("Итоги", 1, 1, 8, 1)], [("слово " * 95, 1, 1, 8, 5)])
        outcome = self.run_check(self.ITEM, deck, "deck.pptx")
        self.assertEqual(outcome.status, "fail")
        self.assertIn("слайд 1 — 100", outcome.detail)
        self.assertIn("слайд 3 — 95", outcome.detail)


if __name__ == "__main__":
    unittest.main()
