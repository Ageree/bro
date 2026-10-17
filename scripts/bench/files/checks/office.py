"""XLSX, DOCX and PPTX kinds.

XLSX values are always read from a copy LibreOffice recalculated (openpyxl writes formulas without cached results);
formulas from the file as sent. DOCX revisions are read from the XML: the accept view drops w:del/w:moveFrom and keeps
w:ins/w:moveTo, the reject view the other way round; a deleted (inserted) paragraph mark joins the paragraph with the next
one in the accept (reject) view.
"""

from __future__ import annotations

import datetime as dt
import fnmatch
import html
import re
import zipfile

from . import CheckError, check, failed, passed
from .context import area

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def _norm(text) -> str:
    text = "" if text is None else str(text)
    text = text.replace(" ", " ").replace(" ", " ").replace("­", "")
    return re.sub(r"\s+", " ", text).strip()


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%d.%m.%Y")
    return _norm(value)


def _number(value):
    """A cell as a number: numbers as they are, money-like strings («12 500,00», «12500,00 руб.») parsed."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        t = value.replace(" ", " ").replace(" ", " ").strip()
        m = re.fullmatch(r"(-?)\s*(\d{1,3}(?:[ ]\d{3})+|\d+)(?:[.,](\d+))?\s*(?:руб\.?|₽|р\.)?", t)
        if m:
            return float(m.group(1) + m.group(2).replace(" ", "") + ("." + m.group(3) if m.group(3) else ""))
    return None


# ------------------------------------------------------------------ workbooks
def workbook(ctx, path, *, values: bool):
    """openpyxl workbook (read only): recalculated values, or the formulas as sent."""
    import openpyxl

    key = ("wb", str(path), values)
    if key not in ctx._cache:
        source = ctx.recalculated(path) if values else path
        try:
            ctx._cache[key] = openpyxl.load_workbook(source, read_only=True, data_only=values)
        except Exception as e:
            raise CheckError(f"XLSX не открывается: {type(e).__name__}: {e}") from e
    return ctx._cache[key]


def grids(ctx, path, *, values: bool = True, sheet_regex: str | None = None, max_rows: int = 200000) -> list:
    """[(sheet name, rows as lists)] — every sheet (or those whose name matches sheet_regex; all when none does)."""
    key = ("grids", str(path), values)
    if key not in ctx._cache:
        out = []
        for ws in workbook(ctx, path, values=values).worksheets:
            rows = []
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= max_rows:
                    break
                rows.append(list(row))
            out.append((ws.title, rows))
        ctx._cache[key] = out
    sheets = ctx._cache[key]
    if sheet_regex:
        chosen = [s for s in sheets if re.search(sheet_regex, s[0])]
        return chosen or sheets
    return sheets


def _cells(rows):
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            if v is not None and v != "":
                yield r, c, v


def _at(rows, r, c):
    if 0 <= r < len(rows) and 0 <= c < len(rows[r]):
        return rows[r][c]
    return None


def _close(a, b, tol, absolute=False) -> bool:
    if a is None:
        return False
    if absolute:
        a, b = abs(a), abs(b)
    return abs(a - b) <= tol


def _cells_with(sheets, value, tol, absolute=False):
    """Every cell holding the number: (sheet name, row, column)."""
    for name, rows in sheets:
        for r, c, v in _cells(rows):
            if _close(_number(v), value, tol, absolute):
                yield name, r, c


def _find_number(sheets, value, tol, absolute=False):
    return next(_cells_with(sheets, value, tol, absolute), None)


def _is_date(v) -> bool:
    if isinstance(v, (dt.datetime, dt.date)):
        return True
    return isinstance(v, str) and re.fullmatch(r"\d{1,2}\.\d{1,2}\.\d{2,4}( \d{1,2}:\d{2}(:\d{2})?)?", v.strip()) is not None


# ------------------------------------------------------------------ xlsx kinds
@check("xlsx_find_value", "xlsx")
def xlsx_find_value(ctx, params, item):
    sheets = grids(ctx, item.path)
    tol = float(params.get("tol", 0.01))
    absolute = bool(params.get("abs_compare"))
    # A forbidden total may also be a number the source prints (a statement's «Всего списаний»): a row that cites it for
    # a reconciliation, its label (first text cell) matched by forbid_skip_row_regex, is not the answer's own total.
    # Only the label: a wrong «Итого расходов» with a note «списания …» beside it still fails.
    skip = params.get("forbid_skip_row_regex")
    rows_of = dict(sheets)
    for bad in params.get("forbid_values", []):
        for name, r, c in _cells_with(sheets, bad, tol, absolute):
            label = next((v for v in rows_of[name][r] if isinstance(v, str) and v.strip()), "")
            if skip and re.search(skip, label):
                continue
            return failed(f"есть запрещённое {bad} ({name}!R{r + 1}C{c + 1})")
    wanted = params["values"] if "values" in params else [params["value"]]
    alts = params.get("alt_values", [])
    found = []
    for value in wanted:
        hit = None
        for candidate in [value, *alts] if len(wanted) == 1 else [value]:
            hit = _find_number(sheets, candidate, tol, absolute)
            if hit:
                break
        if hit is None:
            return failed(f"нет числа {value}" + (f" (и вариантов {alts})" if alts else ""))
        if params.get("require_formula"):
            formula = _at(dict(grids(ctx, item.path, values=False))[hit[0]], hit[1], hit[2])
            if not (isinstance(formula, str) and formula.startswith("=")):
                return failed(f"{value} в {hit[0]}!R{hit[1] + 1}C{hit[2] + 1} — не формула")
        found.append(f"{value} в {hit[0]}")
    return passed("; ".join(found))


@check("table_value", "any")
def table_value(ctx, params, item):
    """The number in a sheet (as xlsx_find_value, its formula asked of a sheet only) or in a table of a Word file:
    «табличку» beside a protocol in Word may be a table in it (17.10: 25 × 2 400 = 60 000 in the .docx)."""
    from .select import group_of

    group = group_of(item)
    if group == "xlsx":
        return xlsx_find_value(ctx, params, item)
    if group != "docx":
        return failed("не таблица")
    import docx

    tol = float(params.get("tol", 0.01))
    for t, table in enumerate(docx.Document(str(item.path)).tables, start=1):
        for row in table.rows:
            for cell in row.cells:
                n = _number(cell.text.strip())
                if n is not None and abs(n - float(params["value"])) <= tol:
                    return passed(f"{params['value']} в таблице {t} документа")
    return failed(f"нет числа {params['value']} в таблицах документа")


def _header_columns(rows, pattern, max_header_rows: int = 30) -> list:
    out = []
    for r, row in enumerate(rows[:max_header_rows]):
        for c, v in enumerate(row):
            if isinstance(v, str) and re.search(pattern, v):
                out.append((r, c))
    return out


@check("xlsx_data_sheet", "xlsx")
def xlsx_data_sheet(ctx, params, item):
    allowed = params.get("rows_any_of") or [params["rows"]]
    seen = []
    for name, rows in grids(ctx, item.path):
        if params.get("date_col_regex"):
            for hr, dc in _header_columns(rows, params["date_col_regex"]):
                acs = [c for r, c in _header_columns(rows, params.get("amount_col_regex", "."))
                       if r == hr and c != dc] if params.get("amount_col_regex") else [None]
                for ac in acs:
                    data = [row for row in rows[hr + 1:] if _is_date(_at([row], 0, dc))]
                    n = len(data)
                    seen.append(f"{name}: {n}")
                    if n not in allowed:
                        continue
                    if not all(isinstance(_at([row], 0, dc), (dt.datetime, dt.date)) for row in data):
                        return failed(f"{name}: даты не датами")
                    if ac is not None and not all(isinstance(_at([row], 0, ac), (int, float)) for row in data):
                        return failed(f"{name}: суммы не числами")
                    return passed(f"{name}: {n} строк")
        else:
            n = _table_rows(rows)
            seen.append(f"{name}: {n}")
            if n in allowed:
                return passed(f"{name}: {n} строк")
    return failed("строк данных: " + ", ".join(seen[:6]) + f" — нужно {allowed}")


def _table_rows(rows) -> int:
    """Rows of the first table: after the first row with ≥ 3 texts, up to a blank row or an «итог»/«всего» row."""
    start = None
    for r, row in enumerate(rows):
        texts = [v for v in row if isinstance(v, str) and v.strip()]
        if len(texts) >= 3:
            start = r
            break
    if start is None:
        return 0
    n = 0
    for row in rows[start + 1:]:
        filled = [v for v in row if v not in (None, "")]
        if len(filled) < 2:
            break
        if any(isinstance(v, str) and re.search(r"(?i)итог|всего", v) for v in row):
            break
        n += 1
    return n


@check("xlsx_unique_values", "xlsx")
def xlsx_unique_values(ctx, params, item):
    best = None
    for name, rows in grids(ctx, item.path):
        for hr, c in _header_columns(rows, params["col_regex"]):
            values = [_norm(_at(rows, r, c)) for r in range(hr + 1, len(rows))]
            values = [v for v in values if v]
            if best is None or len(values) > len(best[1]):
                best = (name, values)
    if best is None:
        return failed(f"нет колонки {params['col_regex']}")
    n = len(set(best[1]))
    if n != params["equals"]:
        return failed(f"{best[0]}: {n} разных значений: {sorted(set(best[1]))[:8]}")
    return passed(f"{best[0]}: {n} разных значений")


def _matches(v, pattern) -> bool:
    return v is not None and re.search(pattern, _cell_text(v)) is not None


def _lines_through(rows, r, c, seen: set) -> list:
    """The row and the column through a cell, each only once per `seen`: a label repeated down a column of 120 000
    source rows read that column once per row, and the check outlived its budget (f-csv, 16.10)."""
    lines = []
    if ("row", r) not in seen:
        lines.append(rows[r] if r < len(rows) else [])
    if ("column", c) not in seen:
        lines.append([_at(rows, i, c) for i in range(len(rows))])
    seen.update({("row", r), ("column", c)})
    return lines


@check("xlsx_matrix_lookup", "xlsx")
def xlsx_matrix_lookup(ctx, params, item):
    sheets = grids(ctx, item.path, sheet_regex=params.get("sheet_regex"))
    tol = float(params.get("tol", 0.01))
    absolute = bool(params.get("abs_compare"))
    if params.get("variants"):
        for variant, wanted in params["variants"].items():
            for name, rows in sheets:
                seen = set()
                for r, c, v in _cells(rows):
                    if not _matches(v, params["row_regex"]):
                        continue
                    for line in _lines_through(rows, r, c, seen):
                        numbers = [_number(x) for x in line]
                        if all(any(_close(n, w, tol, absolute) for n in numbers) for w in wanted.values()):
                            return passed(f"{name}: вариант {variant}")
        return failed(f"нет строки «{params['row_regex']}» со всеми значениями варианта")
    missing = []
    for spec in params["cells"]:
        alts = spec.get("alt_expected", [])
        targets = [spec["expected"]] + (alts if isinstance(alts, list) else [alts])
        if not any(_lookup(sheets, spec, t, tol, absolute) for t in targets):
            missing.append(f"{spec['row_regex']}×{spec.get('col_regex', '*')}={spec['expected']}")
    if missing:
        return failed(f"не найдено {len(missing)} из {len(params['cells'])}: " + "; ".join(missing[:4]))
    return passed(f"{len(params['cells'])} значений на местах")


def _lookup(sheets, spec, target, tol, absolute) -> bool:
    for _, rows in sheets:
        row_hits = [(r, c) for r, c, v in _cells(rows) if _matches(v, spec["row_regex"])]
        if not row_hits:
            continue
        if not spec.get("col_regex"):
            seen = set()
            for r, c in row_hits:
                if any(_close(_number(x), target, tol, absolute) for line in _lines_through(rows, r, c, seen)
                       for x in line):
                    return True
            continue
        col_hits = [(r, c) for r, c, v in _cells(rows) if _matches(v, spec["col_regex"])]
        for ra, ca in row_hits:
            for rb, cb in col_hits:
                if (ra, ca) == (rb, cb):
                    continue
                for value in (_at(rows, ra, cb), _at(rows, rb, ca)):
                    if _close(_number(value), target, tol, absolute):
                        return True
    return False


@check("xlsx_lookup_values", "xlsx")
def xlsx_lookup_values(ctx, params, item):
    sheets = grids(ctx, item.path)
    tol = float(params.get("tol", 0.05))
    for variant in params["expected_any"]:
        ok = True
        for label, want in variant.items():
            hit = False
            for _, rows in sheets:
                seen = set()
                for r, c, v in _cells(rows):
                    if isinstance(v, str) and _norm(label) in _norm(v):
                        if any(_close(_number(x), want, tol) for line in _lines_through(rows, r, c, seen) for x in line):
                            hit = True
                            break
                if hit:
                    break
            if not hit:
                ok = False
                break
        if ok:
            return passed(f"{len(variant)} значений")
    return failed("значения по подписям не найдены ни в одном варианте")


@check("xlsx_cell_formula", "xlsx")
def xlsx_cell_formula(ctx, params, item):
    sheets = grids(ctx, item.path, values=False, sheet_regex=params.get("sheet_regex"))
    names = {n for n, _ in grids(ctx, item.path, values=False)}
    pattern = params.get("formula_regex", ".")
    count, foreign = 0, 0
    for name, rows in sheets:
        for _, _, v in _cells(rows):
            if isinstance(v, str) and v.startswith("=") and re.search(pattern, v):
                count += 1
                if "!" in v and any(other != name and (other in v) for other in names):
                    foreign += 1
    need = int(params.get("min_formula_cells", 1))
    if count < need:
        return failed(f"формул {count} < {need}")
    if params.get("must_reference_other_sheet") and not foreign:
        return failed(f"формул {count}, но ни одна не ссылается на другой лист")
    return passed(f"формул {count}" + (f", на другой лист {foreign}" if params.get("must_reference_other_sheet") else ""))


def _numbers(sheets) -> dict:
    out = {}
    for name, rows in sheets:
        for r, c, v in _cells(rows):
            n = _number(v)
            if n is not None:
                out[(name, r, c)] = n
    return out


@check("xlsx_live_probe", "xlsx")
def xlsx_live_probe(ctx, params, item):
    import openpyxl

    work = ctx.scratch("probe") / f"{item.sha256[:12] or item.path.stem}.xlsx"
    wb = openpyxl.load_workbook(item.path)
    if "deal_no" in params:
        target = None
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if _cell_text(cell.value) == _cell_text(params["deal_no"]):
                        target = (ws, cell.row)
                        break
                if target:
                    break
            if target:
                break
        if target is None:
            return failed(f"сделки {params['deal_no']} нет на листах")
        ws, r = target
        amount = None
        for cell in ws[r]:
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) and abs(cell.value) >= 100:
                amount = cell
        if amount is None:
            return failed(f"в строке сделки {params['deal_no']} нет суммы числом")
        before = _numbers(grids(ctx, item.path))
        amount.value = amount.value + params["delta"]
        wb.save(work)
        after = _numbers(_load_values(ctx, work))
        grown = [k for k, v in after.items() if k in before and k != (ws.title, r - 1, amount.column - 1)
                 and abs(v - before[k] - params["delta"]) <= 0.01]
        if not grown:
            return failed(f"после +{params['delta']} к сделке ничего не пересчиталось")
        return passed(f"пересчитались {len(grown)} ячеек (напр. {grown[0][0]}!R{grown[0][1] + 1}C{grown[0][2] + 1})")
    label = params["set_input_regex"]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not (isinstance(cell.value, str) and re.search(label, cell.value)):
                    continue
                for nr, nc in ((cell.row, cell.column + 1), (cell.row, cell.column + 2), (cell.row + 1, cell.column)):
                    target = ws.cell(nr, nc)
                    v = target.value
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        target.value = params["set_value"] if v < 1 else params.get("alt_set_value", params["set_value"])
                        wb.save(work)
                        hit = _find_number(_load_values(ctx, work), params["expect_value"], float(params.get("tol", 1.0)))
                        if hit:
                            return passed(f"после смены {cell.value} → {target.value}: {params['expect_value']} в {hit[0]}")
                        return failed(f"после смены «{_norm(cell.value)}» на {target.value} нет {params['expect_value']}")
    return failed(f"нет входной ячейки «{label}» с числом рядом")


def _load_values(ctx, path) -> list:
    import openpyxl

    wb = openpyxl.load_workbook(ctx.recalculated(path), read_only=True, data_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]


def _chart_parts(path) -> list:
    with zipfile.ZipFile(path) as z:
        return [z.read(n).decode("utf-8", "replace") for n in z.namelist() if re.match(r"xl/charts/chart\d*\.xml$", n)]


@check("xlsx_chart", "xlsx")
def xlsx_chart(ctx, params, item):
    charts = _chart_parts(item.path)
    if len(charts) < int(params.get("min_charts", 1)):
        return failed(f"диаграмм {len(charts)}")
    pattern = params.get("series_ref_sheet_regex")
    if pattern:
        for xml in charts:
            refs = [html.unescape(r) for r in re.findall(r"<(?:\w+:)?f>([^<]+)</(?:\w+:)?f>", xml)]
            sheets = [re.sub(r"^'|'$", "", ref.split("!")[0]).replace("''", "'") for ref in refs if "!" in ref]
            if any(re.search(pattern, s) for s in sheets):
                return passed(f"диаграмм {len(charts)}, ряды с листа {sheets[0]}")
        return failed(f"диаграмм {len(charts)}, но ряды не с листа «{pattern}»")
    return passed(f"диаграмм {len(charts)}")


@check("xlsx_row_lookup", "xlsx")
def xlsx_row_lookup(ctx, params, item):
    """The row of a cell matching find_regex (with a sum beside it). `or_row_all_of`: a row where each regex matches a
    cell of its own passes too — the operation found by what it is when its marker text was left out."""
    sheets = grids(ctx, item.path)
    for name, rows in sheets:
        for r, c, v in _cells(rows):
            if isinstance(v, str) and re.search(params["find_regex"], v):
                if not params.get("same_row_numeric") or any(_number(x) not in (None, 0.0) and not _is_date(x)
                                                              for i, x in enumerate(rows[r]) if i != c):
                    return passed(f"{name}!R{r + 1}: «{_norm(v)[:40]}»")
                return failed(f"{name}!R{r + 1}: «{_norm(v)[:40]}» без суммы в строке")
    wanted = params.get("or_row_all_of") or []
    for name, rows in sheets if wanted else []:
        for r, row in enumerate(rows):
            texts = [_cell_text(x) for x in row]
            if all(any(re.search(rx, t) for t in texts) for rx in wanted):
                return passed(f"{name}!R{r + 1}: строка без маркера, но с {', '.join(wanted)}")
    return failed(f"«{params['find_regex']}» нет")


@check("xlsx_contains", "xlsx")
def xlsx_contains(ctx, params, item):
    sheets = grids(ctx, item.path)
    texts = [_cell_text(v) for _, rows in sheets for _, _, v in _cells(rows)]
    missing = [s for s in params.get("all_of", []) if not any(_norm(s) in t for t in texts)]
    if missing:
        return failed("нет: " + ", ".join(missing))
    # A note that names the left-out document as left out («…вне дат командировки, в итог не включён») is no
    # inclusion: only the cells such a note is in are skipped, its amount in a cell of its own still counts.
    skip = params.get("none_of_skip_regex")
    kept = [t for t in texts if not (skip and re.search(skip, t))]
    hits = [s for s in params.get("none_of", []) if any(_norm(s) in t for t in kept)]
    if hits:
        return failed("есть: " + ", ".join(hits))
    if params.get("none_of_numeric") or params.get("none_of_cells"):
        # A sheet that lists what was left out («Исключённые документы») is not the act: 17.10 the decoy's payments
        # stood there, named as excluded.
        aside = params.get("skip_sheet_regex")
        for name, rows in sheets:
            if aside and re.search(aside, name):
                continue
            for r, row in enumerate(rows):
                if params.get("act_rows_only") and not (any(_is_date(v) for v in row) and any(_number(v) for v in row)):
                    continue
                for v in row:
                    n = _number(v)
                    if n is not None and any(abs(n - bad) <= 0.01 for bad in params.get("none_of_numeric", [])):
                        return failed(f"{name}!R{r + 1}: число {v} в строке акта")
                    if _cell_text(v) in params.get("none_of_cells", []):
                        return failed(f"{name}!R{r + 1}: {v} в строке акта")
    return passed("ок")


@check("xlsx_top_list", "xlsx")
def xlsx_top_list(ctx, params, item):
    """A column listing the top in order. Columns are ranked by how many of the names they hold, each counted once,
    then the shorter first: a sheet of the source rows names every product thousands of times and must not outrank
    the top itself (f-csv, 16.10). The first column that holds the top passes."""
    want = [_norm(x) for x in params["expected"]]
    candidates = []
    for name, rows in grids(ctx, item.path):
        width = max((len(r) for r in rows), default=0)
        for c in range(width):
            column = [_norm(_at(rows, r, c)) for r in range(len(rows))]
            column = [x for x in column if x]
            distinct = len(set(column) & set(want))
            if distinct >= max(1, len(want) // 2):
                candidates.append((-distinct, len(column), name, column))
    if not candidates:
        return failed("нет колонки со списком товаров")
    k = int(params.get("ordered_prefix", len(want)))
    first_failure = None
    for _, _, name, column in sorted(candidates, key=lambda x: x[:2]):
        start = next(i for i, x in enumerate(column) if x in want)
        got = [x for x in column[start:] if x in want or not x.replace(" ", "").isdigit()][: len(want)]
        if got[:k] != want[:k]:
            first_failure = first_failure or f"{name}: первые {k}: {got[:3]}…, нужно {want[:3]}…"
        elif set(got) != set(want):
            first_failure = first_failure or f"{name}: состав топа не совпал ({len(set(got) & set(want))} из {len(want)})"
        else:
            return passed(f"{name}: порядок и состав совпали")
    return failed(first_failure)


@check("xlsx_text_ok", "xlsx")
def xlsx_text_ok(ctx, params, item):
    texts = [_cell_text(v) for _, rows in grids(ctx, item.path) for _, _, v in _cells(rows) if isinstance(v, str)]
    missing = [s for s in params.get("must_contain", []) if not any(_norm(s) in t for t in texts)]
    if missing:
        return failed("нет: " + ", ".join(missing))
    bad = [t for t in texts if params.get("forbid_regex") and re.search(params["forbid_regex"], t)]
    if bad:
        return failed(f"кракозябры: «{bad[0][:40]}»")
    return passed("кириллица в порядке")


def _names_file(cell: str, old: str) -> bool:
    """The cell names the file `old`, alone or inside an archive: «старое.zip/договор.doc», «старое.zip → договор.doc»."""
    return cell == old or re.search(r"(?:[/\\>→»]|\s-)\s*" + re.escape(old) + "$", cell) is not None


@check("xlsx_inventory", "xlsx")
def xlsx_inventory(ctx, params, item):
    records = ctx.truth_doc("zip_manifest.json")["records"]
    names = {r["path"] for r in records} | {r["name"] for r in records}
    best = None
    for name, rows in grids(ctx, item.path):
        width = max((len(r) for r in rows), default=0)
        for c in range(width):
            column = [_norm(_at(rows, r, c)) for r in range(len(rows))]
            # Distinct names: a column of bare names holds «scan0001.pdf» twice (the root and the nested archive), the
            # column of paths beside it tells them apart (17.10).
            hits = len({x for x in column if x and (x in names or x.split("/")[-1] in names)})
            if best is None or hits > best[0]:
                best = (hits, name, rows, c)
    if best is None or best[0] < 5:
        return failed("нет колонки со старыми именами файлов")
    _, sheet, rows, c = best
    olds = {}
    for r in range(len(rows)):
        x = _norm(_at(rows, r, c))
        if x and (x in names or x.split("/")[-1] in names or "." in x):
            olds.setdefault(x, r)
    n = len(olds)
    if not params.get("min_rows", 0) <= n <= params.get("max_rows", 10 ** 6):
        return failed(f"{sheet}: строк {n}, нужно {params.get('min_rows')}–{params.get('max_rows')}")
    for old, pattern in params.get("must_mark", {}).items():
        r = next((r for x, r in olds.items() if _names_file(x, old)), None)
        if r is None:
            return failed(f"{sheet}: нет строки {old}")
        if not any(re.search(pattern, _cell_text(v)) for i, v in enumerate(rows[r]) if i != c):
            return failed(f"{sheet}: у {old} нет пометки «{pattern}»")
    for old in params.get("must_contain_old_names", []):
        if not any(_names_file(x, old) for x in olds):
            return failed(f"{sheet}: старое имя «{old}» не раскодировано")
    return passed(f"{sheet}: {n} строк")


@check("xlsx_last_payment", "xlsx")
def xlsx_last_payment(ctx, params, item):
    want = dt.date.fromisoformat(params["date"])
    tol = int(params.get("tol_months", 1))
    seen = []
    for name, rows in grids(ctx, item.path):
        width = max((len(r) for r in rows), default=0)
        for dc in range(width):
            dated = [r for r in range(len(rows)) if isinstance(_at(rows, r, dc), (dt.datetime, dt.date))]
            if len(dated) < 12:
                continue
            for pc in range(width):
                if pc == dc:
                    continue
                paid = [r for r in dated if (_number(_at(rows, r, pc)) or 0) > 0.005]
                if not paid:
                    continue
                last = _at(rows, paid[-1], dc)
                last = last.date() if isinstance(last, dt.datetime) else last
                months = (last.year - want.year) * 12 + last.month - want.month
                seen.append((name, len(paid), last))
                if abs(months) <= tol and abs(len(paid) - params["count"]) <= 1:
                    return passed(f"{name}: {len(paid)} платежей, последний {last:%d.%m.%Y}")
    if not seen:
        return failed("нет графика с датами")
    close = sorted(seen, key=lambda s: abs(s[1] - params["count"]))[0]
    return failed(f"ближайший график: {close[0]}, {close[1]} платежей до {close[2]:%d.%m.%Y}")


# ------------------------------------------------------------------ docx
def docx_xml(path, part: str = "word/document.xml"):
    from lxml import etree

    with zipfile.ZipFile(path) as z:
        return etree.fromstring(z.read(part))


def _q(tag: str) -> str:
    return f"{{{W}}}{tag}"


def _inside(el, tags) -> bool:
    p = el.getparent()
    while p is not None:
        if p.tag in tags:
            return True
        p = p.getparent()
    return False


def paragraph_text(p, view: str) -> str:
    """Text of a w:p in the accept or reject view."""
    gone = {_q("del"), _q("moveFrom")} if view == "accept" else {_q("ins"), _q("moveTo")}
    out = []
    for el in p.iter(_q("t"), _q("delText"), _q("tab"), _q("br")):
        if _inside(el, gone):
            continue
        if el.tag == _q("tab"):
            out.append("\t")
        elif el.tag == _q("br"):
            out.append("\n")
        elif el.tag == _q("delText") and view == "accept":
            continue
        else:
            out.append(el.text or "")
    return "".join(out)


def _mark_change(p):
    """'del' / 'ins' when the paragraph mark itself is a tracked change."""
    rpr = p.find("w:pPr/w:rPr", NS)
    if rpr is None:
        return None
    if rpr.find("w:del", NS) is not None:
        return "del"
    if rpr.find("w:ins", NS) is not None:
        return "ins"
    return None


def view_paragraphs(root, view: str) -> list:
    """Paragraph texts of the body in document order; a removed paragraph mark joins the paragraph with the next."""
    out, carry = [], ""
    removed = "del" if view == "accept" else "ins"
    for p in root.iter(_q("p")):
        text = carry + paragraph_text(p, view)
        if _mark_change(p) == removed:
            carry = text
            continue
        carry = ""
        out.append(text)
    if carry:
        out.append(carry)
    return out


def _has_revision(p) -> bool:
    for tag in ("ins", "del", "moveFrom", "moveTo", "rPrChange", "pPrChange"):
        if p.find(f".//w:{tag}", NS) is not None:
            return True
    return False


def _docx_text(path) -> str:
    return "\n".join(_norm(t) for t in view_paragraphs(docx_xml(path), "accept") if _norm(t))


@check("docx_revisions", "docx")
def docx_revisions(ctx, params, item):
    root = docx_xml(item.path)
    ins = [e for e in root.iter(_q("ins")) if e.find(".//w:r", NS) is not None]
    dels = [e for e in root.iter(_q("del")) if e.find(".//w:r", NS) is not None]
    if len(ins) < params.get("min_ins", 0) or len(dels) < params.get("min_del", 0):
        return failed(f"w:ins {len(ins)}, w:del {len(dels)}")
    revisions = list(root.iter(_q("ins"), _q("del")))
    if not revisions:
        return failed("правок нет")
    if params.get("require_author") and any(not (e.get(_q("author")) or "").strip() for e in revisions):
        return failed("у правок нет автора")
    if params.get("require_date") and any(not e.get(_q("date")) for e in revisions):
        return failed("у правок нет даты")
    return passed(f"w:ins {len(ins)}, w:del {len(dels)}, автор {revisions[0].get(_q('author'))}")


def _revised_paragraphs(root) -> list:
    return [(_norm(paragraph_text(p, "accept")), _norm(paragraph_text(p, "reject"))) for p in root.iter(_q("p"))
            if _has_revision(p)]


@check("docx_revision_contains", "docx")
def docx_revision_contains(ctx, params, item):
    """A pair counts when one revised paragraph reads `del` only when the changes are rejected and `ins` (`ins_regex`)
    only when they are accepted (issues/f-docx-tracked-edits.md): the minimal word diff and a whole replacement both pass."""
    root = docx_xml(item.path)
    revised = _revised_paragraphs(root)
    missing = []
    for pair in params.get("pairs", []):
        old, new, new_rx = pair.get("del"), pair.get("ins"), pair.get("ins_regex")
        ok = False
        for accept, reject in revised:
            if old and not (_norm(old) in reject and (_norm(old) not in accept or (new and _norm(old) in _norm(new)))):
                continue
            if new and not (_norm(new) in accept and _norm(new) not in reject):
                continue
            if new_rx and not (re.search(new_rx, accept) and not re.search(new_rx, reject)):
                continue
            ok = True
            break
        if not ok:
            missing.append(f"{old} → {new or new_rx}")
    if missing:
        return failed("нет правки: " + "; ".join(missing))
    for keep in params.get("untouched", []):
        for p in root.iter(_q("p")):
            for tag in ("ins", "del"):
                for el in p.iter(_q(tag)):
                    if _norm(keep) in _norm("".join(t.text or "" for t in el.iter(_q("t"), _q("delText")))):
                        return failed(f"«{keep}» помечено правкой")
        if _norm(keep) not in _docx_text(item.path):
            return failed(f"«{keep}» пропало")
    return passed(f"{len(params.get('pairs', []))} правок на месте")


def _ratio(a: str, b: str) -> float:
    import difflib

    return difflib.SequenceMatcher(None, a, b, autojunk=False).ratio()


def _view_text(path, view: str) -> str:
    return "\n".join(_norm(t) for t in view_paragraphs(docx_xml(path), view) if _norm(t))


@check("docx_accept_reject", "docx")
def docx_accept_reject(ctx, params, item):
    need = float(params.get("min_ratio", 0.995))
    accept = _ratio(_view_text(item.path, "accept"), _view_text(ctx.fixture_path(params["accept_equals"]), "accept"))
    reject = _ratio(_view_text(item.path, "reject"), _view_text(ctx.fixture_path(params["reject_equals"]), "accept"))
    if accept < need or reject < need:
        return failed(f"принять всё = v2 на {accept:.4f}, отклонить = v1 на {reject:.4f}")
    return passed(f"принять {accept:.4f}, отклонить {reject:.4f}")


@check("docx_false_positives", "docx")
def docx_false_positives(ctx, params, item):
    doc = ctx.truth_doc("dogovor_uslug_15_26.json")
    v1 = [p["v1_text"] for p in doc["paragraphs"] if p.get("v1_index") is not None]
    by_index = {p["v1_index"]: p["v1_text"] for p in doc["paragraphs"] if p.get("v1_index") is not None}
    wanted = [_norm(by_index[i]) for i in doc["identical_v1_indices"] if i in by_index] or [_norm(t) for t in v1]
    clean = {_norm(paragraph_text(p, "accept")) for p in docx_xml(item.path).iter(_q("p")) if not _has_revision(p)}
    n = sum(1 for t in wanted if t in clean)
    if n < int(params["min_clean"]):
        return failed(f"чистых {n} из {len(wanted)}")
    return passed(f"чистых {n} из {len(wanted)}")


@check("docx_opens", "docx")
def docx_opens(ctx, params, item):
    pdf = ctx.as_pdf(item.path)
    n = ctx.page_count(pdf)
    return passed(f"LibreOffice: {n} стр.") if n else failed("LibreOffice не дал страниц")


def _num_pr(root_styles, p):
    num = p.find("w:pPr/w:numPr", NS)
    if num is None and root_styles is not None:
        style = p.find("w:pPr/w:pStyle", NS)
        if style is not None:
            sid = style.get(_q("val"))
            for s in root_styles.iter(_q("style")):
                if s.get(_q("styleId")) == sid:
                    num = s.find("w:pPr/w:numPr", NS)
                    break
    if num is None:
        return None
    nid, lvl = num.find("w:numId", NS), num.find("w:ilvl", NS)
    return (nid.get(_q("val")) if nid is not None else None, lvl.get(_q("val")) if lvl is not None else "0")


@check("docx_inserted_paragraph", "docx")
def docx_inserted_paragraph(ctx, params, item):
    root = docx_xml(item.path)
    try:
        styles = docx_xml(item.path, "word/styles.xml")
    except KeyError:
        styles = None
    paras = list(root.find("w:body", NS).iter(_q("p")))
    for i, p in enumerate(paras[:-1]):
        if not re.search(params["after_text_regex"], _norm(paragraph_text(p, "reject"))):
            continue
        new = paras[i + 1]
        text = _norm(paragraph_text(new, "accept"))
        if not re.search(params.get("text_regex", "."), text):
            return failed(f"после пункта идёт «{text[:50]}»")
        if _mark_change(new) != "ins":
            return failed("знак абзаца нового пункта не во w:ins")
        runs = [r for r in new.iter(_q("r")) if r.find("w:t", NS) is not None]
        if not runs or any(not _inside(r, {_q("ins")}) for r in runs):
            return failed("текст нового пункта не целиком во w:ins")
        if params.get("same_numbering_as_previous") and _num_pr(styles, new) != _num_pr(styles, p):
            return failed(f"нумерация {_num_pr(styles, new)} ≠ {_num_pr(styles, p)}")
        return passed(f"«{text[:50]}» во w:ins, нумерация {_num_pr(styles, new)}")
    return failed(f"нет пункта «{params['after_text_regex']}»")


def _canonical(data: bytes) -> bytes:
    from lxml import etree

    return etree.tostring(etree.fromstring(data), method="c14n")


@check("docx_parts_unchanged", "docx")
def docx_parts_unchanged(ctx, params, item):
    original = next((n for t in ctx.case["turns"] for n in t["files"] if n.endswith(".docx")), None)
    if original is None:
        raise CheckError("у кейса нет исходного DOCX")
    with zipfile.ZipFile(ctx.fixture_path(original)) as a, zipfile.ZipFile(item.path) as b:
        names_b = set(b.namelist())
        compared = 0
        for name in a.namelist():
            if not any(fnmatch.fnmatch(name, pattern) for pattern in params["parts"]):
                continue
            if name not in names_b:
                return failed(f"нет части {name}")
            x, y = a.read(name), b.read(name)
            if name.endswith(".xml"):
                x, y = _canonical(x), _canonical(y)
            if x != y:
                return failed(f"часть {name} изменена")
            compared += 1
    if not compared:  # nothing compared proves nothing
        raise CheckError(f"в исходнике {original} нет частей {params['parts']}")
    return passed(f"{compared} частей как в исходнике")


def _tables(path) -> list:
    import docx

    return docx.Document(str(path)).tables


@check("docx_table_shape", "docx")
def docx_table_shape(ctx, params, item):
    seen = []
    for t in _tables(item.path):
        rows, cols = len(t.rows), len(t.columns)
        seen.append(f"{rows}×{cols}")
        if cols == params["cols"] and params.get("rows_min", 0) <= rows <= params.get("rows_max", 10 ** 6):
            return passed(f"таблица {rows}×{cols}")
    return failed("таблицы: " + (", ".join(seen) or "нет"))


def _numeric_cell(s) -> bool:
    return bool(re.fullmatch(r"[\d\s., ]+", str(s).strip()))


@check("docx_table_cells", "docx")
def docx_table_cells(ctx, params, item):
    from .pdf import cer

    tables = [t for t in _tables(item.path) if len(t.columns) == len(params["rows"][0])]
    if not tables:
        return failed("нет таблицы нужной ширины")
    best = None
    for t in tables:
        got = [[_norm(c.text) for c in row.cells] for row in t.rows]
        problems = []
        for want in params["rows"]:
            row = next((g for g in got if g and _norm(g[0]) == _norm(want[0])), None)
            if row is None:
                problems.append(f"нет строки {want[0]}")
                continue
            for g, w in zip(row, want):
                w = _norm(w)
                if _numeric_cell(w) and params.get("numeric_exact"):
                    if re.sub(r"\s", "", g) != re.sub(r"\s", "", w):
                        problems.append(f"строка {want[0]}: «{g}» ≠ «{w}»")
                elif cer(g, w) > float(params.get("text_max_cer", 0.03)):
                    problems.append(f"строка {want[0]}: «{g}» ≈ «{w}»")
        if best is None or len(problems) < len(best):
            best = problems
    if best:
        return failed("; ".join(best[:4]))
    return passed(f"{len(params['rows'])} строк как в оригинале")


@check("docx_text_contains", "docx")
def docx_text_contains(ctx, params, item):
    text = _docx_text(item.path)
    flat = re.sub(r"\s+", " ", text)
    missing = [s for s in params.get("all_of", []) if _norm(s) not in flat]
    patterns = params.get("regex_all", []) + ([params["regex"]] if params.get("regex") else [])
    missing += [p for p in patterns if not re.search(p, flat)]
    if missing:
        return failed("нет: " + "; ".join(missing[:6]))
    return passed("всё на месте")


@check("docx_no_images", "docx")
def docx_no_images(ctx, params, item):
    n = 0
    with zipfile.ZipFile(item.path) as z:
        for name in z.namelist():
            if re.match(r"word/(document|header\d*|footer\d*)\.xml$", name):
                xml = z.read(name).decode("utf-8", "replace")
                n += len(re.findall(r"<w:drawing\b|<v:imagedata\b", xml))
    if n > int(params.get("max_images", 0)):
        return failed(f"картинок {n}")
    return passed(f"картинок {n}")


# ------------------------------------------------------------------ pptx
def presentation(ctx, path):
    from pptx import Presentation

    key = ("pptx", str(path))
    if key not in ctx._cache:
        try:
            ctx._cache[key] = Presentation(str(path))
        except Exception as e:
            raise CheckError(f"PPTX не открывается: {type(e).__name__}: {e}") from e
    return ctx._cache[key]


def _shapes(shapes):
    for s in shapes:
        yield s
        if s.shape_type == 6:  # group
            yield from _shapes(s.shapes)


def _charts(prs) -> list:
    return [s.chart for slide in prs.slides for s in _shapes(slide.shapes) if getattr(s, "has_chart", False) and s.has_chart]


@check("pptx_slide_count", "pptx")
def pptx_slide_count(ctx, params, item):
    n = len(presentation(ctx, item.path).slides)
    if not params.get("min", 0) <= n <= params.get("max", 10 ** 6):
        return failed(f"слайдов {n}")
    return passed(f"слайдов {n}")


@check("pptx_native_charts", "pptx")
def pptx_native_charts(ctx, params, item):
    n = len(_charts(presentation(ctx, item.path)))
    if n < int(params.get("min", 1)):
        return failed(f"нативных диаграмм {n}")
    return passed(f"нативных диаграмм {n}")


def _series(prs) -> list:
    out = []
    for chart in _charts(prs):
        for plot in chart.plots:
            for s in plot.series:
                try:
                    out.append([float(v) if v is not None else None for v in s.values])
                except Exception:
                    continue
    return out


def _same_vector(got, want, scale_free: bool) -> bool:
    if len(got) != len(want) or any(v is None for v in got):
        return False
    for k in ((1.0, 0.001, 1000.0) if scale_free else (1.0,)):
        if all(abs(g - w * k) <= max(0.01 * abs(w * k), 1e-9) for g, w in zip(got, want)):
            return True
    return False


def _shares(values) -> list:
    total = sum(values)
    return sorted(v / total for v in values) if total else []


@check("pptx_chart_series", "pptx")
def pptx_chart_series(ctx, params, item):
    series = _series(presentation(ctx, item.path))
    if not series:
        return failed("в диаграммах нет рядов")
    if params.get("any_chart_series"):
        vectors = [list(v) for v in params["any_chart_series"]]
        matrix = params.get("matrix") or []
        vectors += [list(col) for col in zip(*matrix)] if matrix else []
        for s in series:
            if any(_same_vector(s, v, bool(params.get("scale_free"))) for v in vectors):
                return passed(f"ряд {s}")
        return failed(f"ни один ряд не совпал (ряды: {series[:3]})")
    sets = [params["any_chart_values_set"], *params.get("alt_values_sets", [])]
    tol = float(params.get("norm_tol", 0.002))
    for s in series:
        if any(v is None for v in s):
            continue
        for want in sets:
            if len(s) != len(want):
                continue
            if params.get("normalized"):
                if all(abs(a - b) <= tol for a, b in zip(_shares(s), _shares(want))):
                    return passed(f"доли совпали: {s}")
            elif sorted(s) == sorted(want):
                return passed(f"значения совпали: {s}")
    return failed(f"нет ряда расходов (ряды: {series[:3]})")


@check("pptx_notes", "pptx")
def pptx_notes(ctx, params, item):
    prs = presentation(ctx, item.path)
    slides = list(prs.slides)
    good = 0
    for slide in slides:
        if slide.has_notes_slide:
            words = len(slide.notes_slide.notes_text_frame.text.split())
            good += words >= int(params.get("min_words", 1))
    ratio = good / len(slides) if slides else 0
    if ratio < float(params.get("min_ratio", 0.8)):
        return failed(f"заметки на {good} из {len(slides)} слайдов")
    return passed(f"заметки на {good} из {len(slides)}")


def _slide_text(slide) -> str:
    parts = []
    for s in _shapes(slide.shapes):
        if getattr(s, "has_text_frame", False) and s.has_text_frame:
            parts.append(s.text_frame.text)
        if getattr(s, "has_table", False) and s.has_table:
            parts += [c.text for row in s.table.rows for c in row.cells]
    return "\n".join(parts)


@check("pptx_text_contains", "pptx")
def pptx_text_contains(ctx, params, item):
    text = "\n".join(_slide_text(s) for s in presentation(ctx, item.path).slides)
    if re.search(params["regex"], _norm(text)):
        return passed("найдено")
    return failed(f"нет «{params['regex']}»")


@check("render_no_overflow", "pptx")
def render_no_overflow(ctx, params, item):
    prs = presentation(ctx, item.path)
    sw, sh = prs.slide_width, prs.slide_height
    max_words = int(params.get("max_words", 90))
    wordy = []
    for n, slide in enumerate(prs.slides, start=1):
        # Words as the kit's own check counts them (brokit look.word_count): runs between spaces with a letter or a
        # digit, so «•», «—», «%», «₽» are not words (17.10: 74 words of a slide counted 94).
        words = sum(1 for token in _slide_text(slide).split() if any(ch.isalnum() for ch in token))
        if words > max_words:
            wordy.append(f"слайд {n} — {words}")
        for s in _shapes(slide.shapes):
            if s.left is None or s.width is None:
                continue
            if s.left < -sw * 0.01 or s.top < -sh * 0.01 or s.left + s.width > sw * 1.01 or s.top + s.height > sh * 1.01:
                return failed(f"слайд {n}: «{s.name}» выходит за слайд")
    if wordy:
        return failed(f"больше {max_words} слов: " + ", ".join(wordy))
    # Text frames are no measure of overlap: python-pptx leaves a frame as tall as it was drawn, not as its text, so frames
    # of a valid slide overlap; what the reader sees collide is the rendered words.
    limit = float(params.get("max_word_overlap", 0.2))
    pdf = ctx.as_pdf(item.path)
    for page in range(1, ctx.page_count(pdf) + 1):
        box = ctx.geometry(pdf, page).box
        words = ctx.words(pdf, page)
        for w in words:
            if w.rect[0] < box[0] - 1 or w.rect[2] > box[2] + 1 or w.rect[1] < box[1] - 1 or w.rect[3] > box[3] + 1:
                return failed(f"стр. {page}: «{w.text}» за краем слайда")
        words = sorted(words, key=lambda w: w.rect[0])
        for i, a in enumerate(words):
            for b in words[i + 1:]:
                if b.rect[0] >= a.rect[2]:
                    break
                dx = min(a.rect[2], b.rect[2]) - b.rect[0]
                dy = min(a.rect[3], b.rect[3]) - max(a.rect[1], b.rect[1])
                # Neighbouring lines and kerned neighbours touch by a fraction of a point.
                if dx <= 1 or dy <= 1:
                    continue
                small = min(area(a.rect), area(b.rect)) or 1
                if dx * dy / small > limit:
                    return failed(f"стр. {page}: «{a.text}» и «{b.text}» налезают друг на друга на {dx * dy / small:.0%}")
    return passed(f"{len(prs.slides)} слайдов без переполнения")
